"""Workbook fixture teknis: bukan usulan nilai siswa atau data penelitian."""
import copy,json,subprocess,sys
from pathlib import Path

from openpyxl import load_workbook
import pytest

from aes.io import sha256


def transfer_fixture(df,cfg):
    traits=['isi','organisasi','kebahasaan','mekanik']
    continuations={row.essay_id:{'added_text':' tambahan fixture AI','reference_essay_id':'TECHNICAL_REFERENCE'}
                   for row in df.itertuples() if row.text_completeness=='continuation_unassigned'}
    return {'source_hashes':{'csv':sha256(cfg['csv']),'rubric':sha256(cfg['rubric'])},
      'timestamp_bangkok':'TECHNICAL_DATE','audit_rows':[['scope','Fixture teknis saja']],
      'reviews':[[row.essay_id,t,3,'Alasan fixture teknis','Batas level fixture teknis',
                  ['tambahan fixture AI'] if row.essay_id in continuations else ['fixture teknis'],
                  'Review manusia tetap diperlukan.'] for row in df.itertuples() for t in traits],
      'ai_continuations':continuations,
      'assessment_scopes':{i:{'mechanic_transcription_simulation':True,'teacher_validated':False} for i in df.essay_id}}


def run_export(cfg,transfer,tmp_path):
    path=tmp_path/'private_transfer.json';path.write_text(json.dumps(transfer))
    out=tmp_path/'technical_review.xlsx'
    script=Path(__file__).resolve().parents[1]/'scripts/export_review_excel.py'
    result=subprocess.run([sys.executable,str(script),'--csv',cfg['csv'],'--rubric',cfg['rubric'],
        '--review',str(path),'--output',str(out)],capture_output=True,text=True)
    return result,out


def test_review_opens_numeric_summary_and_keeps_reconstruction_out_of_teacher_source(technical_source,tmp_path):
    df,_,cfg=technical_source
    hashes={key:sha256(cfg[key]) for key in ['csv','rubric']}
    result,out=run_export(cfg,transfer_fixture(df,cfg),tmp_path)
    assert result.returncode==0,result.stderr
    book=load_workbook(out)
    assert book.active.title=='Ringkasan_AI'
    for row in book['Ringkasan_AI'].iter_rows(min_row=2):
        assert [c.value for c in row[3:7]]==[3,3,3,3]
    source_sheet=book['Data_asli']
    cols=[c.value for c in source_sheet[1]]
    for i,row in enumerate(df.itertuples(),2):
        assert source_sheet.cell(i,cols.index('text_raw')+1).value==row.text_raw
        for title in ['Data_asli','Untuk_CSV_Guru']:
            for trait in ['isi','organisasi','kebahasaan','mekanik']:
                assert book[title].cell(i,cols.index('score_'+trait)+1).value is None
    recon=list(book['Rekonstruksi_AI'].values)[1:]
    assert len(recon)==2
    for essay,original,added,combined,reference,status,verified in recon:
        assert combined==original+added and verified=='false'
    assert {key:sha256(cfg[key]) for key in hashes}==hashes


@pytest.mark.parametrize('removed,reason',[('assessment_scopes','simulasi eksplisit'),('ai_continuations','Teks tidak lengkap')])
def test_review_refuses_to_disguise_unverified_or_missing_source(technical_source,tmp_path,removed,reason):
    df,_,cfg=technical_source
    transfer=transfer_fixture(df,cfg);transfer.pop(removed)
    result,out=run_export(cfg,transfer,tmp_path)
    assert result.returncode!=0 and reason in result.stderr
    assert not out.exists()
