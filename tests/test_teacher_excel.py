"""Workbook/inline worker memakai fixture teknis, bukan hasil penelitian siswa."""
import ast
import json
from pathlib import Path
import subprocess
import sys

import nbformat
import numpy as np
import pandas as pd
from openpyxl import load_workbook
import pytest
import torch

from aes.teacher_score import data, experiment, lesson, models, vectors
from aes.teacher_score.excel import ExcelJournal, vector_frame, write_excel
from test_teacher_score import dataset, encoders, TechnicalTokenizer


def test_excel_all_coordinates_strings_and_missing_coverage(tmp_path):
    values = np.arange(2*1024,dtype=np.float32).reshape(2,1024)/100
    frame = vector_frame(values,['fixture_a','fixture_b'])
    frame.insert(1,'word',['=DANGEROUS()','<script>technical</script>'])
    frame.iloc[1,-1] = np.nan
    path=tmp_path/'full.xlsx'
    write_excel(path,{'BGE_full':frame},{'status':'TECHNICAL FIXTURE ONLY'})
    book=load_workbook(path,read_only=True,data_only=False)
    sheet=book['BGE_full'];rows=list(sheet.iter_rows(max_col=1026))
    assert sheet.max_column is None or sheet.max_column == 1026
    assert len(rows[0])==1026 and rows[0][-1].value=='v_1023'
    assert rows[1][1].value=='=DANGEROUS()' and rows[1][1].data_type=='s'
    assert rows[1][-1].value==pytest.approx(float(values[0,-1]))
    assert rows[2][-1].value is None
    book.close()


def test_export_excel_complete_pipeline(dataset,encoders,tmp_path,monkeypatch):
    cfg,bundle=dataset
    monkeypatch.setattr(lesson,'encoder',encoders)
    monkeypatch.setattr(lesson,'tokenizer',lambda *a,**k:TechnicalTokenizer())
    journal=ExcelJournal(tmp_path/'excel',{'status':'TECHNICAL FIXTURE ONLY','source':'synthetic workbook'})
    a,b=lesson.export_tokens(bundle,cfg,journal)
    assert set(a.row_id)==set(r['row_id'] for r in bundle.rows) and b.attention_mask.isin([0,1]).all()
    before=lesson.roberta_pretrained_vectors(bundle,cfg,journal)
    assert before.shape==(24,768)
    bge=models.bge_features(bundle,cfg,tmp_path)
    manifest=data.make_split(bundle,cfg)
    data.dump(tmp_path/'split_manifest.json',manifest)
    events=[]
    def observer(kind,scenario,result):
        events.append((kind,scenario))
        lesson.export_training_event(kind,scenario,result,bundle,cfg,journal)
    report=experiment.train_all(bundle,manifest,bge,cfg,tmp_path,observer=observer)
    assert sum(kind=='cv_fold' for kind,_ in events)==15
    assert sum(kind=='refit' for kind,_ in events)==3
    assert sum(kind=='test' for kind,_ in events)==3
    lesson.export_completed_training(report,bundle,cfg,tmp_path,journal)
    lesson.export_final_features(bundle,bge,cfg,tmp_path,journal)
    result=vectors.vector_lab('alpha technical alpha. Another sentence.',cfg,tmp_path)
    lesson.export_word_sentence_vectors(result,journal,'07_words')
    for role,dimension in [('roberta',768),('bge',1024)]:
        path=journal.folder/'07_words.xlsx'
        word=pd.read_excel(path,sheet_name=role+'_word_contextual')
        with np.load(result[role]['npz'],allow_pickle=False) as full:
            np.testing.assert_allclose(word.filter(regex='^v_').to_numpy(),full['word_contextual'],atol=1e-6)
        assert len(word.filter(regex='^v_').columns)==dimension
    fusion=pd.read_excel(journal.folder/'14_S3_fitur_setelah_training.xlsx',sheet_name='concat')
    roberta=pd.read_excel(journal.folder/'14_S3_fitur_setelah_training.xlsx',sheet_name='roberta')
    projection=pd.read_excel(journal.folder/'14_S3_fitur_setelah_training.xlsx',sheet_name='projection')
    np.testing.assert_allclose(fusion.filter(regex='^v_').to_numpy(),np.concatenate([
        roberta.filter(regex='^v_').to_numpy(),projection.filter(regex='^v_').to_numpy()],axis=1))
    final=pd.read_excel(journal.folder/'12_evaluasi_tiga_skenario.xlsx',sheet_name='TEST_final')
    assert list(final.Skenario)==['S1','S2','S3']
    test=pd.read_excel(journal.folder/'11_S1_test.xlsx',sheet_name='Prediksi_test')
    assert len(test)==manifest['test_rows'] and 'kategori_QWK_AC2' in test.columns
    history=pd.read_excel(journal.folder/'09_S1_fold_1.xlsx',sheet_name='Riwayat_epoch')
    assert {'train_mse_scaled','QWK','AC2','RMSE','Pearson'}<=set(history.columns)
    records=pd.read_excel(journal.folder/'00_indeks_output.xlsx',sheet_name='Daftar_file')
    assert len(records)==len(journal.records)


def test_worker_state_actual_numeric_output_and_error(tmp_path):
    root=Path(__file__).resolve().parents[1];prefix='__TECHNICAL_PROTOCOL__'
    process=subprocess.Popen([sys.executable,str(root/'scripts/aes_lesson_worker.py'),'--prefix',prefix],
                             stdin=subprocess.PIPE,stdout=subprocess.PIPE,stderr=subprocess.STDOUT,text=True)
    def event():
        while True:
            line=process.stdout.readline()
            assert line,'worker closed unexpectedly'
            if line.startswith(prefix):return json.loads(line[len(prefix):])
    def send(code):
        process.stdin.write(json.dumps({'code':code})+'\n');process.stdin.flush()
    try:
        assert event()['kind']=='ready'
        send('value=7');assert event()['kind']=='done'
        send('value+=1; show([{"answer":value}],"fixture")')
        out=event();assert out['kind']=='table' and out['rows']==[[8]]
        assert event()['kind']=='done'
        send('show([{"word":"same","occurrence_indices":[0,2],"same_token_ids":True,"initial_max_abs_difference":0.,"contextual_allclose_1e-6":False,"contextual_cosine":.9}],"Kata berulang")')
        out=event()
        assert out['columns'][1:3]==['Posisi Kata A','Posisi Kata B']
        assert out['rows'][0][1:3]==[1,3] and out['notes']
        assert 'occurrence_indices' not in out['html'] and 'Setelah Konteks' in out['html']
        assert event()['kind']=='done'
        send('1/0');assert 'ZeroDivisionError' in event()['traceback']
        send('assert value==8');assert event()['kind']=='done'
    finally:
        process.terminate();process.wait(timeout=10)


def test_notebook_actual_inline_sources_not_code_printouts():
    root=Path(__file__).resolve().parents[1]
    notebook=nbformat.read(root/'notebooks/03_AES_dataset_guru_300.ipynb',as_version=4)
    sources='\n'.join(c.source for c in notebook.cells if c.cell_type=='code')
    assert 'show_code(' not in sources and 'display(Code' not in sources
    for name in ['SingleScoreModel','Engine','fit','train_all','gwet_ac2','token_vectors','bge_features']:
        assert ('def '+name in sources) or ('class '+name in sources)
    assert 'observer=on_training_stage' in sources
    for cell in notebook.cells:
        if cell.cell_type=='code':ast.parse(cell.source.removeprefix('%%aes\n'))
