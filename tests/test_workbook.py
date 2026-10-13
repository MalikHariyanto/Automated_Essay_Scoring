"""Fixture teknis Excel; bukan label atau esai penelitian."""
import json
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
import pytest

from aes.io import load_data, sha256, targets
from aes.workbook import import_workbook


def fixture_workbook(tmp_path, df):
    book = Workbook()
    ai = book.active
    ai.title = 'Ringkasan_AI'  # Tab pertama sengaja berisi usulan 5, bukan skor guru.
    ai.append(['essay_id', 'ai_usul_isi', 'ai_usul_organisasi', 'ai_usul_kebahasaan', 'ai_usul_mekanik'])
    for essay in df.essay_id:
        ai.append([essay, 5, 5, 5, 5])
    teacher = book.create_sheet('Untuk_CSV_Guru')
    teacher.append(list(df.columns))
    for row in df.itertuples(index=False, name=None):
        teacher.append(list(row))
    path = tmp_path/'technical_review.xlsx'
    book.save(path)
    return path


def test_explicit_teacher_sheet_preserves_missing_targets_and_provenance(technical_source, tmp_path):
    df, rubric, _ = technical_source
    path = fixture_workbook(tmp_path, df)
    original = sha256(path)
    out = tmp_path/'derived.csv'
    report = import_workbook(path, out)
    reloaded = load_data(out)
    assert reloaded.equals(df)
    assert __import__('numpy').isnan(targets(reloaded, rubric)).all()
    assert report['sheet'] == 'Untuk_CSV_Guru'
    assert all(row['ai_usul_isi'] == '5' for row in report['ai_summary'])
    assert report['ai_used_as_teacher_labels'] is False
    assert sha256(path) == original
    assert import_workbook(path, out)['derived_csv_sha256'] == sha256(out)
    with pytest.raises(ValueError, match='Ringkasan_AI bukan label guru'):
        import_workbook(path, tmp_path/'wrong.csv', sheet='Ringkasan_AI')


def test_human_entered_scores_use_teacher_columns_and_formulas_rejected(technical_source, tmp_path):
    df, rubric, _ = technical_source
    path = fixture_workbook(tmp_path, df)
    book = load_workbook(path)
    sheet = book['Untuk_CSV_Guru']
    for name, score in zip(['isi', 'organisasi', 'kebahasaan', 'mekanik'], [1, 2, 3, 4]):
        sheet.cell(2, list(df.columns).index('score_'+name)+1, score)
    book.save(path)
    out = tmp_path/'entered.csv'
    import_workbook(path, out)
    assert targets(load_data(out), rubric)[0].tolist() == [1, 2, 3, 4]
    assert load_data(out).ready_for_training.tolist() == df.ready_for_training.tolist()
    sheet.cell(2, list(df.columns).index('score_isi')+1, '=5')
    book.save(path)
    with pytest.raises(ValueError, match='formula'):
        import_workbook(path, tmp_path/'formula.csv')
    with pytest.raises(ValueError, match='file pengguna'):
        import_workbook(path, Path(technical_source[2]['csv']))


def test_excel_notebook_input_selects_sheet_in_fresh_worker(technical_source, tmp_path, monkeypatch):
    df, _, cfg = technical_source
    root = Path(__file__).resolve().parents[1]
    source = fixture_workbook(tmp_path, df)
    # Nama workbook configurable; tidak bergantung pada nama CSV lama.
    nb = json.loads((root/'notebooks/01_belajar_alur_AES.ipynb').read_text())
    cell = next(''.join(c['source']) for c in nb['cells']
                if c['cell_type']=='code' and ''.join(c['source']).startswith('CSV_PATH ='))
    checkout = tmp_path/'checkout'
    (checkout/'data/source').mkdir(parents=True)
    (checkout/'configs').mkdir()
    (checkout/'configs/default.json').write_bytes((root/'configs/default.json').read_bytes())
    (checkout/'data/source/review_beda_nama.xlsx').write_bytes(source.read_bytes())
    (checkout/'data/source/rubrik_multitrait_DRAF.json').write_bytes(Path(cfg['rubric']).read_bytes())
    output = checkout/'outputs';output.mkdir()
    import os, hashlib
    namespace = dict(ROOT=checkout, OUTPUT=output, IN_COLAB=False, MODE='inspect',
                     INPUT_FORMAT='xlsx', EXCEL_FILENAME='review_beda_nama.xlsx', EXCEL_SHEET='Untuk_CSV_Guru',
                     Path=Path, os=os, json=json, hashlib=hashlib)
    def worker(args):
        args = [root/'scripts/import_workbook.py', *args[1:]]
        subprocess.run([sys.executable, *map(str,args)], check=True, capture_output=True, text=True)
    namespace['run_aes'] = worker
    exec(compile(cell,'excel_input_cell','exec'), namespace)
    actual = json.loads((output/'notebook_config.json').read_text())
    assert actual['input_workbook']['sheet'] == 'Untuk_CSV_Guru'
    assert load_data(actual['csv']).equals(df)
    assert namespace['workbook_import']['ai_summary'][0]['ai_usul_isi'] == '5'
    assert not (checkout/'data/source/dataset_aes_multitrait_DRAF.csv').exists()
