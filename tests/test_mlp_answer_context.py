"""Penyambungan konteks harus menolak ID, teks atau label yang berbeda."""
import io
import json
from pathlib import Path
import sys
from zipfile import ZipFile

from openpyxl import Workbook
import pytest

ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'scripts'))
from mlp_answer_context import build_context


def files(tmp_path,change=None):
    dataset=tmp_path/'TECHNICAL_ONLY.xlsx';w=Workbook();w.active.title='Sheet1'
    w.active.append(['soal','response_text','jawaban_referensi','rater'])
    w.active.append(['technical question','technical answer','technical key',80]);w.save(dataset)
    result=Workbook();s=result.active;s.title='Data dan Nilai Guru'
    s.append(['ID Jawaban','Jawaban Siswa','Nilai Guru']);s.append(['J001','technical answer',80])
    if change:s.cell(2,change[0]).value=change[1]
    blob=io.BytesIO();result.save(blob);archive=tmp_path/'TECHNICAL_RESULTS.zip'
    with ZipFile(archive,'w') as z:z.writestr('01_Data_dan_Tokenisasi.xlsx',blob.getvalue())
    return dataset,archive


def test_matching_rows_make_private_context_and_public_metadata(tmp_path):
    blob,metadata=build_context(*files(tmp_path));data=json.loads(blob)
    assert metadata['all_ids_answers_labels_match'] and metadata['rows']==1
    assert data['prompts'][data['links']['J001']]['reference']=='technical key'
    assert 'prompts' not in metadata and 'question' not in json.dumps(metadata)


@pytest.mark.parametrize('change',[(1,'wrong ID'),(2,'wrong answer'),(3,81)])
def test_mismatched_identity_answer_or_label_is_rejected(tmp_path,change):
    with pytest.raises(ValueError,match='tidak cocok'):build_context(*files(tmp_path,change))
