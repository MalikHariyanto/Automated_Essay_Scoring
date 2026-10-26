"""Tiga input E6; fixture teknis, bukan hasil siswa atau pretrained penuh."""
import ast
import hashlib
import inspect
import json
import os
from pathlib import Path
import subprocess
import sys
import time
import types

import nbformat
from nbclient import NotebookClient
from jupyter_client import KernelManager
from jupyter_client.kernelspec import KernelSpecManager
import numpy as np
import pytest
import torch

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'tests'))
import test_context_pair_notebook as context
from test_teacher_score import TechnicalTokenizer,make_book,encoders

NOTEBOOK=ROOT/'notebooks/experiments/AES_300_E6_Tiga_Input_CPU.ipynb'


class ThreeTechnicalTokenizer(context.PairTechnicalTokenizer):
    padding_side='right'


@pytest.fixture
def runtime(monkeypatch,tmp_path):
    monkeypatch.setattr(context,'NOTEBOOK',NOTEBOOK)
    ns=context.runtime();path=tmp_path/'TECHNICAL_FIXTURE_ONLY.xlsx';make_book(path)
    cfg=ns['config']({'dataset':str(path),'output':str(tmp_path/'results'),'expected_rows':24,
        'device':'cpu','epochs':2,'roberta_max_length':32,'roberta_overlap':4,
        'bge_max_length':32,'bge_overlap':4,'context_max_tokens':12,'question_max_tokens':4,
        'effective_batch':4,'micro_batch':1,'dropout':0.})
    ns['tokenizer']=lambda *a,**k:ThreeTechnicalTokenizer()
    return ns,cfg,ns['read_dataset'](cfg)


def test_clean_standalone_source_identity_and_deterministic_builder():
    raw=NOTEBOOK.read_bytes();n=json.loads(raw)
    assert (ROOT/'aes-learning/notebooks/experiments'/NOTEBOOK.name).read_bytes()==raw
    assert len(n['cells'])==47 and n['metadata']['aes_experiment']['reused_controls']==[]
    fingerprint=n['metadata']['aes_standalone']['code_sha256'];codes=[];identities=[]
    for cell in n['cells']:
        if cell['cell_type']!='code':continue
        assert cell['outputs']==[] and cell['execution_count'] is None
        s=''.join(cell['source']).removeprefix('%%aes\n')
        tree=ast.parse(s);codes.append(s.replace(fingerprint,'__NOTEBOOK_CODE_SHA256__'))
        for node in tree.body:
            if isinstance(node,ast.Assign) and any(isinstance(t,ast.Name) and t.id=='NOTEBOOK_CODE_SHA256' for t in node.targets):identities.append(ast.literal_eval(node.value))
        for node in ast.walk(tree):
            if isinstance(node,ast.ImportFrom):assert not node.module or not node.module.startswith('aes')
        assert not any(x in s for x in ['googleapiclient','prepare_reused_controls','DriveCheckpointCleaner'])
    assert len(codes)==30 and len(identities)==2 and all(v==fingerprint for v in identities)
    assert hashlib.sha256('\n'.join(codes).encode()).hexdigest()==fingerprint
    subprocess.run([sys.executable,str(ROOT/'scripts/build_e6_three_input_notebook.py')],check=True,cwd=ROOT,capture_output=True)
    assert NOTEBOOK.read_bytes()==raw


def test_read_dataset_requires_all_three_text_fields(runtime,tmp_path):
    from openpyxl import load_workbook
    ns,cfg,bundle=runtime
    assert bundle.summary['input_features']==['soal','response_text','jawaban_referensi']
    for name in ('soal','jawaban_referensi'):
        book=load_workbook(cfg['dataset']);sheet=book.active;header=[c.value for c in sheet[1]]
        sheet.cell(2,header.index(name)+1).value=None;p=tmp_path/(name+'.xlsx');book.save(p);book.close()
        with pytest.raises(ValueError,match=name):ns['read_dataset']({**cfg,'dataset':str(p)})


def test_bge_context_changes_cache_but_teacher_never_changes_features(runtime,encoders):
    ns,cfg,bundle=runtime
    def stable(cfg,role):
        with torch.random.fork_rng():
            torch.manual_seed(170)
            return encoders(cfg,role)
    ns['encoder']=stable
    old=ns['bge_features'](bundle,cfg,Path(cfg['output'])/'run').numpy()
    assert np.allclose(np.linalg.norm(old,axis=1),1)
    bundle.y[:]=999
    for row in bundle.rows:row['score']=-999
    np.testing.assert_array_equal(old,ns['bge_features'](bundle,cfg,Path(cfg['output'])/'run').numpy())
    bundle.rows[0]['question']='different unfamiliar context question'
    bundle.rows[0]['reference']='different reference content words'
    changed=ns['bge_features'](bundle,cfg,Path(cfg['output'])/'run').numpy()
    assert not np.array_equal(old[0],changed[0])
    np.testing.assert_array_equal(old[1:],changed[1:])
    assert len(list((Path(cfg['output'])/'cache').glob('bge_three_inputs_*.npz')))==2


@pytest.mark.parametrize('scenario',['S1','S2','S3'])
def test_three_input_forward_and_trainable_heads(runtime,encoders,scenario):
    ns,cfg,bundle=runtime;ns['encoder']=encoders
    bge=ns['bge_features'](bundle,cfg,Path(cfg['output'])/'run')
    model=ns['SingleScoreModel'](scenario,encoders(cfg,'roberta') if scenario!='S2' else None).eval()
    engine=ns['Engine'](model,bundle.texts,bge,cfg,'cpu',roberta_tokenizer=ThreeTechnicalTokenizer(),contexts=ns['contexts_from_bundle'](bundle))
    prediction,features=engine.forward([0],return_features=True)
    assert prediction.shape==(1,1)
    if scenario!='S2':
        previous=features['roberta'].detach().clone()
        engine.contexts=[('different question','unfamiliar reference words')]*len(bundle.rows)
        _,changed=engine.forward([0],return_features=True)
        assert not torch.equal(previous,changed['roberta'])
        assert all(p.requires_grad for p in model.roberta.parameters())
    prediction.square().mean().backward()
    assert sum(p.grad.abs().sum().item() for p in model.head.parameters() if p.grad is not None)>0
    groups={g['name']:g['lr'] for g in ns['optimizer_groups'](model,cfg)}
    assert groups.get('roberta')==({'S1':1e-5,'S3':5e-6}.get(scenario))
    assert groups['head']==1e-4
    if scenario=='S3':assert features['concat'].shape==(1,1024) and groups['projection']==1e-4


def test_both_token_exports_and_word_vectors_cover_answer_with_context(runtime,encoders):
    ns,cfg,bundle=runtime
    for role in ('roberta','bge'):
        tok=ThreeTechnicalTokenizer()
        frame,coverage=ns['pair_token_frames'](bundle,cfg,tok,role)
        assert set(frame['input_segment']) >= {'Soal/referensi','Jawaban siswa','Token khusus'}
        assert all(c['answer_tokens_uncovered']==0 for c in coverage)
        model=encoders(cfg,role).eval();config={**cfg,'pair_context':ns['contexts_from_bundle'](bundle)[0]}
        trace=ns['token_vectors']('answer '*55,model,tok,tok,role,config,'cpu')
        assert np.all(trace['counts']>0) and trace['contextual'].shape[0]==55
        if role=='bge':assert np.isclose(np.linalg.norm(trace['cls']),1)


def test_notebook_all_cells_and_resume_with_random_transformers(tmp_path):
    """Seluruh 30 cell, 15 CV fits dan 3 refit; bukan bobot pretrained penelitian."""
    notebook=nbformat.read(NOTEBOOK,as_version=4)
    isolated=tmp_path/'standalone';isolated.mkdir();fixture=isolated/'TECHNICAL_FIXTURE_ONLY.xlsx';make_book(fixture)
    persistent=tmp_path/'PERSISTENT_TECHNICAL_ONLY';workspace=isolated/'workspace'
    guard="""
import builtins
_original_import=builtins.__import__
def _no_project_import(name,*args,**kwargs):
    if name=='aes' or name.startswith('aes.'):raise ImportError('Standalone fixture forbids project imports')
    return _original_import(name,*args,**kwargs)
builtins.__import__=_no_project_import
"""
    tiny="""
from transformers import RobertaConfig,RobertaModel,XLMRobertaConfig,XLMRobertaModel
def encoder(cfg,role):
    config,model,dim,heads=(RobertaConfig,RobertaModel,768,12) if role=='roberta' else (XLMRobertaConfig,XLMRobertaModel,1024,16)
    with torch.random.fork_rng():
        torch.manual_seed(170)
        result=model(config(vocab_size=64,hidden_size=dim,num_hidden_layers=1,num_attention_heads=heads,
            intermediate_size=16,max_position_embeddings=64,pad_token_id=1,bos_token_id=0,eos_token_id=2,
            hidden_dropout_prob=0.,attention_probs_dropout_prob=0.),add_pooling_layer=False)
    result.requires_grad_(role=='roberta');return result
def tokenizer(*a,**k):return ThreeTechnicalTokenizer()
"""
    for cell in notebook.cells:
        if cell.cell_type!='code':continue
        if 'EPOCHS,FOLDS=25,5' in cell.source:
            cell.source=cell.source.replace('EPOCHS,FOLDS=25,5','EPOCHS,FOLDS=1,5').replace('expected_rows=300','expected_rows=24')
            cell.source=cell.source.replace('CONTEXT_MAX_TOKENS,QUESTION_MAX_TOKENS=192,64','CONTEXT_MAX_TOKENS,QUESTION_MAX_TOKENS=12,4')
            cell.source=cell.source.replace('context_max_tokens=CONTEXT_MAX_TOKENS,question_max_tokens=QUESTION_MAX_TOKENS,','context_max_tokens=CONTEXT_MAX_TOKENS,question_max_tokens=QUESTION_MAX_TOKENS,roberta_max_length=32,roberta_overlap=4,bge_max_length=32,bge_overlap=4,')
        if 'def encoder(cfg, role):' in cell.source:
            cell.source+='\n'+inspect.getsource(TechnicalTokenizer)+'\n'+inspect.getsource(context.PairTechnicalTokenizer)+'\n'+inspect.getsource(ThreeTechnicalTokenizer)+tiny
            cell.source=cell.source.replace('context.PairTechnicalTokenizer','PairTechnicalTokenizer')
        if cell.source.startswith('# Penghubung proses Python:'):
            node=next(n for n in ast.parse(cell.source).body if isinstance(n,ast.Assign))
            cell.source='WORKER_SOURCE='+repr(guard+ast.literal_eval(node.value))+'\n'
    notebook.cells.insert(0,nbformat.v4.new_code_cell(guard));path=isolated/'TECHNICAL_notebook.ipynb';nbformat.write(notebook,path)
    kernels=tmp_path/'kernels';spec=kernels/'aes_fixture';spec.mkdir(parents=True)
    (spec/'kernel.json').write_text(json.dumps({'argv':[sys.executable,'-I','-m','ipykernel_launcher','-f','{connection_file}'],'display_name':'AES fixture','language':'python'}))
    env={**os.environ,'AES300_WORKSPACE':str(workspace),'AES300_DATASET':str(fixture),
         'AES300_PERSISTENT_DIR':str(persistent),'AES300_PYTHON':sys.executable,
         'HF_HUB_OFFLINE':'1','TRANSFORMERS_OFFLINE':'1','OMP_NUM_THREADS':'2','MKL_NUM_THREADS':'2','IPYTHONDIR':str(tmp_path/'ipython')}
    env.pop('PYTHONPATH',None);env.pop('AES300_AUDIT_ONLY',None);started=time.time()
    for attempt in (1,2):
        current=nbformat.read(path,as_version=4);manager=KernelManager(kernel_name='aes_fixture',kernel_spec_manager=KernelSpecManager(kernel_dirs=[str(kernels)]))
        try:NotebookClient(current,timeout=360,km=manager,resources={'metadata':{'path':str(isolated)}}).execute(env={**env,'AES300_WORKSPACE':str(workspace if attempt==1 else isolated/'fresh_runtime')})
        finally:
            if manager.has_kernel:manager.shutdown_kernel(now=True)
        pointer=json.loads(((workspace if attempt==1 else isolated/'fresh_runtime')/'outputs/dataset300_download.json').read_text())
        run=Path(pointer['run_folder']);report=json.loads((run/'final_report.json').read_text())
        assert pointer['excel_count']==5 and len(report['cv'])==15 and report['test_evaluations_per_scenario']==1
        hashes={str(p):hashlib.sha256(p.read_bytes()).hexdigest() for p in run.glob('S*/final.pt')}
        if attempt==1:first_hashes=hashes
        else:assert hashes==first_hashes
    import pandas as pd
    output=run/'hasil_ringkas';data=pd.ExcelFile(output/'01_Data_dan_Tokenisasi.xlsx')
    for title in ['Token RoBERTa','Token BGE-M3']:
        frame=pd.read_excel(data,sheet_name=title);assert {'Soal/referensi','Jawaban siswa','Token khusus'} <= set(frame['Bagian Input'])
    for scenario in ('S1','S3'):
        vectors=json.loads((run/f'vectors/{scenario}/report.json').read_text())
        for role in ('roberta','bge'):assert vectors[role]['input_policy']=='soal/referensi berpasangan dengan jawaban' and vectors[role]['uncovered_tokens']==0
    assert not list(workspace.rglob('*.py'))
    proof={'scope':'TECHNICAL RANDOM TRANSFORMER FIXTURE; NO RESEARCH RESULTS','code_cells':30,'cv_fits':15,'refits':3,'reports':5,'resume_fresh_runtime':True,'elapsed_seconds':round(time.time()-started,2)}
    (isolated/'three_input_execution_proof.json').write_text(json.dumps(proof,indent=2));print(json.dumps(proof))
