"""Penyajian ulang fixture; angka penelitian tidak dibuat atau diubah di sini."""
import hashlib
import logging
from pathlib import Path
import zipfile

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from aes.teacher_score.excel import write_excel,vector_frame
from aes.teacher_score.models import ExpectedPoolerNotice
from aes.teacher_score.presentation import build_reports,repack_archive


@pytest.fixture
def saved_outputs(tmp_path):
    folder=tmp_path/'old';folder.mkdir()
    metadata={'evidence_scope':'TECHNICAL FIXTURE ONLY; NOT STUDENT SCORES'}
    def save(name,sheets):write_excel(folder/(name+'.xlsx'),sheets,metadata)
    ids=['row_0001','row_0002']
    save('01_konfigurasi',{'Parameter':pd.DataFrame([{'name':'epochs','value':5},{'name':'folds','value':5}])})
    save('02_dataset_dan_cleaning',{'Sebelum_sesudah':pd.DataFrame({'row_id':ids,'teks_model':['=literal fixture','synthetic second answer'],'skor_guru':[80,90]})})
    save('03_split_dan_cv',{'Split_setiap_baris':pd.DataFrame({'row_id':ids,'role':['test','development']}),
                           'Fold_setiap_baris':pd.DataFrame({'row_id':ids,'fold':[1,1],'role':['validation','train']})})
    token=pd.DataFrame({'row_id':['row_0001','row_0001'],'position':[0,1],'window':[0,0],
                        'token':['<s>','<pad>'],'token_id':[0,1],'attention_mask':[1,0]})
    save('04_tokenisasi',{'RoBERTa_input':token,'BGE_input':token.drop(columns='window')})
    roberta=np.arange(1536,dtype=float).reshape(2,768)/1000
    bge=np.arange(2048,dtype=float).reshape(2,1024)/1000
    save('05_roberta_pretrained',{'Vektor_jawaban_768':vector_frame(roberta,ids)})
    save('06_bge_frozen',{'Vektor_jawaban_1024':vector_frame(bge,ids)})
    for scenario in ['S1','S3']:
        save(f'14_{scenario}_fitur_setelah_training',{'roberta':vector_frame(roberta+.1,ids),
            'concat':vector_frame(bge+.2,ids),'dropout_eval':vector_frame(roberta,ids)})
    base={}
    for role,dim in [('roberta',768),('bge',1024)]:
        word=vector_frame(np.ones((2,dim)),[0,1]);word.insert(1,'word',['alpha','alpha']);word.insert(2,'covered',[True,True])
        sentence=vector_frame(np.zeros((1,dim)),[0]);sentence.insert(1,'text',['alpha alpha.'])
        for name in ['word_initial','word_contextual']:base[role+'_'+name]=word
        base[role+'_sentence_cls']=sentence
        base[role+'_kata_berulang']=pd.DataFrame([{'word':'alpha','occurrence_indices':[0,1],
            'same_token_ids':True,'initial_max_abs_difference':0.,'contextual_allclose_1e-6':False,'contextual_cosine':.9}])
    save('07_vektor_kata_kalimat_pretrained',base)
    for s in ['S1','S3']:save(f'15_{s}_vektor_kata_setelah_training',base)
    evaluations=[];cv=[]
    for index,scenario in enumerate(['S1','S2','S3']):
        predictions=np.array([81.125+index,88.375-index])
        save(f'11_{scenario}_test',{'Prediksi_test':pd.DataFrame({'row_id':ids,'skor_guru':[80,90],
            'prediksi_mentah':predictions,'selisih_mentah':predictions-np.array([80,90])})})
        evaluations.append({'Skenario':scenario,'QWK':.7+index/100,'AC2':.8,'RMSE':2.125+index,
                            'Pearson':.9,'out_of_range_fraction':.1,'undefined_reasons':{}})
        for fold in range(1,6):
            cv.append({'Skenario':scenario,'Fold':fold,'QWK':.75,'AC2':.8,'RMSE':2.5,'Pearson':.85})
            history=pd.DataFrame([{'epoch':1,'train_mse_scaled':.1,'RMSE':4.,'QWK':.4},
                                  {'epoch':2,'train_mse_scaled':.05,'RMSE':3.,'QWK':.6}])
            save(f'09_{scenario}_fold_{fold}',{'Riwayat_epoch':history})
        save(f'10_{scenario}_refit',{'Riwayat_epoch':pd.DataFrame([{'epoch':1,'train_mse_scaled':.04}])})
    means=pd.DataFrame([{'Skenario':s,**{k+'_'+a:v for k in ['QWK','AC2','RMSE','Pearson'] for a,v in [('mean',.75),('std',.03),('count',5)]}} for s in ['S1','S2','S3']])
    save('12_evaluasi_tiga_skenario',{'TEST_final':pd.DataFrame(evaluations),'CV_semua_fold':pd.DataFrame(cv),'CV_mean_SD_n':means})
    save('13_prediksi_satu_jawaban',{'Input_privat':pd.DataFrame([{'text':'alpha alpha.'}]),
                                   'Tiga_skenario':pd.DataFrame([{'Skenario':s,'prediksi_mentah':80.125,'kategori_salinan':80,'di_luar_rentang':False} for s in ['S1','S2','S3']])})
    return folder,roberta,bge


def test_five_reports_copy_numbers_and_combine_folds(saved_outputs,tmp_path):
    source,roberta,bge=saved_outputs
    hashes={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in source.glob('*.xlsx')}
    result=build_reports(source,tmp_path/'new')
    assert len(result['files'])==5 and not result['training_run'] and not result['metrics_recomputed']
    with zipfile.ZipFile(result['archive']) as bundle:assert len(bundle.namelist())==5
    folder=Path(result['archive']).parent
    for path in result['files']:
        book=load_workbook(path,read_only=True,data_only=False)
        assert book.sheetnames[0]=='Panduan'
        assert 'TECHNICAL FIXTURE' in str(list(book['Panduan'].values)[1][1])
        assert not any('dropout' in name or 'GELU' in name or 'normalized' in name for name in book.sheetnames)
        book.close()
    copied=pd.read_excel(folder/'02_Vektor_Jawaban.xlsx',sheet_name='RoBERTa Sebelum Latihan')
    assert copied.columns[0]=='ID Jawaban' and copied.iloc[0,0]=='J001'
    assert copied.columns[-1]=='Komponen 768'
    np.testing.assert_allclose(copied.iloc[:,1:].to_numpy(),roberta,atol=0,rtol=1e-14)
    final=pd.read_excel(folder/'04_Prediksi_dan_Evaluasi.xlsx',sheet_name='Perbandingan Model')
    assert final['RMSE (poin)'].tolist()==[2.125,3.125,4.125]
    assert final.QWK.tolist()==[.7,.71,.72]
    pairs=pd.read_excel(folder/'04_Prediksi_dan_Evaluasi.xlsx',sheet_name='Nilai Guru dan Prediksi')
    assert pairs['Prediksi S1'].tolist()==[81.125,88.375]
    assert pairs['Selisih S1'].tolist()==[1.125,-1.625]
    history=pd.read_excel(folder/'05_Riwayat_Pelatihan.xlsx',sheet_name='RoBERTa (S1)')
    assert len(history)==11 and set(history.Tahap)=={'Validasi silang','Latihan final'}
    words=pd.read_excel(folder/'03_Vektor_Kata_dan_Kalimat.xlsx',sheet_name='Kata RoBERTa Sebelum')
    assert words.columns[0]=='Urutan Kata' and words.iloc[:,0].tolist()==[1,2]
    assert 'Komponen 768' in words.columns
    tokens=pd.read_excel(folder/'01_Data_dan_Tokenisasi.xlsx',sheet_name='Token RoBERTa')
    assert len(tokens)==1 and tokens['Token'].tolist()==['<s>']
    data=load_workbook(folder/'01_Data_dan_Tokenisasi.xlsx',read_only=True,data_only=False)
    assert data['Data dan Nilai Guru']['B2'].data_type=='s'
    data.close()
    assert hashes=={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in source.glob('*.xlsx')}


def test_archive_converter_and_partial_results(saved_outputs,tmp_path):
    source,_,_=saved_outputs
    archive=tmp_path/'old.zip'
    with zipfile.ZipFile(archive,'w') as bundle:
        for p in source.glob('*.xlsx'):bundle.write(p,arcname='nested/'+p.name)
        bundle.writestr('../irrelevant.txt','must not be extracted')
    result=repack_archive(archive,tmp_path/'from_zip')
    assert len(result['files'])==5 and not (tmp_path/'irrelevant.txt').exists()
    partial=tmp_path/'partial';partial.mkdir()
    p=source/'02_dataset_dan_cleaning.xlsx';(partial/p.name).write_bytes(p.read_bytes())
    result=build_reports(partial,tmp_path/'partial_result')
    assert len(result['files'])==1 and Path(result['files'][0]).name=='01_Data_dan_Tokenisasi.xlsx'


def test_only_expected_pooler_notice_is_filtered():
    filter=ExpectedPoolerNotice()
    def record(message):return logging.LogRecord('transformers.modeling_utils',logging.WARNING,'fixture',1,message,(),None)
    expected="Some weights of the model checkpoint at BAAI/bge-m3 were not used when initializing XLMRobertaModel: ['pooler.dense.bias', 'pooler.dense.weight']"
    assert not filter.filter(record(expected))
    assert filter.filter(record(expected.replace("'pooler.dense.weight'","'encoder.layer.0.weight'")))
    assert filter.filter(record('Some weights of XLMRobertaModel were not initialized: pooler.dense.weight'))
    assert filter.filter(record('CUDA out of memory'))
