"""Technical fixtures only: no real student scores or pretrained accuracy claims."""
import copy
import hashlib
import inspect
import json
import os
from pathlib import Path
import subprocess
import sys
import types
import zipfile

import nbformat
from nbclient import NotebookClient
from jupyter_client import KernelManager
from jupyter_client.kernelspec import KernelSpecManager
import numpy as np
from openpyxl import load_workbook
import pytest

ROOT = Path(__file__).resolve().parents[1]
NOTEBOOK = ROOT/'notebooks/experiments/AES_300_BGE_M3_Ridge_Tiga_Input_Colab.ipynb'
sys.path.insert(0, str(ROOT/'tests'))
from test_teacher_score import TechnicalTokenizer, make_book
from test_context_pair_notebook import PairTechnicalTokenizer


def runtime(tmp_path):
    n = nbformat.read(NOTEBOOK, as_version=4)
    mod = types.ModuleType('ridge_three_input_technical'); sys.modules[mod.__name__] = mod
    ns = mod.__dict__; ns['ROOT'] = tmp_path
    for i in (7, 10, 11, 14, 17, 18, 21, 23, 24, 25, 26, 27):
        exec(compile(n.cells[i].source.removeprefix('%%aes\n'), '<ridge-standalone>', 'exec'), ns)
    ns['show'] = lambda *a, **k: None
    return ns


def fixture(tmp_path):
    ns = runtime(tmp_path); path = tmp_path/'TECHNICAL_FIXTURE_ONLY.xlsx'; make_book(path)
    cfg = dict(dataset=str(path), output=str(tmp_path/'results'), sheet='Sheet1', expected_rows=24,
               score_min=0, score_max=100, folds=5, test_fraction=.2, split_seed=2026, split_candidates=128,
               ridge_alphas=[.01, .1, 1., 10., 100.])
    bundle = ns['read_dataset'](cfg); manifest = ns['make_split'](bundle, cfg)
    rng = np.random.default_rng(18); X = rng.normal(size=(24, 1024)).astype(np.float32)
    X /= np.linalg.norm(X, axis=1, keepdims=True)
    return ns, cfg, bundle, manifest, X


def test_standalone_identity_deterministic_builder_and_unchanged_old_experiments():
    raw = NOTEBOOK.read_bytes(); n = nbformat.reads(raw.decode(), as_version=4); nbformat.validate(n)
    assert (ROOT/'aes-learning/notebooks/experiments'/NOTEBOOK.name).read_bytes() == raw
    identity = n.metadata.aes_standalone.code_sha256
    codes = [c.source.removeprefix('%%aes\n').replace(identity, '__RIDGE_CODE_HASH__') for c in n.cells if c.cell_type=='code']
    assert hashlib.sha256('\n'.join(codes).encode()).hexdigest() == identity
    assert all(c.execution_count is None and not c.outputs for c in n.cells if c.cell_type=='code')
    assert n.metadata.aes_experiment.input_columns == ['soal','jawaban_referensi','response_text']
    assert n.metadata.aes_experiment.BGE_frozen and n.metadata.aes_experiment.test_evaluation_enabled
    assert not n.metadata.aes_experiment.research_results_included
    code = '\n'.join(codes)
    assert 'from aes.' not in code and 'import bge_ridge_three_input' not in code
    assert 'AdamW' not in code and 'StandardScaler(' not in code
    tracked = [p for p in (ROOT/'notebooks/experiments').glob('*.ipynb') if p!=NOTEBOOK]
    before = {str(p): hashlib.sha256(p.read_bytes()).hexdigest() for p in tracked}
    subprocess.run([sys.executable, 'scripts/build_bge_ridge_three_input_notebook.py'], cwd=ROOT, check=True, capture_output=True)
    assert NOTEBOOK.read_bytes() == raw
    assert before == {str(p): hashlib.sha256(p.read_bytes()).hexdigest() for p in tracked}


def test_selection_never_uses_outer_validation_or_test_and_same_historical_split(tmp_path):
    ns, cfg, bundle, m, X = fixture(tmp_path)
    old = json.loads((ROOT/'notebooks/experiments/AES_300_BGE_Frozen_Readout_Colab.ipynb').read_text())
    legacy = dict(ns); exec(old['cells'][14]['source'].removeprefix('%%aes\n'), legacy)
    prior = legacy['make_split'](bundle, cfg)
    assert all(m[k] == prior[k] for k in ('development','test','folds'))
    f = m['folds'][0]; blocks = [fold['validation'] for fold in m['folds']]
    expected = ns['select_ridge_alpha'](X.astype(float), bundle.y, f['train'], blocks, cfg['ridge_alphas'], 'TEST_FIXTURE')
    changed = bundle.y.copy(); changed[f['validation']] = -999999.; changed[m['test']] = 999999.
    assert ns['select_ridge_alpha'](X.astype(float), changed, f['train'], blocks, cfg['ridge_alphas'], 'TEST_FIXTURE') == expected
    # Changing held-out labels must not change final alpha, OOF, or fitted coefficients.
    a = ns['run_ridge_three_input'](bundle, X, m, cfg, tmp_path/'a')
    other = copy.deepcopy(bundle); other.y[m['test']] = 100 - other.y[m['test']]
    b = ns['run_ridge_three_input'](other, X, m, cfg, tmp_path/'b')
    assert a['alpha_search'] == b['alpha_search']
    assert a['model_sha256'] == b['model_sha256']
    assert a['contract']['selected_alpha'] == b['contract']['selected_alpha']
    assert [r for r in a['predictions'] if r['partition']=='DEVELOPMENT_OOF'] == [r for r in b['predictions'] if r['partition']=='DEVELOPMENT_OOF']


def test_final_model_export_reloads_and_partitions_are_not_mixed(tmp_path):
    ns, cfg, bundle, m, X = fixture(tmp_path)
    bundle.rows[0]['text'] = '=TECHNICAL_LITERAL_NOT_A_FORMULA'
    result = ns['run_ridge_three_input'](bundle, X, m, cfg, tmp_path/'run')
    assert len(result['predictions']) == len(bundle.rows)
    for role, indices in [('DEVELOPMENT_OOF', m['development']), ('TEST_HISTORICAL_HOLDOUT', m['test'])]:
        rows = [r for r in result['predictions'] if r['partition']==role]
        assert len(rows) == len(indices) and {r['row_id'] for r in rows} == {bundle.rows[i]['row_id'] for i in indices}
    with np.load(tmp_path/'run/ridge_final_model.npz', allow_pickle=False) as model:
        p = X[m['test']].astype(float) @ model['coef'] + model['intercept']
    actual = [r['prediction'] for r in result['predictions'] if r['partition']=='TEST_HISTORICAL_HOLDOUT']
    np.testing.assert_allclose(p, actual, atol=1e-12)
    final = next(r for r in result['evaluation'] if r['partition']=='TEST_HISTORICAL_HOLDOUT')
    errors = p - bundle.y[m['test']]
    assert final['MAE'] == pytest.approx(np.abs(errors).mean())
    assert final['RMSE'] == pytest.approx(np.sqrt(np.mean(errors**2)))
    assert final['within5'] == int((np.abs(errors)<=5).sum())
    info = ns['export_ridge_three_input'](bundle, X, m, cfg, result)
    with zipfile.ZipFile(info['archive']) as z:
        proof = json.loads(z.read('manifest.json'))
        assert len([name for name in z.namelist() if name.endswith('.xlsx')]) == 5
        assert proof['test_evaluated'] and not proof['test_used_for_selection']
        assert all(hashlib.sha256(z.read(name)).hexdigest()==sha for name, sha in proof['files'].items())
    b = load_workbook(tmp_path/'run/hasil_excel'/ns['RIDGE_BOOKS'][0], data_only=False)
    header = [c.value for c in b['Data'][1]]; cell = b['Data'].cell(2, header.index('response_text')+1)
    assert cell.value == '=TECHNICAL_LITERAL_NOT_A_FORMULA' and cell.data_type == 's'; b.close()
    b = load_workbook(tmp_path/'run/hasil_excel'/ns['RIDGE_BOOKS'][2], read_only=True)
    assert {'Prediksi_Development_OOF', 'Prediksi_Test'} <= set(b.sheetnames); b.close()


def test_invalid_fold_and_raw_prediction_metrics_are_not_hidden(tmp_path):
    ns, cfg, bundle, m, X = fixture(tmp_path)
    bad = copy.deepcopy(m); bad['folds'][0]['validation'].append(m['test'][0])
    with pytest.raises(ValueError, match='Fold'):
        ns['validate_ridge_inputs'](bundle, X, bad, cfg)
    bad = copy.deepcopy(m); bad['folds'][1]['validation'] = bad['folds'][0]['validation'][:]
    with pytest.raises(ValueError): ns['validate_ridge_inputs'](bundle, X, bad, cfg)
    actual = ns['ridge_metric_row']('TECHNICAL', np.array([100., 0.]), np.array([110., -10.]), cfg)
    assert actual['RMSE'] == 10 and actual['MAE'] == 10 and actual['within5'] == 0
    assert actual['out_of_range_fraction'] == 1.


def test_all_notebook_cells_with_random_frozen_transformer(tmp_path):
    """Execute the actual cell sequence with a tiny random model, no pretrained download."""
    n = nbformat.read(NOTEBOOK, as_version=4)
    isolated = tmp_path/'standalone'; isolated.mkdir()
    dataset = isolated/'TECHNICAL_FIXTURE_ONLY.xlsx'; make_book(dataset)
    tiny = '''
from transformers import XLMRobertaConfig, XLMRobertaModel
class ThreeTechnicalTokenizer(PairTechnicalTokenizer):
    padding_side='right'
def tokenizer(*args, **kwargs):
    return ThreeTechnicalTokenizer()
def encoder(cfg, role):
    assert role=='bge'
    with torch.random.fork_rng():
        torch.manual_seed(170)
        result=XLMRobertaModel(XLMRobertaConfig(vocab_size=64,hidden_size=1024,num_hidden_layers=1,
            num_attention_heads=16,intermediate_size=16,max_position_embeddings=64,pad_token_id=1,
            bos_token_id=0,eos_token_id=2,hidden_dropout_prob=0.,attention_probs_dropout_prob=0.),add_pooling_layer=False)
    result.requires_grad_(False)
    return result
'''
    for cell in n.cells:
        if cell.cell_type!='code': continue
        if 'EXPECTED_ROWS = 300' in cell.source:
            cell.source = cell.source.replace('EXPECTED_ROWS = 300', 'EXPECTED_ROWS = 24').replace('bge_max_length=1024, bge_overlap=128', 'bge_max_length=32, bge_overlap=4').replace('context_max_tokens=192, question_max_tokens=64', 'context_max_tokens=12, question_max_tokens=4')
        if 'def encoder(cfg, role):' in cell.source:
            cell.source += '\n'+inspect.getsource(TechnicalTokenizer)+'\n'+inspect.getsource(PairTechnicalTokenizer)+tiny
    kernels = tmp_path/'kernels'; spec = kernels/'aes_ridge_fixture'; spec.mkdir(parents=True)
    (spec/'kernel.json').write_text(json.dumps(dict(argv=[sys.executable,'-I','-m','ipykernel_launcher','-f','{connection_file}'], display_name='AES Ridge fixture', language='python')))
    workspace = isolated/'workspace'
    env = dict(os.environ, AES300_WORKSPACE=str(workspace), AES300_DATASET=str(dataset), AES300_PYTHON=sys.executable,
               HF_HUB_OFFLINE='1', TRANSFORMERS_OFFLINE='1', OMP_NUM_THREADS='2', MKL_NUM_THREADS='2', IPYTHONDIR=str(tmp_path/'ipython'))
    env.pop('PYTHONPATH', None); env.pop('AES300_AUDIT_ONLY', None)
    manager = KernelManager(kernel_name='aes_ridge_fixture', kernel_spec_manager=KernelSpecManager(kernel_dirs=[str(kernels)]))
    try:
        NotebookClient(n, timeout=240, km=manager, resources={'metadata':{'path':str(isolated)}}).execute(env=env)
    finally:
        if manager.has_kernel: manager.shutdown_kernel(now=True)
    info = json.loads((workspace/'outputs/bge_ridge_download.json').read_text())
    assert len(info['books']) == 5 and info['test_evaluated']
    run = Path(info['run_folder']); result = json.loads((run/'ridge_result.json').read_text())
    encoding = json.loads((run/'bge_encoding.json').read_text())
    assert encoding['frozen'] and encoding['answer_tokens_uncovered'] == 0
    assert len(result['predictions']) == 24 and not result['BGE_encoder_trained']
    assert all(c.execution_count is not None for c in n.cells if c.cell_type=='code')
    evidence = ROOT/'outputs/bge_ridge_three_input_verification'; evidence.mkdir(parents=True, exist_ok=True)
    nbformat.write(n, evidence/'TECHNICAL_executed.ipynb')
    (evidence/'proof.json').write_text(json.dumps(dict(scope='SYNTHETIC24 + RANDOM FROZEN XLM-R; NO PRETRAINED AES RESULTS',
        code_sha256=nbformat.read(NOTEBOOK, as_version=4).metadata.aes_standalone.code_sha256,
        code_cells=sum(c.cell_type=='code' for c in n.cells), test_evaluated_fixture=True,
        books=5, source_notebook_outputs_empty=True), indent=2))
