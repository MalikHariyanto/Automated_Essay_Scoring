"""Fixture teknis untuk latihan AI; bukan esai/label penelitian."""
import json
from pathlib import Path
import numpy as np
import pandas as pd
import pytest
import torch
from openpyxl import Workbook, load_workbook
from aes.ai_learning import SCOPE, prepare_ai_learning, collect_report
from aes.audit import EligibilityError
from aes.io import load_data, sha256
from aes.trainer import train_run
from test_trainer import TinyEncoder, TinyTokenizer

@pytest.fixture
def ai_source(technical_source):
    df,r,cfg=technical_source
    # Empat kelompok kata berbeda, agar fixture tidak mengarang independensi esai siswa.
    words=['alpha','bravo','charlie','delta']
    df['text_raw']=[(' '+words[i%4])*15+str(i) for i in range(10)]
    df['text_clean']=df.text_raw
    df.to_csv(cfg['csv'],index=False)
    w=Workbook();ws=w.active;ws.title='Data_asli'
    ws.append(list(df.columns))
    for row in df.itertuples(index=False,name=None):ws.append(list(row))
    refs=w.create_sheet('Ringkasan_AI');refs.append(['essay_id','status_usulan']+['ai_usul_'+t['name'] for t in r['traits']])
    for i,row in df.iterrows():refs.append([row.essay_id,'AI_RECONSTRUCTION_SIMULATION' if i>=8 else 'AI_TECHNICAL_FIXTURE']+[(i+j)%5+1 for j in range(4)])
    ws=w.create_sheet('JSON_asli');ws['B3']=sha256(cfg['rubric'])
    path=Path(cfg['csv']).parent/'technical_ai.xlsx';w.save(path)
    cfg['input_workbook']={'path':str(path),'sha256':sha256(path)}
    cfg['ai_learning']['enabled']=True
    return cfg


def test_ai_references_separate_from_teacher_and_reconstruction(ai_source):
    cfg=ai_source;before={k:sha256(cfg[k]) for k in ['csv','rubric']}
    df,r,y,train,valid,m=prepare_ai_learning(cfg)
    assert y.shape==(10,4) and np.isfinite(y[train+valid]).all()
    assert np.isnan(y[8:]).all() and len(train)+len(valid)==8
    assert not set(df.iloc[train].group_id)&set(df.iloc[valid].group_id)
    assert m['scope']==SCOPE and len(m['excluded'])==2
    assert all(df['score_'+t['name']].eq('').all() for t in r['traits'])
    assert not r['confirmed'] and df.ready_for_training.eq('false').all()
    assert before=={k:sha256(cfg[k]) for k in before}


def test_ai_requires_explicit_mode_and_matching_book(ai_source):
    cfg=ai_source;cfg['ai_learning']['enabled']=False
    with pytest.raises(EligibilityError,match='eksplisit'):prepare_ai_learning(cfg)
    cfg['ai_learning']['enabled']=True;cfg['input_workbook']['sha256']='wrong'
    with pytest.raises(EligibilityError,match='cocok'):prepare_ai_learning(cfg)


def test_ai_refuses_missing_candidate_and_changed_text(ai_source):
    cfg=ai_source;p=cfg['input_workbook']['path'];w=load_workbook(p)
    w['Ringkasan_AI']['C2']=None;w.save(p);cfg['input_workbook']['sha256']=sha256(p)
    with pytest.raises(EligibilityError,match='belum berupa'):prepare_ai_learning(cfg)
    w['Ringkasan_AI']['C2']=3;w.save(p);cfg['input_workbook']['sha256']=sha256(p)
    df=load_data(cfg['csv']);df.loc[0,'text_raw']='CHANGED TECHNICAL TEXT';df.to_csv(cfg['csv'],index=False)
    with pytest.raises(EligibilityError,match='Teks berubah'):prepare_ai_learning(cfg)


def test_ai_refuses_duplicate_leakage_between_groups(ai_source):
    cfg=ai_source;df=load_data(cfg['csv']);df.loc[4,'group_id']='OTHER_GROUP';df.to_csv(cfg['csv'],index=False)
    with pytest.raises(EligibilityError,match='lintas group'):prepare_ai_learning(cfg)


@pytest.mark.parametrize('scenario',['S1','S2','S3'])
def test_ai_backward_actual_update_and_checkpoint_scope(scenario,ai_source,monkeypatch):
    import aes.trainer as tr
    cfg=ai_source;df,r,y,train,valid,m=prepare_ai_learning(cfg)
    monkeypatch.setattr(tr,'load_roberta',lambda _:TinyEncoder())
    monkeypatch.setattr(tr.AutoTokenizer,'from_pretrained',lambda *a,**k:TinyTokenizer())
    bge=torch.nn.functional.normalize(torch.randn(len(df),1024),dim=1);frozen=bge.clone()
    resolved={'bge':{'checkpoint':'BAAI/bge-m3','revision':'fixture'},'roberta':{'checkpoint':'flax-community/indonesian-roberta-base','revision':'fixture'}}
    result=train_run(scenario,'A',train,None,df.text_raw.tolist(),y,bge,resolved,cfg,r,42,
                     Path(cfg['output'])/scenario,epochs=2,early_stop=False,training_scope=SCOPE)
    assert result['steps_with_nonzero_learning_rate']==1
    assert all(result['parameter_sample_changes'].values())
    assert torch.equal(frozen,bge)
    p=torch.load(result['checkpoint'],weights_only=True)
    assert p['training']['scope']==SCOPE and p['training']['validation_essay_count']==0
    assert p['training']['monitor']=='fixed_epochs_no_validation'


def test_one_optimizer_step_does_not_have_zero_lr_for_entire_run(technical_source):
    _,r,cfg=technical_source
    result=train_run('S2','A',[0,1],None,['fixture']*2,np.full((2,4),3.),
                     torch.randn(2,1024),{},cfg,r,42,Path(cfg['output'])/'one_step',
                     epochs=1,early_stop=False,training_scope=SCOPE)
    assert result['steps_with_nonzero_learning_rate']==1
    assert result['parameter_sample_changes']['head']


def test_ai_report_and_notebook_trace_use_ai_label(ai_source,capsys):
    cfg=ai_source;df,r,y,train,valid,m=prepare_ai_learning(cfg)
    predictions={s:(y[valid]+.2).tolist() for s in ['S1','S2','S3']}
    report=collect_report(m,r,y,valid,{},predictions,{})
    report['reused_completed_run']=False
    assert len(report['rows'])==12 and all(row['RMSE']==pytest.approx(.2) for row in report['rows'])
    assert report['reference_source']=='AI_PROVISIONAL_NOT_TEACHER'
    output=Path(cfg['output']);output.mkdir(parents=True,exist_ok=True)
    (output/'ai_learning_report.json').write_text(json.dumps(report))
    root=Path(__file__).resolve().parents[1];nb=json.loads((root/'notebooks/02_latihan_acuan_AI.ipynb').read_text())
    cell=next(''.join(c['source']) for c in nb['cells'] if c['cell_type']=='code' and ''.join(c['source']).startswith("run_aes([ROOT/'scripts/notebook_step.py', 'report'"))
    tables=[]
    exec(compile(cell,'ai_fixture_ui','exec'),dict(ROOT=root,OUTPUT=output,CONFIG_PATH='FIXTURE',MODE='ai_learning',TRACE_SCENARIO='S3',TRACE_TRAIT='isi',json=json,
         run_aes=lambda _:None,show_table=lambda cols,rows:tables.append((cols,rows))))
    assert any('reference_AI' in cols for cols,rows in tables)
    assert 'ACUAN AI' in capsys.readouterr().out.upper()


def test_normal_inference_refuses_ai_checkpoint(monkeypatch):
    import aes.inference as inf
    monkeypatch.setattr(inf,'load_checkpoint',lambda *a:(None,{'training':{'trained':True,'scope':SCOPE}}))
    with pytest.raises(ValueError,match='latihan acuan AI'):inf.predict('technical.pt',['fixture'])
