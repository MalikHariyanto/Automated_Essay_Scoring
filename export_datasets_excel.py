"""Convert the two explicitly requested public datasets into real Excel workbooks.

No training, target rescaling, or modification of active AES research inputs.
"""
import argparse
import csv
import hashlib
import json
import re
import shlex
import sys
from pathlib import Path
import zipfile

import xlsxwriter
from openpyxl import load_workbook

ASAP_SHA = 'cf04127eaff3224ce4aab64654e62e5738d3d672c61bb998546dd1674c60d312'
ELLIPSE_SHA = '782344e99668a3ff508d7410c0eb6e36da70f3b28f81c96e367f1ca04924b06c'
ASAP_REPO = 'https://github.com/robert1ridley/cross-prompt-trait-scoring'
ELLIPSE_REPO = 'https://github.com/scrosseye/ELLIPSE-Corpus'
TRAITS = ['content', 'organization', 'style', 'conventions']
ELL_TRAITS = ['Cohesion', 'Syntax', 'Vocabulary', 'Phraseology', 'Grammar', 'Conventions']


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read_source(path, encoding, delimiter, expected_sha):
    assert sha(path) == expected_sha, 'Source bytes differ from the verified public dataset'
    with path.open(encoding=encoding, newline='') as f:
        reader = csv.DictReader(f, delimiter=delimiter)
        rows = list(reader)
    assert all(None not in row for row in rows), 'Unexpected field count'
    return reader.fieldnames, rows


def value(text, numeric):
    if text == '':
        return None
    if numeric and re.fullmatch(r'[+-]?\d+(?:\.\d+)?', text):
        number = float(text)
        return int(number) if number.is_integer() else number
    return text


def write_table(book, name, headers, rows, *, text_col=None):
    sheet = book.add_worksheet(name)
    number = book.add_format({'num_format': '0.0#', 'valign': 'top'})
    integer = book.add_format({'num_format': '0', 'valign': 'top'})
    wrap = book.add_format({'text_wrap': True, 'valign': 'top'})
    plain = book.add_format({'valign': 'top'})
    sheet.freeze_panes(1, 1)
    sheet.hide_gridlines(2)
    sheet.set_row(0, 32)
    sheet.set_default_row(24)
    sheet.set_column(0, 0, 17, plain)
    sheet.set_column(1, len(headers)-1, 20, plain)
    if text_col is not None:
        sheet.set_column(text_col, text_col, 88, wrap)
    sheet.add_table(0, 0, len(rows), len(headers)-1, {
        'name': re.sub(r'[^A-Za-z0-9_]', '_', name)+'_Table',
        'columns': [{'header': x} for x in headers],
        'style': 'Table Style Medium 2',
    })
    for n, row in enumerate(rows, 1):
        if text_col is not None:
            sheet.set_row(n, 66)
        for col, cell in enumerate(row):
            if cell is None:
                continue
            if isinstance(cell, (int, float)):
                sheet.write_number(n, col, cell, integer if float(cell).is_integer() else number)
            else:
                assert len(cell) <= 32767, 'Excel cell limit would truncate text'
                assert not re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', cell), 'Invalid XML text'
                # Always store source text as text, including leading =, +, - and @.
                sheet.write_string(n, col, cell, wrap if col == text_col else plain)
    sheet.set_landscape()
    sheet.repeat_rows(0)
    sheet.set_tab_color('#4F6793')
    return sheet


def instructions(book, title, notes):
    sheet = book.add_worksheet('Petunjuk')
    sheet.hide_gridlines(2)
    sheet.set_column('A:A', 4)
    sheet.set_column('B:B', 112)
    heading = book.add_format({'font_size': 20, 'bold': True, 'font_color': '#223057', 'text_wrap': True})
    text = book.add_format({'text_wrap': True, 'valign': 'top', 'font_size': 11})
    sheet.write_string(1, 1, title, heading)
    sheet.set_row(1, 45)
    for row, note in enumerate(notes, 3):
        sheet.write_string(row, 1, note, text)
        sheet.set_row(row, 55)
    sheet.set_tab_color('#9A1316')


def workbook(path):
    return xlsxwriter.Workbook(path, {'strings_to_formulas': False, 'strings_to_urls': False})


def preserve_carriage_returns(path):
    # Use XML character references for CR rather than Excel's escaped string form.
    # Both are legal OOXML, but this also preserves values in openpyxl/pandas.
    # A literal source token escaped as _x005F_x000D_ must remain literal.
    temporary = path.with_suffix('.tmp.xlsx')
    with zipfile.ZipFile(path) as original, zipfile.ZipFile(temporary, 'w') as fixed:
        for info in original.infolist():
            payload = original.read(info.filename)
            if info.filename == 'xl/sharedStrings.xml':
                payload = re.sub(rb'(?<!_x005F)_x000D_', b'&#13;', payload)
            fixed.writestr(info, payload)
    temporary.replace(path)


def verify(path, expectations):
    result = {}
    with zipfile.ZipFile(path) as z:
        assert '[Content_Types].xml' in z.namelist() and 'xl/workbook.xml' in z.namelist()
        assert not any(x.endswith('vbaProject.bin') for x in z.namelist())
    book = load_workbook(path, read_only=True, data_only=False)
    for name, headers, expected in expectations:
        sheet = book[name]
        iterator = sheet.iter_rows()
        assert [c.value for c in next(iterator)] == headers, name
        count = 0
        for actual, wanted in zip(iterator, expected, strict=True):
            assert [c.value for c in actual] == wanted, (name, count+2, 'value mismatch')
            assert all(c.data_type != 'f' for c in actual), (name, count+2, 'unexpected formula')
            count += 1
        assert count == len(expected)
        result[name] = {'data_rows': count, 'columns': len(headers), 'all_cells_match': True, 'source_text_not_interpreted_as_formulas': True}
    book.close()
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--asap', type=Path, required=True)
    parser.add_argument('--ellipse', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    ahead, asap = read_source(args.asap, 'latin-1', '\t', ASAP_SHA)
    ehead, ellipse = read_source(args.ellipse, 'utf-8-sig', ',', ELLIPSE_SHA)
    assert len(asap) == 12976 and len(ellipse) == 3911
    selected = [row for row in asap if row['essay_set'] == '7']
    assert len(selected) == 1569
    aoriginal = [[value(row[h], h not in ('essay', 'essay_id')) for h in ahead] for row in asap]
    eoriginal = [[value(row[h], h not in ('full_text', 'text_id_kaggle', 'gender', 'race_ethnicity', 'SES', 'task', 'prompt', 'set')) for h in ehead] for row in ellipse]
    afocus_headers = ['essay_id', 'essay']+[t+'_0_6' for t in TRAITS]
    afocus, araters = [], []
    rater_headers = ['essay_id']+[f'rater{r}_{t}_0_3' for t in TRAITS for r in (1, 2)]
    for row in selected:
        scores, raters = [], []
        for i in range(1, 5):
            first = int(row[f'rater1_trait{i}'])
            second = int(row[f'rater2_trait{i}'])
            assert 0 <= first <= 3 and 0 <= second <= 3
            scores.append(first+second)
            raters += [first, second]
        afocus.append([row['essay_id'], row['essay']]+scores)
        araters.append([row['essay_id']]+raters)
    efocus_headers = ['text_id_kaggle', 'full_text']+ELL_TRAITS
    efocus = [[row['text_id_kaggle'], row['full_text']]+[float(row[t]) for t in ELL_TRAITS] for row in ellipse]
    assert all(1 <= score <= 5 for row in efocus for score in row[2:])

    apath = args.output/'ASAP_Train_dan_Prompt7.xlsx'
    with workbook(apath) as book:
        instructions(book, 'ASAP — data Excel dan empat dimensi Prompt 7', [
            'Tab Prompt7_4dimensi: 1.569 esai dengan empat skor terpisah. Teks dan ID tetap sama dengan sumber. Tab ini dibuka pertama kali.',
            'content, organization, style, dan conventions adalah nama dimensi sumber. style tidak otomatis sama dengan seluruh definisi kebahasaan penelitian pengguna.',
            'Skor pada tab Prompt7_4dimensi adalah jumlah dua penilai manusia: rater1 + rater2, mengikuti arrange_data.py pada sumber. Rentang 0–6 per dimensi, bukan 1–5. Ini pengolahan skor manusia, bukan penilaian AI baru.',
            'Tab Penilai_Prompt7 menyimpan masing-masing skor penilai 0–3 secara terpisah agar penggabungannya dapat diperiksa.',
            'Tab ASAP_semua_asli: seluruh 12.976 baris dan 28 kolom sumber TSV, termasuk prompt lain. Nilai kosong tetap kosong. Tidak semua prompt memiliki empat dimensi Prompt 7.',
            'Bahasa Inggris. File ini ekspor dataset publik eksternal dan tidak mengganti input atau rubrik aktif penelitian AES pengguna.',
            'Sumber: '+ASAP_REPO+' | commit 981ec606f039c5b3ee163b5b474e4329d9ebbd17. Pengolahan dimensi: arrange_data.py.',
            'Dataset asli: https://www.kaggle.com/c/asap-aes/data. Atribusi dan ketentuan sumber/Kaggle tetap berlaku; tidak ada lisensi baru yang diberikan.',
            'Sumber TSV dibaca dengan encoding latin-1 mengikuti kode penelitian sumber. Tidak ada pembulatan, normalisasi skala, pembuatan label kosong, atau penilaian ulang.'
        ])
        focus = write_table(book, 'Prompt7_4dimensi', afocus_headers, afocus, text_col=1)
        write_table(book, 'Penilai_Prompt7', rater_headers, araters)
        write_table(book, 'ASAP_semua_asli', ahead, aoriginal, text_col=2)
        focus.activate()
    epath = args.output/'ELLIPSE_Train.xlsx'
    with workbook(epath) as book:
        instructions(book, 'ELLIPSE — train Excel dan enam dimensi', [
            'Tab Skor_6dimensi: 3.911 esai, teks asli, ID, dan enam skor analitik manusia. Tab ini dibuka pertama kali.',
            'Dimensi: Cohesion, Syntax, Vocabulary, Phraseology, Grammar, Conventions. Dataset tidak memiliki skor content/isi terpisah.',
            'Skor asli 1–5, termasuk setengah poin seperti 2,5 dan 3,5. Semua nilai dipertahankan tanpa pembulatan ke bilangan bulat atau pemilihan empat dimensi otomatis.',
            'Tab Train_asli: seluruh 3.911 baris dan 26 kolom CSV sumber, termasuk metadata dan Overall yang sudah ada di sumber. Tidak ada skor total baru yang dihitung.',
            'Bahasa Inggris. Data train saja; file test dan arsip penilai mentah tidak digabungkan. Ekspor ini tidak mengubah input penelitian aktif pengguna.',
            'Sumber pembuat: '+ELLIPSE_REPO+' | commit dc3b8f0b3b4332fc9f64302c4ccfc4ed582f4b43. Sumber CSV: ELLIPSE_Final_github_train.csv.',
            'Sitasi: Crossley et al. (2023), Measuring second language proficiency using the English Language Learner Insight, Proficiency and Skills Evaluation (ELLIPSE) Corpus. International Journal of Learner Corpus Research 9(2), 248–269.',
            'Lisensi sumber dan adaptasi dataset ini: CC BY-NC-SA 4.0. https://creativecommons.org/licenses/by-nc-sa/4.0/ | Atribusi, nonkomersial, dan ShareAlike tetap berlaku.',
            'Angka fitur dan skor disimpan sebagai angka Excel. Teks esai disimpan sebagai teks, termasuk jika diawali karakter =, +, -, atau @.'
        ])
        focus = write_table(book, 'Skor_6dimensi', efocus_headers, efocus, text_col=1)
        write_table(book, 'Train_asli', ehead, eoriginal, text_col=1)
        focus.activate()

    preserve_carriage_returns(apath)
    preserve_carriage_returns(epath)
    checks = {
        apath.name: verify(apath, [('Prompt7_4dimensi', afocus_headers, afocus), ('Penilai_Prompt7', rater_headers, araters), ('ASAP_semua_asli', ahead, aoriginal)]),
        epath.name: verify(epath, [('Skor_6dimensi', efocus_headers, efocus), ('Train_asli', ehead, eoriginal)]),
    }
    report = {
        'scope': 'Public external datasets converted to Excel at user request; no research training or active-data changes.',
        'command': shlex.join([sys.executable, *sys.argv]),
        'cwd': str(Path.cwd()),
        'exit_code': 0,
        'source_digests': {'ASAP_TSV': ASAP_SHA, 'ELLIPSE_CSV': ELLIPSE_SHA},
        'workbooks': {p.name: {'bytes': p.stat().st_size, 'sha256': sha(p), 'sheets_checked': checks[p.name]} for p in (apath, epath)},
        'conversion': {'asap_prompt7': 'Two human rater scores summed per trait, per source arrange_data.py; 0–6 scale.', 'ellipse': 'Original six human analytic scores preserved, including halves.', 'missing_labels_filled': False, 'original_texts_modified': False, 'scale_normalized': False},
        'passed': ['All 12,976 ASAP source rows and all 3,911 ELLIPSE train rows preserved', 'Every exported data cell verified by independently reopening Excel with openpyxl', 'All 1,569 Prompt 7 four-trait sums and individual raters verified', 'No essay truncation, Excel formulas, macros or automatic text hyperlinks'],
        'failed': [], 'skipped': [], 'xfailed': [],
        'versions': {'xlsxwriter': xlsxwriter.__version__, 'openpyxl': __import__('openpyxl').__version__},
    }
    (args.output/'verification.json').write_text(json.dumps(report, ensure_ascii=False, indent=2)+'\n')
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
