"""Latihan berangka dari usulan AI, terpisah dari training/evaluasi guru."""
import copy
import hashlib
import json
from pathlib import Path

import numpy as np
import pandas as pd
import torch
from openpyxl import load_workbook
from sklearn.model_selection import GroupShuffleSplit

from .audit import EligibilityError, audit
from .checkpoint import load_checkpoint
from .cleaning import clean
from .encoders import BGEEncoder, cached_bge, resolve_models
from .io import dump, load_data, load_rubric, sha256
from .metric_trace import metric_trace
from .metrics import evaluate, ordinal_prediction
from .trainer import BatchEngine, device_name, train_with_oom_retry
from .workbook import sheet_records

SCOPE = 'ai_reference_learning_only'
VERSION = 1
SCENARIOS = ['S1','S2','S3']


def prepare_ai_learning(cfg):
    options=cfg.get('ai_learning',{})
    if not options.get('enabled'):
        raise EligibilityError('Pilih MODE=ai_learning secara eksplisit untuk latihan memakai acuan AI.')
    book_info=cfg.get('input_workbook',{})
    source=Path(book_info.get('path',''))
    if not source.is_file() or sha256(source)!=book_info.get('sha256'):
        raise EligibilityError('Workbook AI belum tersedia/cocok. Impor Excel lengkap lewat cell input terlebih dahulu.')
    df=load_data(cfg['csv']);rubric=load_rubric(cfg['rubric'])
    if df.essay_id.duplicated().any():raise EligibilityError('ID esai duplikat perlu diperiksa.')
    if df.rubric_version.ne(rubric['version']).any():raise EligibilityError('Versi rubrik dataset berbeda.')
    book=load_workbook(source,read_only=True,data_only=False)
    try:
        if not {'Ringkasan_AI','Data_asli','JSON_asli'}.issubset(book.sheetnames):
            raise EligibilityError('Pakai workbook review lengkap dengan Ringkasan_AI, Data_asli dan JSON_asli.')
        if book['JSON_asli']['B3'].value!=sha256(cfg['rubric']):
            raise EligibilityError('Rubrik aktif berbeda dari rubrik sumber usulan AI.')
        original={r['essay_id']:r for r in sheet_records(book,'Data_asli')[1]}
        records=sheet_records(book,'Ringkasan_AI')[1]
    finally:book.close()
    if len({r['essay_id'] for r in records})!=len(records):raise EligibilityError('ID pada ringkasan AI duplikat.')
    references={r['essay_id']:r for r in records}
    if set(references)!=set(df.essay_id):raise EligibilityError('ID workbook dan data aktif tidak cocok.')
    y=np.full((len(df),4),np.nan);included=[];excluded=[]
    for index,row in df.iterrows():
        ref=references[row.essay_id]
        if row.essay_id not in original or original[row.essay_id]['text_raw']!=row.text_raw:
            raise EligibilityError('Teks berubah sejak review AI; usulan lama tidak boleh dipasangkan.')
        if row.get('text_completeness')=='continuation_unassigned' or ref.get('status_usulan')=='AI_RECONSTRUCTION_SIMULATION':
            excluded.append({'essay_id':row.essay_id,'reason':'Teks asli terpotong/versi rekonstruksi; tidak masuk latihan.'});continue
        if not row.text_raw.strip():raise EligibilityError('Teks latihan kosong.')
        if not ref.get('status_usulan','').startswith('AI_'):raise EligibilityError('Sumber angka tidak dilabeli AI.')
        for j,t in enumerate(rubric['traits']):
            value=pd.to_numeric(ref.get('ai_usul_'+t['name'],''),errors='coerce')
            if not np.isfinite(value) or value not in t['categories']:
                raise EligibilityError(f"{row.essay_id}: usulan AI {t['name']} belum berupa kategori 1–5 lengkap.")
            y[index,j]=float(value)
        included.append(index)
    eligible=df.iloc[included]
    if len(eligible)<4 or eligible.group_id.eq('').any() or eligible.group_id.nunique()<2:
        raise EligibilityError('Latihan perlu minimal empat teks dan dua kelompok kemiripan yang tidak kosong.')
    checked=audit(eligible,rubric);lookup=eligible.set_index('essay_id').group_id
    pairs=checked['exact_duplicates']+[p['ids'] for p in checked['near_duplicates']]
    if any(lookup[a]!=lookup[b] for a,b in pairs):
        raise EligibilityError('Kandidat duplikasi lintas group perlu ditinjau sebelum latihan; tidak beralih ke split random.')
    fraction=options.get('validation_fraction',.5);seed=options.get('split_seed',2026)
    if fraction!=.5:raise ValueError('Latihan kecil ini memakai holdout 50/50; protokol penelitian tidak diubah.')
    ti,vi=next(GroupShuffleSplit(n_splits=1,test_size=fraction,random_state=seed).split(eligible,groups=eligible.group_id))
    train=[included[i] for i in ti];valid=[included[i] for i in vi]
    if len(train)<2 or len(valid)<2:raise EligibilityError('Train dan validation latihan masing-masing perlu minimal dua esai.')
    epochs=options.get('epochs',2)
    if not 2<=epochs<=5:raise ValueError('Latihan acuan AI memakai 2–5 epoch tetap, tanpa tuning pada validation.')
    manifest={'scope':SCOPE,'version':VERSION,'dataset_sha256':sha256(cfg['csv']),
        'rubric_sha256':sha256(cfg['rubric']),'workbook_sha256':sha256(source),
        'train':df.iloc[train].essay_id.tolist(),'validation':df.iloc[valid].essay_id.tolist(),
        'groups':dict(zip(eligible.essay_id,eligible.group_id)),
        'group_mapping_status':'proposed_similarity_groups_NOT_verified_student_identity',
        'excluded':excluded,'validation_fraction':fraction,'split_seed':seed,'train_seed':42,'epochs':epochs,
        'candidate':'A_fixed_before_validation','models':cfg['models'],'preprocessing':cfg['preprocessing'],
        'scenarios':cfg['scenarios'],'training_settings':cfg['training'],'device':device_name(),
        'rubric_confirmed':rubric['confirmed'],'teacher_labels_modified':False,
        'research_protocol_80_20_25_folds_executed':False}
    assert not set(eligible.iloc[ti].group_id)&set(eligible.iloc[vi].group_id)
    return df,rubric,y,train,valid,manifest


def collect_report(manifest,rubric,y,valid,results,predictions,flows):
    report={'scope':SCOPE,'reference_source':'AI_PROVISIONAL_NOT_TEACHER','manifest':manifest,
        'note':'Latihan meniru usulan AI pada transkripsi draf; bukan akurasi terhadap guru atau hasil penelitian.',
        'rows':[],'metric_traces':{},'training':results,'predictions':predictions,'trained_flows':flows}
    for s in SCENARIOS:
        p=np.asarray(predictions[s],dtype=float)
        metrics=evaluate(y[valid],p,rubric)
        report['metric_traces'][s]=metric_trace(y[valid],p,rubric,manifest['validation'])
        for t in rubric['traits']:
            m=metrics[t['name']]
            report['rows'].append({'scenario':s,'trait':t['name'],**m,'status':SCOPE})
    return report


def run_ai_learning(cfg):
    df,rubric,y,train,valid,manifest=prepare_ai_learning(cfg)
    key=hashlib.sha256(json.dumps(manifest,sort_keys=True).encode()).hexdigest()[:16]
    base=Path(cfg['output'])/'ai_learning'/key
    source_hashes={k:sha256(cfg[k]) for k in ['csv','rubric']}
    report_path=base/'report.json'
    if cfg['ai_learning'].get('reuse_completed',True) and report_path.exists():
        previous=json.loads(report_path.read_text())
        if previous.get('scope')!=SCOPE or previous.get('manifest')!=manifest:raise ValueError('Laporan latihan berbeda sumber/scope.')
        for result in previous['training'].values():
            if sha256(result['checkpoint'])!=result['checkpoint_sha256']:raise ValueError('Checkpoint latihan berubah.')
        report=collect_report(manifest,rubric,y,valid,previous['training'],previous['predictions'],previous['trained_flows'])
        report['reused_completed_run']=True
        dump(Path(cfg['output'])/'ai_learning_report.json',report)
        print('Memakai hasil latihan acuan AI yang sudah selesai dan cocok hash; tidak melatih ulang.',flush=True)
        return report
    dump(base/'manifest.json',manifest)
    print(f'ACUAN AI, bukan guru: {len(train)} esai train; {len(valid)} validation; {len(manifest["excluded"])} ditunda.',flush=True)
    resolved=resolve_models(cfg)
    texts=[clean(t,**cfg['preprocessing']) for t in df.text_raw]
    encoder=BGEEncoder(resolved['bge'],device_name())
    assert not encoder.encoder.model.training and not any(p.requires_grad for p in encoder.encoder.model.parameters())
    bge,_=cached_bge(df,texts,encoder,cfg);del encoder
    frozen_before=bge.clone();results={};predictions={};flows={}
    for s in SCENARIOS:
        print(f'{s}: memuat model awal dan menjalankan {manifest["epochs"]} epoch latihan acuan AI.',flush=True)
        local=copy.deepcopy(cfg)
        result=train_with_oom_retry(scenario=s,candidate='A',train_ids=train,valid_ids=None,
            texts=texts,y=y,bge=bge,resolved=resolved,cfg=local,rubric=rubric,seed=42,out=base/s,
            epochs=manifest['epochs'],early_stop=False,training_scope=SCOPE)
        if result['steps_with_nonzero_learning_rate']<1 or not all(result['parameter_sample_changes'].values()):
            raise RuntimeError('Latihan belum terbukti memperbarui sampel parameter trainable; hasil tidak diklaim selesai.')
        model,payload=load_checkpoint(result['checkpoint'],device_name(),cfg['rubric'],cfg['preprocessing'])
        if payload['training'].get('scope')!=SCOPE:raise ValueError('Checkpoint bukan scope latihan AI.')
        engine=BatchEngine(model,texts,bge,resolved,cfg,device_name())
        pred=engine.predict(valid,cfg['training']['micro_batch'])
        with torch.no_grad():
            output,features=engine.forward(valid[:1],return_features=True)
        flows[s]={'essay_id':manifest['validation'][0],'status':'AI_TRAINED_LEARNING_ONLY',
            'shapes':{**{k:list(v.shape) for k,v in features.items()},'output':list(output.shape)},
            'vector_first_five':{k:v[0,:5].cpu().tolist() for k,v in features.items()},
            'continuous':output[0].cpu().tolist(),
            'categories':[int(ordinal_prediction(output[0,j].cpu().item(),t)) for j,t in enumerate(rubric['traits'])]}
        result['checkpoint_sha256']=sha256(result['checkpoint'])
        results[s]=result;predictions[s]=pred.tolist()
        frame=pd.DataFrame(pred,columns=['pred_'+t['name'] for t in rubric['traits']])
        frame.insert(0,'essay_id',manifest['validation'])
        for j,t in enumerate(rubric['traits']):frame['reference_ai_'+t['name']]=y[valid,j]
        frame.to_csv(base/s/'validation_against_ai.csv',index=False)
        del model,engine,features,output
        if torch.cuda.is_available():torch.cuda.empty_cache()
        assert torch.equal(frozen_before,bge),'Vektor BGE frozen berubah.'
        print(f'{s} selesai; prediksi validation dibandingkan dengan acuan AI, bukan guru.',flush=True)
    assert source_hashes=={k:sha256(cfg[k]) for k in source_hashes}
    report=collect_report(manifest,rubric,y,valid,results,predictions,flows)
    report['reused_completed_run']=False
    dump(report_path,report);dump(Path(cfg['output'])/'ai_learning_report.json',report)
    return report
