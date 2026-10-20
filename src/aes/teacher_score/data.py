"""Input privat, audit label, kelompok kebocoran, dan split tanpa mengubah sumber."""
from dataclasses import dataclass
import hashlib
import json
import math
from pathlib import Path
import unicodedata

import numpy as np
from openpyxl import load_workbook
from sklearn.model_selection import GroupKFold, GroupShuffleSplit

from aes.cleaning import clean
from . import SCOPE

DEFAULTS = {
    'sheet': 'Sheet1', 'expected_rows': 300,
    # Rentang dinyatakan notebook unggahan, bukan diambil dari min/max test.
    'score_min': 0, 'score_max': 100,
    'test_fraction': .2, 'folds': 5, 'seed': 42, 'split_seed': 2026, 'split_candidates': 128,
    # Batas dan aturan berhenti sama untuk S1, S2, dan S3; bukan wajib 25 epoch.
    'epochs': 25, 'early_stopping_patience': 5, 'effective_batch': 8, 'micro_batch': 1,
    'encoder_lr': 1e-5, 'projection_lr': 1e-4, 'head_lr': 1e-4, 'weight_decay': .01,
    'warmup_ratio': .1, 'gradient_clip': 1., 'dropout': .1,
    'roberta_max_length': 512, 'roberta_overlap': 128,
    'bge_max_length': 1024, 'bge_batch': 4, 'device': 'auto',
    'roberta': {'checkpoint': 'flax-community/indonesian-roberta-base',
                'revision': 'cc5ecee7de8532880ad08d802292a6f42865bc83'},
    'bge': {'checkpoint': 'BAAI/bge-m3',
            'revision': '5617a9f61b028005a4858fdac845db406aefb181'},
    'vector_source': 'pretrained', 'trace_position': 0,
}
REQUIRED = ['Email address', 'response_number', 'response_text', 'soal',
            'jawaban_referensi', 'rater']


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def dump(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False)+'\n', encoding='utf-8')


def config(values):
    result = {**DEFAULTS, **values}
    if result['score_min'] >= result['score_max']:
        raise ValueError('score_min harus lebih kecil dari score_max')
    if any(not isinstance(result[k], int) for k in ('score_min', 'score_max')):
        raise ValueError('Workflow ini memakai kategori skor integer yang dinyatakan terlebih dahulu.')
    if not 0 < result['test_fraction'] < 1 or result['folds'] < 2 or result['split_candidates'] < 1:
        raise ValueError('Split/folds tidak sah')
    if result['epochs'] < 1 or not 1 <= result['micro_batch'] <= result['effective_batch']:
        raise ValueError('Epoch/batch tidak sah')
    if type(result['early_stopping_patience']) is not int or result['early_stopping_patience'] < 1:
        raise ValueError('early_stopping_patience harus bilangan bulat positif.')
    if result['vector_source'] not in ('pretrained', 'S1', 'S3'):
        raise ValueError('vector_source harus pretrained, S1, atau S3')
    if result['device'] not in ('auto', 'cpu', 'cuda'):
        raise ValueError('device harus auto, cpu, atau cuda')
    if not 0 <= result['warmup_ratio'] < 1:
        raise ValueError('warmup_ratio tidak sah')
    if any(not math.isfinite(result[k]) or result[k] <= 0
           for k in ('encoder_lr', 'projection_lr', 'head_lr', 'gradient_clip', 'bge_batch')):
        raise ValueError('Learning rate, gradient clip, dan batch BGE harus positif.')
    if not 0 <= result['dropout'] < 1 or result['weight_decay'] < 0:
        raise ValueError('Dropout/weight decay tidak sah.')
    if not 2 < result['roberta_max_length'] <= 512 or not 0 <= result['roberta_overlap'] < result['roberta_max_length']-2:
        raise ValueError('Window RoBERTa melebihi batas atau overlap tidak sah.')
    if not 2 < result['bge_max_length'] <= 8192:
        raise ValueError('Batas BGE harus 3–8192 token termasuk special token.')
    return result


@dataclass
class Bundle:
    rows: list
    y: np.ndarray
    groups: np.ndarray
    source_sha256: str
    summary: dict

    @property
    def texts(self):
        return [r['text'] for r in self.rows]


def read_dataset(cfg):
    path = Path(cfg['dataset'])
    source_hash = digest(path)
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        if cfg['sheet'] not in book.sheetnames:
            raise ValueError(f"Tab {cfg['sheet']} tidak tersedia: {book.sheetnames}")
        data = iter(book[cfg['sheet']].iter_rows(values_only=True))
        header = [str(x).strip() if x is not None else None for x in next(data)]
        if any(header.count(x) != 1 for x in REQUIRED):
            raise ValueError('Enam header wajib hilang atau duplikat: '+', '.join(REQUIRED))
        rows = []
        for excel_row, values in enumerate(data, 2):
            if all(v is None for v in values):
                continue  # Hanya baris format kosong, bukan label/jawaban yang hilang.
            raw = {k: values[i] for i, k in enumerate(header) if k is not None}
            if any(raw[k] is None or not str(raw[k]).strip() for k in REQUIRED):
                raise ValueError(f'Baris {excel_row}: input/label wajib kosong; tidak diisi otomatis.')
            score = raw['rater']
            if isinstance(score, bool) or not isinstance(score, (int, float)) or not math.isfinite(score):
                raise ValueError(f'Baris {excel_row}: rater harus angka skor guru, bukan teks/ID.')
            if int(score) != score or not cfg['score_min'] <= score <= cfg['score_max']:
                raise ValueError(f'Baris {excel_row}: skor bukan kategori integer dalam rentang yang dinyatakan.')
            if not isinstance(raw['response_text'], str):
                raise ValueError(f'Baris {excel_row}: jawaban bukan teks.')
            rows.append({'row_id': f'row_{excel_row-1:04d}', 'source_row': excel_row,
                         'text': clean(raw['response_text']), 'raw_text': raw['response_text'],
                         'email_group': unicodedata.normalize('NFC', str(raw['Email address'])).strip().casefold(),
                         'question': str(raw['soal']), 'reference': str(raw['jawaban_referensi']),
                         'response_number': str(raw['response_number']), 'score': int(score)})
    finally:
        book.close()
    if cfg.get('expected_rows') is not None and len(rows) != cfg['expected_rows']:
        raise ValueError(f"Diharapkan {cfg['expected_rows']} baris, ditemukan {len(rows)}; tidak dipotong otomatis.")
    if not rows:
        raise ValueError('Dataset kosong')
    parent = list(range(len(rows)))

    def find(i):
        while parent[i] != i:
            parent[i] = parent[parent[i]]
            i = parent[i]
        return i

    def union(a, b):
        parent[find(b)] = find(a)

    email_first, duplicate_sets = {}, {}
    for i, row in enumerate(rows):
        email = row['email_group']
        if email in email_first:
            union(email_first[email], i)
        email_first[email] = i
        # Normalisasi ini HANYA audit kesamaan, bukan teks masukan encoder.
        normalized = ' '.join(row['text'].casefold().split())
        key = hashlib.sha256(normalized.encode()).hexdigest()
        duplicate_sets.setdefault(key, []).append(i)
    for indices in duplicate_sets.values():
        for i in indices[1:]:
            union(indices[0], i)
    roots = {root: i for i, root in enumerate(sorted({find(i) for i in range(len(rows))}))}
    groups = np.asarray([roots[find(i)] for i in range(len(rows))], dtype=int)
    y = np.asarray([r['score'] for r in rows], dtype=float)
    summary = {'scope': SCOPE, 'rows': len(rows), 'target_columns': ['rater'], 'model_outputs': 1,
               'score_range_observed': [float(y.min()), float(y.max())],
               'score_range_declared': [cfg['score_min'], cfg['score_max']],
               'scale_basis': '0–100 mengikuti notebook unggahan; deskriptor rubrik guru tidak tersedia.',
               'email_groups': len(email_first), 'connected_groups': len(roots),
               'questions': len({r['question'] for r in rows}),
               'duplicate_sets': sum(len(v) > 1 for v in duplicate_sets.values()),
               'cross_email_duplicate_sets': sum(len({rows[i]['email_group'] for i in v}) > 1 for v in duplicate_sets.values()),
               'source_sha256': source_hash, 'four_trait_labels_available': False,
               'four_rubric_report': 'Arsitektur empat output memungkinkan, tetapi empat skor tidak dapat ditentukan dari satu skor guru. Perlu anotasi manusia per dimensi; skor tidak disalin/dipecah otomatis.',
               'input_features': ['response_text'],
               'grouping_note': 'Email dipakai sebagai kelompok pencegah kebocoran; tidak membuktikan identitas terverifikasi. Duplikat lintas email disatukan. Pertanyaan/referensi/email/nilai tidak masuk fitur.'}
    return Bundle(rows, y, groups, source_hash, summary)


def make_split(bundle, cfg):
    indices = np.arange(len(bundle.rows))
    splitter = GroupShuffleSplit(n_splits=cfg['split_candidates'], test_size=cfg['test_fraction'], random_state=cfg['split_seed'])
    # Pilih kedekatan JUMLAH BARIS dengan target 80:20, bukan skor/hasil model.
    # GroupShuffleSplit membagi kelompok; ukuran kelompok tidak selalu sama.
    candidates = list(splitter.split(indices, groups=bundle.groups))
    development, test = min(candidates, key=lambda pair: abs(len(pair[1])/len(indices)-cfg['test_fraction']))
    if len(set(bundle.groups[development])) < cfg['folds']:
        raise ValueError('Kelompok training tidak cukup untuk CV yang diminta.')
    folds = []
    for number, (a, b) in enumerate(GroupKFold(n_splits=cfg['folds']).split(development, groups=bundle.groups[development]), 1):
        training, validation = development[a], development[b]
        assert not set(bundle.groups[training]) & set(bundle.groups[validation])
        folds.append({'fold': number, 'train': training.tolist(), 'validation': validation.tolist()})
    assert not set(bundle.groups[development]) & set(bundle.groups[test])
    manifest = {'scope': SCOPE, 'source_sha256': bundle.source_sha256,
                'development': development.tolist(), 'test': test.tolist(), 'folds': folds,
                'development_rows': len(development), 'test_rows': len(test),
                'development_groups': len(set(bundle.groups[development])),
                'test_groups': len(set(bundle.groups[test])),
                'target_test_row_fraction': cfg['test_fraction'],
                'candidate_group_fraction': cfg['test_fraction'],
                'split_candidates': cfg['split_candidates'],
                'selection_uses_labels': False,
                'actual_test_row_fraction': len(test)/len(indices),
                'policy': 'Holdout kelompok: pilih kandidat deterministik terdekat target proporsi baris berdasarkan ukuran kelompok saja, tanpa label/hasil model. CV hanya di development. Email/duplikat tidak menyeberang. Mengukur kelompok baru pada pertanyaan ini, bukan unseen-prompt.',
                'test_used_for_selection': False}
    manifest['split_sha256'] = hashlib.sha256(json.dumps(manifest, sort_keys=True).encode()).hexdigest()
    return manifest


def run_folder(cfg, bundle):
    package = Path(__file__).parent
    code_files = list(package.glob('*.py'))+[package.parent/'metrics.py', package.parent/'windows.py', package.parent/'cleaning.py']
    code_hash = hashlib.sha256(b''.join(p.name.encode()+p.read_bytes() for p in sorted(code_files))).hexdigest()
    semantic = {k: v for k, v in cfg.items() if k not in ('dataset', 'output', 'trace_position', 'vector_source', 'new_answer', 'focus_word')}
    key = hashlib.sha256(json.dumps({'source': bundle.source_sha256, 'config': semantic,
                                     'code': code_hash}, sort_keys=True).encode()).hexdigest()[:16]
    folder = Path(cfg['output'])/key
    folder.mkdir(parents=True, exist_ok=True)
    dump(folder/'effective_config.json', cfg)
    return folder
