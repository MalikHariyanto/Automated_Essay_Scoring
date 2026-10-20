"""Workbook privat lengkap: angka per koordinat, provenance, dan indeks tahapan."""
import json
import math
from pathlib import Path
import re

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill


def excel_value(value):
    if isinstance(value, np.generic):
        value = value.item()
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, default=str)
    if isinstance(value, float) and not math.isfinite(value):
        return None  # Tidak tercakup/undefined: tetap kosong, bukan nol.
    if value is pd.NA:
        return None
    return value


def write_excel(path, sheets, metadata):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook(write_only=True)
    used = set()
    counts = {}
    all_sheets = {'Info': pd.DataFrame([{'field': k, 'value': v} for k, v in metadata.items()]), **sheets}
    for name, frame in all_sheets.items():
        frame = frame if isinstance(frame, pd.DataFrame) else pd.DataFrame(frame)
        if len(frame.columns) > 16384:
            raise ValueError('Terlalu banyak kolom untuk Excel; tidak dipotong diam-diam.')
        headers = [str(x) for x in frame.columns]
        chunks = max(1, math.ceil(len(frame)/1048575))
        for part in range(chunks):
            base = re.sub(r'[\\/*?:\[\]]', '_', name)
            suffix = f'_{part+1}' if chunks > 1 else ''
            title = base[:31-len(suffix)]+suffix
            if title in used:
                raise ValueError('Nama sheet bertabrakan setelah dipendekkan: '+title)
            used.add(title)
            sheet = book.create_sheet(title)
            sheet.freeze_panes = 'B2'
            heading = []
            for text in headers:
                cell = WriteOnlyCell(sheet, value=text)
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='172035')
                cell.data_type = 's'
                heading.append(cell)
            sheet.append(heading)
            subset = frame.iloc[part*1048575:(part+1)*1048575]
            for row in subset.itertuples(index=False, name=None):
                values = []
                for value in row:
                    value = excel_value(value)
                    if isinstance(value, str):
                        if len(value) > 32767:
                            raise ValueError('Teks melebihi batas satu sel Excel; ekspor dihentikan agar tidak terpotong.')
                        cell = WriteOnlyCell(sheet, value=value)
                        cell.data_type = 's'  # Teks siswa '=...' tidak dijalankan sebagai formula.
                        values.append(cell)
                    else:
                        values.append(value)
                sheet.append(values)
            counts[title] = {'rows': len(subset), 'columns': len(headers)}
    temporary = path.with_name(path.stem+'.pending.xlsx')
    book.save(temporary)
    temporary.replace(path)
    return counts


def vector_frame(values, ids, prefix='v'):
    values = np.asarray(values)
    if values.ndim != 2 or len(values) != len(ids):
        raise ValueError('Vektor perlu [N,D] dan ID berjumlah N.')
    frame = pd.DataFrame(values, columns=[f'{prefix}_{i:04d}' for i in range(values.shape[1])])
    frame.insert(0, 'row_id', list(ids))
    return frame


class ExcelJournal:
    def __init__(self, folder, metadata, verbose=False):
        self.folder = Path(folder)
        self.folder.mkdir(parents=True, exist_ok=True)
        self.metadata = metadata
        self.verbose = verbose
        file = self.folder/'index.json'
        self.records = json.loads(file.read_text()) if file.exists() else {}

    def save(self, name, sheets, stage, **extra):
        meta = {**self.metadata, 'stage': stage, 'status': 'ACTUAL EXECUTION', **extra}
        path = self.folder/(name+'.xlsx')
        counts = write_excel(path, sheets, meta)
        self.records[name] = {'file': path.name, 'stage': stage, 'status': meta['status'],
                              'sheets_and_shapes': counts, 'source': extra.get('source', '')}
        (self.folder/'index.json').write_text(json.dumps(self.records, ensure_ascii=False, indent=2)+'\n')
        write_excel(self.folder/'00_indeks_output.xlsx', {'Daftar_file': pd.DataFrame(self.records.values())},
                    {**self.metadata, 'stage': 'Daftar keluaran aktual; hanya tahap yang selesai dicatat'})
        if self.verbose:
            print(f'Rincian tersimpan: {path.name}', flush=True)
        return path
