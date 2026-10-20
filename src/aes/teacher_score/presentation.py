"""Presentasi hasil yang sudah dihitung; tanpa torch, training atau hitung ulang metrik."""
from dataclasses import dataclass
import itertools
import json
import math
from pathlib import Path
import re
import shutil
import tempfile
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


LABELS = {'row_id':'ID Jawaban','word':'Kata','covered':'Tercakup','word_index':'Urutan Kata',
          'token_id':'ID Token','token':'Token','position':'Posisi Token','window':'Bagian Teks',
          'attention_mask':'Dibaca Model','skor_guru':'Nilai Guru','prediksi_mentah':'Prediksi Skor',
          'kategori_salinan':'Skor Dibulatkan','di_luar_rentang':'Di Luar Rentang',
          'Skenario':'Skenario','Fold':'Fold','AC2':'Gwet AC2','RMSE':'RMSE (poin)',
          'out_of_range_fraction':'Di Luar Rentang (%)','Status':'Keterangan',
          'teks_model':'Jawaban Siswa','teks_asli':'Teks Asli','score':'Nilai Guru',
          'role':'Penggunaan Data','connected_group':'Kelompok Jawaban',
          'response_text':'Jawaban Siswa','rater':'Nilai Guru','soal':'Soal',
          'sentence_index':'Urutan Kalimat','text':'Kalimat','fold':'Fold','scenario':'Skenario',
          'encoder':'Encoder','vector':'Panjang Vektor','head':'Lapisan Prediksi','output':'Jumlah Skor',
          'used_for':'Digunakan untuk','id':'Kode','title':'Judul'}
MODELS = {'S1':'RoBERTa','S2':'BGE-M3','S3':'Fusion RoBERTa + BGE-M3'}
ROLES = {'development':'Data latih (termasuk validasi silang)','test':'Data uji',
         'train':'Latih','validation':'Validasi'}


def friendly_header(name):
    match = re.fullmatch(r'v_(\d+)', str(name))
    return 'Komponen '+str(int(match[1])+1) if match else LABELS.get(name, name)


def friendly_value(name, value):
    if hasattr(value,'item'):
        try:value=value.item()
        except ValueError:pass
    if name == 'row_id' and isinstance(value,str):
        match = re.fullmatch(r'row_(\d+)',value)
        return f'J{int(match[1]):03d}' if match else value
    if name == 'role':
        return ROLES.get(value,value)
    if name == 'attention_mask' and value in (0,1):
        return 'Ya' if value else 'Tidak'
    if name in ('word_index','sentence_index','position','window') and isinstance(value,(int,float)):
        return value+1
    if isinstance(value,bool):
        return 'Ya' if value else 'Tidak'
    return value


@dataclass
class Table:
    headers: list
    rows: object


class SavedOutputs:
    def __init__(self,folder):
        folder=Path(folder)
        self.folder=folder/'excel' if (folder/'excel').is_dir() else folder
        self.files={p.name:p for p in self.folder.glob('*.xlsx') if not p.name.endswith('.pending.xlsx')}
        self.names={}

    def has(self,file,sheet=None):
        if file not in self.files:return False
        if sheet is None:return True
        if file not in self.names:
            book=load_workbook(self.files[file],read_only=True,data_only=True)
            try:self.names[file]=book.sheetnames
            finally:book.close()
        return sheet in self.names[file]

    def records(self,file,sheet):
        if not self.has(file,sheet):return
        book=load_workbook(self.files[file],read_only=True,data_only=True)
        try:
            iterator=book[sheet].iter_rows(values_only=True)
            headers=next(iterator,())
            for values in iterator:
                yield dict(zip(headers,values))
        finally:book.close()

    def headers(self,file,sheet):
        book=load_workbook(self.files[file],read_only=True,data_only=True)
        try:return list(next(book[sheet].iter_rows(values_only=True),()))
        finally:book.close()

    def copy(self,file,sheet,columns=None,filter_row=None):
        headers=columns or self.headers(file,sheet)
        def values():
            for row in self.records(file,sheet):
                if filter_row and not filter_row(row):continue
                yield [friendly_value(key,row.get(key)) for key in headers]
        return Table([friendly_header(x) for x in headers],values())


def guide(description,items):
    return Table(['Bagian','Cara Membaca'],iter([
        ['Tentang file ini',description],
        ['Asal angka','Disalin dari hasil Colab yang sudah tersimpan. Tidak ada training, prediksi atau metrik yang dihitung ulang.'],
        ['ID Jawaban','J001 berarti baris jawaban pertama. ID ini hanya penghubung antarfile, bukan fitur model.'],
        ['Angka kosong','Berarti hasil tidak tersedia atau tidak dapat dihitung; bukan angka nol.'],
        *items]))


def write_report(path,sheets):
    book=Workbook(write_only=True)
    for title,table in sheets.items():
        if len(title)>31:raise ValueError('Nama sheet terlalu panjang: '+title)
        if len(table.headers)>16384:raise ValueError('Jumlah kolom melebihi batas Excel.')
        sheet=book.create_sheet(title);sheet.freeze_panes='B2'
        sheet.sheet_format.defaultColWidth=16
        sheet.column_dimensions['A'].width=23
        if title=='Panduan':sheet.column_dimensions['B'].width=110
        for i,heading in enumerate(table.headers,1):
            if heading in ('Jawaban Siswa','Teks Asli','Kalimat','Cara Membaca','Penjelasan','Apa Artinya','Catatan'):
                sheet.column_dimensions[get_column_letter(i)].width=70
        header=[]
        for value in table.headers:
            cell=WriteOnlyCell(sheet,value=value);cell.data_type='s'
            cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='172035')
            cell.alignment=Alignment(wrap_text=True,vertical='top');header.append(cell)
        sheet.append(header)
        count=1
        for row in table.rows:
            count+=1
            if count>1048576:raise ValueError('Terlalu banyak baris; data tidak dipotong diam-diam.')
            cells=[]
            for value in row:
                if isinstance(value,(dict,list,tuple)):value=json.dumps(value,ensure_ascii=False)
                if isinstance(value,float) and not math.isfinite(value):value=None
                cell=WriteOnlyCell(sheet,value=value)
                if isinstance(value,str):
                    if len(value)>32767:raise ValueError('Teks melebihi batas satu sel Excel.')
                    cell.data_type='s'
                    cell.alignment=Alignment(wrap_text=True,vertical='top')
                elif isinstance(value,float):cell.number_format='0.0000'
                cells.append(cell)
            sheet.append(cells)
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
    temporary=path.with_name(path.stem+'.pending.xlsx');book.save(temporary);temporary.replace(path)
    return path


def data_report(source):
    sheets={'Panduan':guide('Data, pembagian data dan contoh token yang benar-benar masuk model.',[
        ['Tokenisasi','Token bukan selalu satu kata. ID Token adalah nomor token pada kosakata model.'],
        ['Bagian Teks','Jawaban panjang dapat dibaca dalam beberapa bagian. Bagian bernomor mulai 1.'],
        ['Token yang ditampilkan','Token dengan attention mask 1 ditampilkan. Padding diabaikan; special token tetap ditampilkan.']])}
    file='02_dataset_dan_cleaning.xlsx'
    if source.has(file,'Sebelum_sesudah'):
        sheets['Data dan Nilai Guru']=source.copy(file,'Sebelum_sesudah',['row_id','teks_model','skor_guru'])
    file='03_split_dan_cv.xlsx'
    if source.has(file,'Split_setiap_baris'):
        sheets['Pembagian Data']=source.copy(file,'Split_setiap_baris',['row_id','role'])
    if source.has(file,'Fold_setiap_baris'):
        sheets['Pembagian Lima Fold']=source.copy(file,'Fold_setiap_baris',['fold','row_id','role'])
        sheets['Pembagian Lima Fold'].headers[0]='Fold'
    for raw,title in [('RoBERTa_input','Token RoBERTa'),('BGE_input','Token BGE-M3')]:
        file='04_tokenisasi.xlsx'
        if source.has(file,raw):
            keys=['row_id']+(['window'] if raw=='RoBERTa_input' else [])+['position','token','token_id']
            sheets[title]=source.copy(file,raw,keys,lambda r:r.get('attention_mask',1)==1)
    return sheets if len(sheets)>1 else None


def answer_vectors_report(source):
    mappings=[('05_roberta_pretrained.xlsx','Vektor_jawaban_768','RoBERTa Sebelum Latihan'),
              ('06_bge_frozen.xlsx','Vektor_jawaban_1024','BGE-M3'),
              ('14_S1_fitur_setelah_training.xlsx','roberta','RoBERTa S1 Setelah Latihan'),
              ('14_S3_fitur_setelah_training.xlsx','roberta','RoBERTa S3 Setelah Latihan'),
              ('14_S3_fitur_setelah_training.xlsx','concat','Fusion S3')]
    sheets={'Panduan':guide('Satu baris adalah vektor satu jawaban. Semua komponen angka disimpan.',[
        ['Komponen 1, 2, ...','Satu komponen adalah satu angka dalam vektor; tidak mempunyai arti tetap seperti isi atau tata bahasa.'],
        ['Nama kolom lama','Komponen 1 adalah v_0000 pada file lama. Hanya namanya berubah, angkanya tetap.'],
        ['RoBERTa','768 angka per jawaban. Bandingkan ID Jawaban yang sama sebelum dan sesudah latihan.'],
        ['BGE-M3','1024 angka per jawaban. Bobot encoder tetap, sehingga tidak perlu disalin berulang untuk S2 dan S3.'],
        ['Fusion S3','Gabungan 768 angka RoBERTa dan 256 angka hasil proyeksi BGE. Hasilnya 1024 komponen.'],
        ['Batas perbandingan','Jangan menghitung cosine langsung antarruang RoBERTa dan BGE. Vektor bukan skor siswa.']])}
    for file,raw,title in mappings:
        if source.has(file,raw):sheets[title]=source.copy(file,raw)
    return sheets if len(sheets)>1 else None


def repeated_word_rows(records,model):
    for row in records:
        positions=row.get('occurrence_indices')
        if isinstance(positions,str):positions=json.loads(positions)
        positions=positions or []
        same_ids=row.get('same_token_ids')
        initial=row.get('initial_max_abs_difference')
        same_initial=None if initial is None else abs(initial)<=1e-6
        same_context=row.get('contextual_allclose_1e-6')
        if not same_ids:explanation='Ejaan sama, tetapi ID subword berbeda. Bandingkan vektor dengan memperhatikan tokenisasinya.'
        elif same_context is None:explanation='Hasil perbandingan vektor akhir belum tersedia.'
        elif same_context:explanation='ID token sama; vektor akhir juga sama dalam toleransi numerik 0,000001.'
        else:explanation='ID token sama, tetapi vektor akhir berbeda karena konteks dan posisi.'
        yield [model,row.get('word'),positions[0]+1 if positions else None,positions[1]+1 if len(positions)>1 else None,
               friendly_value('',same_ids),friendly_value('',same_initial),friendly_value('',same_context),
               row.get('contextual_cosine'),explanation]


def word_vectors_report(source):
    base='07_vektor_kata_kalimat_pretrained.xlsx'
    conditions=[(base,'roberta','RoBERTa Sebelum'),(base,'bge','BGE-M3'),
                ('15_S1_vektor_kata_setelah_training.xlsx','roberta','RoBERTa S1'),
                ('15_S3_vektor_kata_setelah_training.xlsx','roberta','RoBERTa S3')]
    sheets={'Panduan':guide('Vektor kata dan kalimat untuk satu jawaban contoh yang sama.',[
        ['Kata dan kalimat','Sheet Kata berisi vektor setelah konteks. Sheet Kalimat berasal dari pembacaan kalimat secara terpisah.'],
        ['Sebelum Konteks','Lookup embedding kata sebelum penambahan konteks; hanya perlu disajikan sekali untuk tiap encoder pretrained.'],
        ['Kata Berulang','Bandingkan dua posisi kata di model yang sama. Kesamaan numerik memakai toleransi 0,000001.'],
        ['Tercakup',"Nilai 'Tidak' berarti kata melewati batas input; komponen kosong bukan vektor nol."],
        ['BGE-M3','Ditampilkan sekali karena encodernya frozen. Rincian internal tetap tersimpan pada hasil sumber.']])}
    if source.has('13_prediksi_satu_jawaban.xlsx','Input_privat'):
        sheets['Jawaban Contoh']=source.copy('13_prediksi_satu_jawaban.xlsx','Input_privat')
        sheets['Jawaban Contoh'].headers=['Jawaban Siswa']
    comparisons=[]
    def vector_table(file,sheet,unit):
        copied=source.copy(file,sheet)
        copied.headers[0]='Urutan Kata' if unit=='word' else 'Urutan Kalimat'
        return Table(copied.headers,([r[0]+1,*r[1:]] for r in copied.rows))
    for file,role,label in conditions:
        for key,title in [('word_contextual','Kata '+label),('sentence_cls','Kalimat '+label)]:
            if source.has(file,role+'_'+key):sheets[title]=vector_table(file,role+'_'+key,'word' if key.startswith('word') else 'sentence')
        if source.has(file,role+'_kata_berulang'):
            comparisons.extend(repeated_word_rows(source.records(file,role+'_kata_berulang'),label))
    for role,label in [('roberta','RoBERTa'),('bge','BGE')]:
        if source.has(base,role+'_word_initial'):
            sheets[label+' Sebelum Konteks']=vector_table(base,role+'_word_initial','word')
    if comparisons:
        sheets['Kata Berulang']=Table(['Model','Kata','Posisi Kata 1','Posisi Kata 2','ID Token Sama',
            'Vektor Awal Sama','Vektor Akhir Sama','Kemiripan Vektor Akhir','Apa Artinya'],comparisons)
    return sheets if any(name.startswith('Kata ') for name in sheets) else None


def evaluation_report(source):
    file='12_evaluasi_tiga_skenario.xlsx'
    sheets={'Panduan':guide('Buka Perbandingan Model untuk hasil akhir; buka Nilai Guru dan Prediksi untuk melihat setiap jawaban uji.',[
        ['S1, S2, S3','S1 = RoBERTa; S2 = BGE-M3; S3 = fusion keduanya.'],
        ['Data uji','Angka test untuk laporan akhir. Pemilihan model/epoch memakai validasi silang, bukan tuning dari test.'],
        ['Skor mentah','Disimpan persis dari hasil sumber. Tampilan empat desimal tidak mengubah angka yang tersimpan.'],
        ['Ringkasan Validasi','Rata-rata ± SD antar-fold; bukan interval kepercayaan. Nilai mentah tiap-fold ada di Hasil Tiap Fold.']])}
    if source.has(file,'TEST_final'):
        def final_rows():
            for r in source.records(file,'TEST_final'):
                reasons=r.get('undefined_reasons') or ''
                if isinstance(reasons,str) and reasons.startswith('{'):
                    reasons=json.loads(reasons)
                if isinstance(reasons,dict):reasons='; '.join(f'{k}: {v}' for k,v in reasons.items())
                outside=r.get('out_of_range_fraction')
                yield [r.get('Skenario'),MODELS.get(r.get('Skenario')),r.get('RMSE'),r.get('Pearson'),r.get('QWK'),r.get('AC2'),
                       None if outside is None else outside*100,reasons]
        sheets['Perbandingan Model']=Table(['Skenario','Model','RMSE (poin)','Pearson','QWK','Gwet AC2','Di Luar Rentang (%)','Catatan'],final_rows())
    if source.has(file,'CV_semua_fold'):
        sheets['Hasil Tiap Fold']=source.copy(file,'CV_semua_fold',['Skenario','Fold','RMSE','Pearson','QWK','AC2'])
    if source.has(file,'CV_mean_SD_n'):
        def summaries():
            for row in source.records(file,'CV_mean_SD_n'):
                values=[]
                for metric in ('RMSE','Pearson','QWK','AC2'):
                    mean,sd,n=row.get(metric+'_mean'),row.get(metric+'_std'),row.get(metric+'_count')
                    values.append('Tidak dapat dihitung' if mean is None else f'{mean:.4f} ± {sd:.4f} (n={n})' if sd is not None else f'{mean:.4f} (n={n})')
                yield [row.get('Skenario'),*values]
        sheets['Ringkasan Validasi']=Table(['Skenario','RMSE (rata-rata ± SD)','Pearson (rata-rata ± SD)','QWK (rata-rata ± SD)','Gwet AC2 (rata-rata ± SD)'],summaries())
    predicted={}
    for scenario in ('S1','S2','S3'):
        f=f'11_{scenario}_test.xlsx'
        for row in source.records(f,'Prediksi_test'):
            item=predicted.setdefault(row['row_id'],{'teacher':row['skor_guru']})
            if item['teacher']!=row['skor_guru']:raise ValueError('Nilai guru antarfile test berbeda.')
            item[scenario]=row['prediksi_mentah']
            item['selisih_'+scenario]=row.get('selisih_mentah',row.get('selisih'))
    if predicted:
        def per_answer():
            for identifier,row in predicted.items():
                yield [friendly_value('row_id',identifier),row['teacher'],*[row.get(s) for s in ('S1','S2','S3')],
                       *[row.get('selisih_'+s) for s in ('S1','S2','S3')]]
        sheets['Nilai Guru dan Prediksi']=Table(['ID Jawaban','Nilai Guru','Prediksi S1','Prediksi S2','Prediksi S3','Selisih S1','Selisih S2','Selisih S3'],per_answer())
    if source.has('13_prediksi_satu_jawaban.xlsx','Tiga_skenario'):
        sheets['Prediksi Satu Jawaban']=source.copy('13_prediksi_satu_jawaban.xlsx','Tiga_skenario',
                           ['Skenario','prediksi_mentah','kategori_salinan','di_luar_rentang'])
        sheets['Jawaban Contoh']=source.copy('13_prediksi_satu_jawaban.xlsx','Input_privat')
        sheets['Jawaban Contoh'].headers=['Jawaban Siswa']
    if len(sheets)==1:return None
    sheets['Arti Metrik']=Table(['Metrik','Cara Membaca','Catatan'],[
        ['RMSE','Lebih kecil berarti kesalahan prediksi lebih kecil.','Satuannya poin skor, bukan dibatasi 0–1.'],
        ['Pearson','Hubungan linear antara skor guru dan prediksi.','Tinggi belum menjamin nilai prediksi tepat.'],
        ['QWK','Kesepakatan skor kategori dengan penalti kuadratik.','Dibaca bersama RMSE dan Pearson.'],
        ['Gwet AC2','Kesepakatan kategori dengan koreksi peluang Gwet.','Dapat berbeda jauh dari QWK; daftar kategori memengaruhi hasil.']])
    return sheets


def training_report(source):
    sheets={'Panduan':guide('Riwayat latihan per skenario. Lima fold digabung ke satu sheet untuk setiap model.',[
        ['Validasi silang','Model baru dilatih pada tiap fold, dengan satu fold untuk validasi.'],
        ['Latihan final','Latihan ulang memakai seluruh development dengan jumlah epoch yang dipilih dari CV.'],
        ['Loss Training (MSE)','Target dinormalisasi saat loss dihitung. Angka ini berbeda satuan dari RMSE dalam poin skor.'],
        ['Epoch','Satu putaran seluruh data latih pada fold/tahap itu.']])}
    for scenario,label in [('S1','RoBERTa (S1)'),('S2','BGE-M3 (S2)'),('S3','Fusion (S3)')]:
        files=sorted(f for f in source.files if re.fullmatch(f'09_{scenario}_fold_\\d+\\.xlsx',f))
        final=f'10_{scenario}_refit.xlsx'
        if source.has(final):files.append(final)
        if not files:continue
        def history(files=files):
            for f in files:
                match=re.search(r'_fold_(\d+)',f)
                for row in source.records(f,'Riwayat_epoch'):
                    yield ['Validasi silang' if match else 'Latihan final',int(match[1]) if match else None,
                           row.get('epoch'),row.get('train_mse_scaled'),row.get('RMSE'),row.get('Pearson'),row.get('QWK'),row.get('AC2')]
        sheets[label]=Table(['Tahap','Fold','Epoch','Loss Training (MSE)','RMSE Validasi (poin)','Pearson Validasi','QWK Validasi','Gwet AC2 Validasi'],history())
    if len(sheets)==1:return None
    if source.has('08_arsitektur.xlsx','Spesifikasi'):
        sheets['Model yang Dibandingkan']=source.copy('08_arsitektur.xlsx','Spesifikasi')
    settings={'epochs':'Maksimum epoch per fold','folds':'Jumlah fold','effective_batch':'Batch efektif',
              'micro_batch':'Micro batch','encoder_lr':'Learning rate encoder','head_lr':'Learning rate head',
              'score_min':'Skor minimum','score_max':'Skor maksimum','roberta_max_length':'Batas token RoBERTa',
              'bge_max_length':'Batas token BGE-M3'}
    if source.has('01_konfigurasi.xlsx','Parameter'):
        sheets['Pengaturan Eksperimen']=Table(['Pengaturan','Nilai'],([settings[r['name']],r.get('value')] for r in source.records('01_konfigurasi.xlsx','Parameter') if r.get('name') in settings))
    if source.has('01_rujukan_paper.xlsx','Rujukan'):
        keys=['title','authors','year','used_for','url','limitation']
        sheets['Rujukan Penelitian']=source.copy('01_rujukan_paper.xlsx','Rujukan',keys)
        sheets['Rujukan Penelitian'].headers=['Judul','Penulis','Tahun','Digunakan untuk','Tautan','Batas Penerapan']
    return sheets


def build_reports(folder,destination):
    """Hanya membaca Excel lama; model/checkpoint/dataset asli tidak diperlukan."""
    source=SavedOutputs(folder);destination=Path(destination);destination.mkdir(parents=True,exist_ok=True)
    if not source.files:raise ValueError('Tidak ada Excel hasil lama di folder sumber.')
    report_types=[('01_Data_dan_Tokenisasi.xlsx',data_report),('02_Vektor_Jawaban.xlsx',answer_vectors_report),
                  ('03_Vektor_Kata_dan_Kalimat.xlsx',word_vectors_report),
                  ('04_Prediksi_dan_Evaluasi.xlsx',evaluation_report),('05_Riwayat_Pelatihan.xlsx',training_report)]
    files=[]
    scope=''
    for info in source.records('01_konfigurasi.xlsx','Info'):
        if info.get('field')=='evidence_scope':scope=str(info.get('value',''))
    for name,builder in report_types:
        sheets=builder(source)
        if sheets:
            if scope:
                sheets['Panduan'].rows=itertools.chain([['Cakupan Sumber',scope]],sheets['Panduan'].rows)
            path=write_report(destination/name,sheets);files.append(path)
            print('Siap: '+name,flush=True)
    if not files:raise ValueError('Workbook sumber belum memuat tahap yang dikenali.')
    archive=destination/'Hasil_AES_Ringkas.zip'
    with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED) as bundle:
        for path in files:bundle.write(path,arcname=path.name)
    result={'archive':str(archive.resolve()),'files':[str(p.resolve()) for p in files],
            'source_files':len(source.files),'training_run':False,'metrics_recomputed':False}
    (destination/'ringkasan.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    return result


def repack_archive(archive,destination):
    with tempfile.TemporaryDirectory(prefix='aes_reports_') as folder:
        seen=set()
        with zipfile.ZipFile(archive) as bundle:
            for entry in bundle.infolist():
                name=Path(entry.filename).name
                if entry.is_dir() or not re.fullmatch(r'\d\d_[^/\\]+\.xlsx',name):continue
                if name in seen:raise ValueError('Nama workbook ganda di ZIP: '+name)
                seen.add(name)
                with bundle.open(entry) as src,(Path(folder)/name).open('wb') as dst:
                    shutil.copyfileobj(src,dst)
        return build_reports(folder,destination)
