"""Impor worksheet guru ke CSV turunan; usulan AI hanya untuk inspeksi."""
import csv
import json
from pathlib import Path

from openpyxl import load_workbook

from .io import dump, load_data, sha256


TEACHER_SHEET = 'Untuk_CSV_Guru'


def sheet_records(workbook, name):
    worksheet = workbook[name]
    values = list(worksheet.iter_rows())
    if not values:
        raise ValueError(f'Tab {name} kosong.')
    for row in values:
        if any(cell.data_type == 'f' for cell in row):
            raise ValueError(f'Tab {name} memuat formula; gunakan nilai eksplisit hasil review.')
    headers = [cell.value for cell in values[0]]
    if any(not isinstance(h, str) or not h for h in headers) or len(set(headers)) != len(headers):
        raise ValueError(f'Header tab {name} kosong atau duplikat.')

    def text(value):
        if value is None:
            return ''
        if isinstance(value, bool):
            return str(value).lower()
        return str(value)

    records = [dict(zip(headers, [text(cell.value) for cell in row])) for row in values[1:]
               if any(cell.value is not None for cell in row)]
    return headers, records


def import_workbook(source, destination, sheet=TEACHER_SHEET):
    """Tidak memilih tab pertama, mengubah status, atau memetakan ai_usul_* ke score_*."""
    source, destination = Path(source).resolve(), Path(destination).resolve()
    if sheet != TEACHER_SHEET:
        raise ValueError(f'Pakai tab {TEACHER_SHEET}; Ringkasan_AI bukan label guru.')
    if source == destination or destination.suffix.lower() != '.csv':
        raise ValueError('Tujuan harus CSV turunan yang berbeda dari workbook sumber.')
    metadata_path = destination.with_suffix('.import.json')
    if destination.exists():
        previous = json.loads(metadata_path.read_text()) if metadata_path.exists() else {}
        if previous.get('workbook') != str(source) or previous.get('derived_csv_sha256') != sha256(destination):
            raise ValueError('Tujuan sudah berisi file pengguna; pilih path CSV turunan lain.')
    original_hash = sha256(source)
    workbook = load_workbook(source, read_only=True, data_only=False)
    try:
        if sheet not in workbook.sheetnames:
            raise ValueError(f'Tab wajib {sheet} tidak tersedia. Tab ditemukan: {workbook.sheetnames}')
        headers, records = sheet_records(workbook, sheet)
        needed = {'essay_id', 'text_raw', 'text_clean', 'group_id', 'rubric_version',
                  'scoring_status', 'transcription_status', 'ready_for_training',
                  'score_isi', 'score_organisasi', 'score_kebahasaan', 'score_mekanik'}
        if needed - set(headers):
            raise ValueError(f'Tab guru kurang kolom: {sorted(needed - set(headers))}')
        ai_summary = sheet_records(workbook, 'Ringkasan_AI')[1] if 'Ringkasan_AI' in workbook.sheetnames else []
        if ai_summary and not all('essay_id' in row for row in ai_summary):
            raise ValueError('Ringkasan_AI tidak memiliki essay_id.')
        tabs = workbook.sheetnames[:]
    finally:
        workbook.close()
    destination.parent.mkdir(parents=True, exist_ok=True)
    with destination.open('w', encoding='utf-8-sig', newline='') as handle:
        writer = csv.DictWriter(handle, fieldnames=headers)
        writer.writeheader()
        writer.writerows(records)
    df = load_data(destination)
    assert sha256(source) == original_hash, 'Workbook sumber berubah saat impor.'
    report = {
        'workbook': str(source), 'workbook_sha256': original_hash, 'sheet': sheet,
        'available_sheets': tabs, 'derived_csv': str(destination),
        'derived_csv_sha256': sha256(destination), 'rows': len(df),
        'teacher_score_source': f'{sheet}.score_*', 'ai_used_as_teacher_labels': False,
        'source_unchanged': True, 'ai_summary_status': 'AI_NOT_TEACHER_FOR_REVIEW_ONLY',
        'ai_summary': ai_summary,
    }
    dump(metadata_path, report)
    return report
