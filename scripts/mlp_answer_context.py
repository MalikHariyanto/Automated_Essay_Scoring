"""Sambungkan konteks privat ke ID E6 setelah semua teks/label cocok."""
import hashlib
import json
import unicodedata
from zipfile import ZipFile


def build_context(dataset, results):
    from openpyxl import load_workbook
    from build_mlp_study import table
    book=load_workbook(dataset,read_only=True,data_only=True)
    try:
        it=book['Sheet1'].iter_rows(values_only=True);headers=next(it)
        rows=[dict(zip(headers,row)) for row in it if any(v is not None for v in row)]
    finally:book.close()
    with ZipFile(results) as archive:
        saved=table(archive.read('01_Data_dan_Tokenisasi.xlsx'),'Data dan Nilai Guru')
    if len(rows)!=len(saved):raise ValueError('Jumlah jawaban dataset dan E6 berbeda.')
    def clean(text):
        text=unicodedata.normalize('NFC',text).replace('\r\n','\n').replace('\r','\n')
        return ''.join(c for c in text if c in '\n\t' or unicodedata.category(c)!='Cc')
    prompts,links={},{}
    for i,(row,old) in enumerate(zip(rows,saved),1):
        id=f'J{i:03d}'
        if old['ID Jawaban']!=id or old['Jawaban Siswa']!=clean(row['response_text']) or old['Nilai Guru']!=row['rater']:
            raise ValueError('ID, jawaban atau label E6 tidak cocok; konteks tidak disambungkan.')
        q,r=row['soal'],row['jawaban_referensi']
        if not all(isinstance(v,str) and v.strip() for v in (q,r)):
            raise ValueError('Soal/referensi tidak lengkap.')
        key=hashlib.sha256(json.dumps([q,r],ensure_ascii=False).encode()).hexdigest()[:20]
        prompts[key]={'question':q,'reference':r};links[id]=key
    payload={'scope':'PRIVATE DISPLAY CONTEXT; E6 MODEL REMAINS ANSWER-ONLY',
             'prompts':prompts,'links':links}
    raw=json.dumps(payload,ensure_ascii=False,separators=(',',':')).encode()
    metadata={'sha256':hashlib.sha256(raw).hexdigest(),'rows':len(links),'prompt_pairs':len(prompts),
              'dataset_sha256':hashlib.sha256(dataset.read_bytes()).hexdigest(),
              'e6_results_sha256':hashlib.sha256(results.read_bytes()).hexdigest(),
              'all_ids_answers_labels_match':True,'e6_input_policy':'response_text only'}
    return raw,metadata
