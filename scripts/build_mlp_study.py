"""Buku MLP: source commit tetap, ringkasan publik dan HTML pribadi lengkap.

Tidak menjalankan kode notebook atau training. Berkas pribadi hanya ditulis
ke outputs/, bukan ke docs/. XLSX asli dibaca lokal oleh penjelajah browser.
"""
import argparse
import ast
import base64
import hashlib
import html
import io
import json
from pathlib import Path
import re
import subprocess
import tokenize
import zipfile
import xml.etree.ElementTree as ET

from mlp_study_content import CHAPTERS, GLOSSARY, PINS, FUNCTION_HINTS

ROOT = Path(__file__).resolve().parents[1]
COMMIT = '15bc7bbb2aee91ecaf44fb91a0139a8c21f64a20'
NOTEBOOK_PATH = 'notebooks/03_AES_dataset_guru_300.ipynb'
SNAPSHOT = ROOT/'docs/mlp_study_snapshot.json'
TARGET = ROOT/'docs/belajar_MLP_Spider.html'
BOOKS = ['01_Data_dan_Tokenisasi.xlsx','02_Vektor_Jawaban.xlsx','03_Vektor_Kata_dan_Kalimat.xlsx',
         '04_Prediksi_dan_Evaluasi.xlsx','05_Riwayat_Pelatihan.xlsx']


def source(cell):
    return cell['source'] if isinstance(cell['source'], str) else ''.join(cell['source'])


def sha(data):
    return hashlib.sha256(data).hexdigest()


def table(blob, sheet):
    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    rows = book[sheet].iter_rows(values_only=True)
    header = next(rows)
    result = [dict(zip(header, row)) for row in rows if any(v is not None for v in row)]
    book.close()
    return result


def workbook_inventory(blob):
    archive = zipfile.ZipFile(io.BytesIO(blob))
    ns = {'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
          'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
    rels = ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))
    targets = {r.attrib['Id']: r.attrib['Target'] for r in rels}
    result = []
    for sheet in ET.fromstring(archive.read('xl/workbook.xml')).find('s:sheets', ns):
        target = targets[sheet.attrib['{'+ns['r']+'}id']]
        path = target.lstrip('/') if target.startswith('/') else 'xl/'+target
        text = archive.read(path).decode()
        first = re.search(r'<row\b[^>]*>.*?</row>', text, re.S)
        cells = ET.fromstring(first[0]) if first else []
        headers = [''.join(c.itertext()) for c in cells]
        row_numbers = [int(m[1]) for m in re.finditer(r'<row\b[^>]*\br="(\d+)"', text)]
        result.append({'name': sheet.attrib['name'], 'rows': max(row_numbers, default=1)-1,
                       'columns': len(headers), 'headers': headers[:6]})
    return result


def extract_snapshot(notebook_bytes, zip_bytes):
    original = json.loads(notebook_bytes)
    committed = json.loads(subprocess.check_output(['git','show',COMMIT+':'+NOTEBOOK_PATH], cwd=ROOT))
    if len(original['cells']) != len(committed['cells']) or any(
        c['cell_type'] != p['cell_type'] or source(c) != source(p)
        for c,p in zip(original['cells'],committed['cells']) if c['cell_type']=='code' or p['cell_type']=='code'
    ):
        raise ValueError('Source unggahan berbeda dari MLP commit 15bc7bb; perlu audit penjelasan baru.')
    archive = zipfile.ZipFile(io.BytesIO(zip_bytes))
    if set(archive.namelist()) != set(BOOKS):
        raise ValueError('ZIP harus memuat lima workbook hasil MLP yang diketahui.')
    blobs = {name: archive.read(name) for name in BOOKS}
    metrics = table(blobs[BOOKS[3]], 'Perbandingan Model')
    folds = table(blobs[BOOKS[3]], 'Hasil Tiap Fold')
    history = {s: table(blobs[BOOKS[4]], name) for s,name in
               [('S1','RoBERTa (S1)'),('S2','BGE-M3 (S2)'),('S3','Fusion (S3)')]}
    fold_summary = table(blobs[BOOKS[4]], 'Ringkasan Lima Fold')
    # Whitelist laporan penjelasan/agregat; tidak menanam tabel siswa atau vektor ke GitHub.
    public_sheets = {}
    for index, names in {0:['Panduan'],1:['Panduan'],2:['Panduan'],
                         3:['Panduan','Ringkasan Validasi','Arti Metrik'],
                         4:['Panduan','Model yang Dibandingkan','Pengaturan Eksperimen']}.items():
        for name in names:
            public_sheets[f'{index}|{name}'] = table(blobs[BOOKS[index]], name)
    pairs = table(blobs[BOOKS[3]], 'Nilai Guru dan Prediksi')
    matrices = {}
    for scenario in ('S1','S2','S3'):
        matrix = [[0 for _ in range(11)] for _ in range(11)]
        for row in pairs:
            matrix[min(10, int(row['Nilai Guru'])//10)][min(10, int(row['Pembulatan '+scenario])//10)] += 1
        matrices[scenario] = matrix
    # Hanya agregat dan tabel training/model yang tidak memuat jawaban/ID/nilai individual.
    native = {str(i): original['cells'][i].get('outputs',[]) for i in (12,18,32)}
    native['39'] = [original['cells'][39]['outputs'][i] for i in (1,2,3)]
    return {
        'scope':'SAVED MLP RUN; NO NEW TRAINING', 'commit':COMMIT,
        'notebook_sha256':sha(notebook_bytes),'zip_sha256':sha(zip_bytes),
        'source_notebook_sha256':sha(subprocess.check_output(['git','show',COMMIT+':'+NOTEBOOK_PATH],cwd=ROOT)),
        'rows':300,'development':240,'test':len(pairs),'metrics':metrics,'folds':folds,
        'history':history,'foldSummary':fold_summary,'matrices':matrices,'native':native,
        'publicSheets':public_sheets,
        'books':[{'file':name,'sha256':sha(blobs[name]),'sheets':workbook_inventory(blobs[name])} for name in BOOKS],
        'saved_outputs':sum(len(c.get('outputs',[])) for c in original['cells']),
        'runtime_source_matches_commit':True,
    }


def highlighted_lines(text):
    lines = text.split('\n')
    spans = [[] for _ in lines]
    try:
        for token in tokenize.generate_tokens(io.StringIO(text).readline):
            kind = ('comment' if token.type==tokenize.COMMENT else 'string' if token.type==tokenize.STRING
                    else 'number' if token.type==tokenize.NUMBER else 'keyword' if token.type==tokenize.NAME and token.string in
                    {'def','class','if','elif','else','for','in','return','with','as','import','from','True','False','None','not','and','or','raise','try','except','finally','while','yield'} else None)
            if not kind:continue
            for row in range(token.start[0]-1,min(token.end[0],len(lines))):
                a = token.start[1] if row==token.start[0]-1 else 0
                b = token.end[1] if row==token.end[0]-1 else len(lines[row])
                if b>a:spans[row].append((a,b,kind))
    except (tokenize.TokenError, IndentationError):
        pass
    result=[]
    for line,parts in zip(lines,spans):
        position=0; output=''
        for a,b,kind in sorted(parts):
            if a<position:continue
            output+=html.escape(line[position:a])+f'<span class="syn-{kind}">'+html.escape(line[a:b])+'</span>'
            position=b
        result.append(output+html.escape(line[position:]))
    return result


def content(snapshot):
    notebook=json.loads(subprocess.check_output(['git','show',COMMIT+':'+NOTEBOOK_PATH],cwd=ROOT))
    cells=[];annotations=[];missing=[]
    for i,cell in enumerate(notebook['cells']):
        if cell['cell_type']!='code':continue
        text=source(cell); lines=text.split('\n')
        notes=[]
        for index,anchor,title,body,term in PINS:
            if index!=i:continue
            found=next((j+1 for j,line in enumerate(lines) if anchor in line),None)
            if found is None:missing.append([index,anchor]);continue
            notes.append({'line':found,'title':title,'body':body,'term':term,'kind':'fundamental'})
        tree=ast.parse(text.removeprefix('%%aes\n'))
        offset=1 if text.startswith('%%aes\n') else 0
        definitions=[node for node in tree.body if isinstance(node,(ast.FunctionDef,ast.ClassDef))]
        for node in definitions:
            line=node.lineno+offset
            if any(abs(n['line']-line)<=1 for n in notes):continue
            doc=ast.get_docstring(node)
            body=FUNCTION_HINTS.get(node.name) or (doc.split('\n')[0] if doc else
                 'Definisi ini menyiapkan '+('kelas' if isinstance(node,ast.ClassDef) else 'fungsi')+' '+node.name+'. Definisi belum menjalankan tahap; cari pemanggilan namanya di bagian bawah cell atau cell berikutnya. Baca bersama alur dan keluaran bab ini.')
            notes.append({'line':line,'title':node.name+'()', 'body':body,'term':None,'kind':'function'})
        notes.sort(key=lambda n:n['line'])
        for number,note in enumerate(notes):note['id']=f'c{i}-n{number}'
        chapter=next((j for j,c in enumerate(CHAPTERS) if i in c['cells'] and j!=0),0)
        cells.append({'index':i,'ordinal':len(cells)+1,'source':text,'htmlLines':highlighted_lines(text),
                      'title':definitions[0].name if definitions else 'Pengaturan & eksekusi',
                      'chapter':chapter,'notes':notes,'sha256':sha(text.encode())})
        annotations.extend(notes)
    if missing:
        raise ValueError('Anchor belum ditemukan: '+json.dumps(missing))
    return {'chapters':CHAPTERS,'glossary':GLOSSARY,'cells':cells,'snapshot':snapshot,
            'comparison':json.loads((ROOT/'docs/mlp_comparison_snapshot.json').read_text()),
            'frozen':json.loads((ROOT/'docs/mlp_frozen_snapshot.json').read_text()),
            'lora':json.loads((ROOT/'docs/mlp_lora_snapshot.json').read_text()),
            'context':json.loads((ROOT/'docs/mlp_context_snapshot.json').read_text()),
            'allResults':json.loads((ROOT/'docs/mlp_all_results_snapshot.json').read_text()),
            'markdown':[{'index':i,'source':source(c)} for i,c in enumerate(notebook['cells']) if c['cell_type']=='markdown'],
            'coverage':{'code_cells':len(cells),'code_lines':sum(len(c['htmlLines']) for c in cells),
                        'annotations':len(annotations),'manual_pins_found':len(PINS)-len(missing),
                        'manual_pins_total':len(PINS)}}


def build(args):
    notebook_bytes=Path(args.notebook).read_bytes() if args.notebook else None
    zip_bytes=Path(args.results).read_bytes() if args.results else None
    if bool(notebook_bytes)!=bool(zip_bytes):raise ValueError('Sediakan --notebook dan --results bersamaan.')
    if notebook_bytes:
        snapshot=extract_snapshot(notebook_bytes,zip_bytes)
        SNAPSHOT.write_text(json.dumps(snapshot,ensure_ascii=False,indent=2,allow_nan=False)+'\n')
    else:snapshot=json.loads(SNAPSHOT.read_text())
    data=content(snapshot)
    template=(ROOT/'scripts/templates/mlp_study.html').read_text()
    js=(ROOT/'scripts/templates/mlp_study.js').read_text()
    if js.count('/*__EXCEL_EXPLORER__*/') != 1:
        raise ValueError('Anchor penjelajah Excel harus tersedia tepat satu kali.')
    js=js.replace('/*__EXCEL_EXPLORER__*/',(ROOT/'scripts/templates/mlp_excel.js').read_text())
    js+='\n'+(ROOT/'scripts/templates/mlp_comparison.js').read_text()
    js+='\n'+(ROOT/'scripts/templates/mlp_metrics.js').read_text()
    js+='\n'+(ROOT/'scripts/templates/mlp_frozen.js').read_text()
    js+='\n'+(ROOT/'scripts/templates/mlp_lora.js').read_text()
    js+='\n'+(ROOT/'scripts/templates/mlp_dashboard.js').read_text()
    js+='\n'+(ROOT/'scripts/templates/mlp_compare_fullscreen.js').read_text()
    css=(ROOT/'scripts/templates/mlp_study.css').read_text()
    css+='\n'+(ROOT/'scripts/templates/mlp_excel.css').read_text()
    css+='\n'+(ROOT/'scripts/templates/mlp_comparison.css').read_text()
    css+='\n'+(ROOT/'scripts/templates/mlp_metrics.css').read_text()
    css+='\n'+(ROOT/'scripts/templates/mlp_frozen.css').read_text()
    css+='\n'+(ROOT/'scripts/templates/mlp_scroll.css').read_text()
    css+='\n'+(ROOT/'scripts/templates/mlp_compare_fullscreen.css').read_text()
    css+='\n'+(ROOT/'scripts/templates/mlp_dashboard.css').read_text()
    vendor=(ROOT/'scripts/vendor/mlp_jszip.min.js').read_text()
    for name,marker in [('paper_splash.svg','__SPLASH__'),('paper_city.svg','__CITY__')]:
        css=css.replace(marker,'data:image/svg+xml;base64,'+base64.b64encode((ROOT/'scripts/assets'/name).read_bytes()).decode())
    document=(template.replace('__STYLE__',css).replace('__SYMBOLS__',(ROOT/'scripts/assets/paper_spider_symbols.svg').read_text())
              .replace('__DATA__',json.dumps(data,ensure_ascii=False,allow_nan=False).replace('<',chr(92)+'u003c'))
              .replace('__VENDOR__',vendor).replace('__SCRIPT__',js))
    TARGET.write_text(document.replace('__PRIVATE_PAYLOAD__',''))
    private=None
    if notebook_bytes:
        bundle=io.BytesIO()
        with zipfile.ZipFile(bundle,'w') as z:
            z.writestr('notebook.ipynb',notebook_bytes,compress_type=zipfile.ZIP_DEFLATED)
            z.writestr('results.zip',zip_bytes,compress_type=zipfile.ZIP_STORED)
            if args.baseline:
                baseline=Path(args.baseline).read_bytes()
                if sha(baseline)!=data['comparison']['baseline']['notebook_sha256']:
                    raise ValueError('Notebook awal tidak cocok dengan snapshot pembanding.')
                z.writestr('baseline.ipynb',baseline,compress_type=zipfile.ZIP_DEFLATED)
            if args.baseline_evaluation:
                evaluation=Path(args.baseline_evaluation).read_bytes()
                if sha(evaluation)!=data['comparison']['baseline']['eval_xlsx_sha256']:
                    raise ValueError('Workbook evaluasi awal berbeda dari sumber terverifikasi.')
                z.writestr('baseline_eval.xlsx',evaluation,compress_type=zipfile.ZIP_STORED)
            if getattr(args,'frozen_results',None):
                frozen=Path(args.frozen_results).read_bytes()
                if sha(frozen)!=data['frozen']['zip_sha256']:raise ValueError('ZIP double frozen tidak cocok snapshot.')
                z.writestr('frozen_results.zip',frozen,compress_type=zipfile.ZIP_STORED)
            if getattr(args,'experiment5_results',None):
                latest=Path(args.experiment5_results).read_bytes()
                if sha(latest)!=data['lora']['zip_sha256']:raise ValueError('ZIP E5 berbeda dari snapshot.')
                z.writestr('experiment5_results.zip',latest,compress_type=zipfile.ZIP_STORED)
            if getattr(args,'context_results',None):
                context=Path(args.context_results).read_bytes()
                if sha(context)!=data['context']['zip_sha256']:raise ValueError('ZIP konteks pair berbeda dari snapshot.')
                z.writestr('context_results.zip',context,compress_type=zipfile.ZIP_STORED)
        private=ROOT/'outputs/mlp_study/belajar_MLP_Spider_Privat.html'
        private.parent.mkdir(parents=True,exist_ok=True)
        private.write_text(document.replace('__PRIVATE_PAYLOAD__',base64.b64encode(bundle.getvalue()).decode()))
    print(json.dumps({'public':str(TARGET),'private':str(private) if private else None,
                      'coverage':data['coverage'],'saved_outputs':snapshot['saved_outputs']},ensure_ascii=False))


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--notebook');p.add_argument('--results');p.add_argument('--baseline');p.add_argument('--baseline-evaluation');p.add_argument('--frozen-results');p.add_argument('--experiment5-results');p.add_argument('--context-results')
    build(p.parse_args())
