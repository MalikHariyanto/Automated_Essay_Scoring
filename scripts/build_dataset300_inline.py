"""Bangun notebook AES mandiri: seluruh kode runtime tersedia langsung di cell."""
import ast
import hashlib
import json
from pathlib import Path
import re
import textwrap

ROOT = Path(__file__).resolve().parents[1]

WORKER_SOURCE = r'''
import json, sys, traceback, types
from pathlib import Path

prefix, workspace = sys.argv[1:3]
module = types.ModuleType('__aes_notebook__')
sys.modules[module.__name__] = module
namespace = module.__dict__
namespace['ROOT'] = Path(workspace)

def event(kind, **values):
    print(prefix+json.dumps({'kind':kind, **values},ensure_ascii=False,allow_nan=False),flush=True)

def show(frame, title='', rows=8, columns=8):
    # Fungsi tampilan diisi oleh cell notebook; tidak diimpor dari file proyek.
    import pandas as pd
    frame = frame if isinstance(frame,pd.DataFrame) else pd.DataFrame(frame)
    view = namespace['preview_table'](frame.columns,frame.head(rows).to_dict('records'),title,
                                    row_count=len(frame),max_rows=rows,max_columns=columns)
    view['rows'] = json.loads(pd.DataFrame(view['rows']).to_json(orient='values',force_ascii=False,default_handler=str))
    event('table',**view,html=namespace['render_table'](view))

namespace['show'] = show
event('ready',python=sys.version.split()[0])
for line in sys.stdin:
    try:
        payload = json.loads(line)
        if payload.get('command') == 'close':break
        exec(compile(payload['code'],payload.get('filename','<aes-cell>'),'exec'),namespace)
    except BaseException:
        event('error',traceback=traceback.format_exc())
    else:
        event('done')
'''.strip()

COMMENTS = {
    'spec = cfg[role]': 'Ambil nama checkpoint dan revision model yang dipilih: roberta atau bge.',
    'result = AutoModel.from_pretrained(': 'Muat arsitektur dan bobot pretrained. Ini baru inisialisasi encoder, belum latihan memakai nilai guru.',
    'expected = 768 if role': 'Periksa panjang vektor: RoBERTa 768 angka, BGE-M3 1024 angka.',
    'result.requires_grad_(': 'RoBERTa boleh belajar; bobot encoder BGE tidak diperbarui.',
    'result = AutoTokenizer.from_pretrained(': 'Muat tokenizer dari checkpoint yang sama agar ID token cocok dengan kosakata model.',
    'self.scenario, self.roberta =': 'Simpan skenario dan encoder RoBERTa. S2 memakai vektor BGE yang sudah dihitung, jadi tidak memuat RoBERTa.',
    'roberta.requires_grad_(True)': 'Fine-tuning: bobot RoBERTa boleh berubah ketika optimizer menjalankan langkah latihan.',
    'ids = tokenizer.encode(': 'Tokenizer mengubah teks menjadi ID subword, bukan menghitung skor.',
    'windows, mapping, _ = batch_windows(': 'Pecah jawaban panjang menjadi window; mapping menghubungkan setiap window dengan jawaban asalnya.',
    'self.head = nn.Linear(768, 1)': 'S1: 768 koordinat RoBERTa dipetakan ke satu prediksi kontinu.',
    'self.head = nn.Sequential(': 'S2: MLP belajar memberi skor di atas fitur BGE yang tetap.',
    'self.projection = nn.Sequential(': 'S3: padatkan 1024 angka BGE menjadi 256 angka sebelum digabung.',
    "features['projection'] =": 'Proyeksi BGE ikut belajar, walaupun bobot encoder BGE tetap frozen.',
    "features['roberta_layernorm'] =": 'Normalisasi tambahan pada 768 komponen RoBERTa sebelum fusion; vektor asli tetap disimpan.',
    "features['concat'] = torch.cat(": 'Sambung 768 angka RoBERTa setelah LayerNorm dan 256 angka proyeksi BGE menjadi vektor fusion 1024 angka.',
    'cls = self.roberta(': 'Baca token pada setiap window, lalu ambil vektor token pertama (CLS) sebagai ringkasannya.',
    'essay = mean_by_essay(': 'Rata-ratakan vektor CLS semua window milik jawaban yang sama.',
    'prediction = self.head(': 'Lapisan prediksi memetakan vektor menjadi satu angka; saat latihan angka ini masih memakai skala target 0–1.',
    'model.eval()': 'Mode pembacaan: dropout dinonaktifkan agar ekstraksi vektor dan prediksi konsisten.',
    'self.model.eval()': 'Mode prediksi: nonaktifkan dropout. Pembacaan di bawah memakai no_grad.',
    'with torch.no_grad():': 'Tidak menghitung gradien untuk tahap pembacaan ini; bukan langkah latihan.',
    'output_vectors.append(torch.nn.functional.normalize(': 'Normalkan panjang vektor BGE dengan L2; ini mengatur vektor, bukan menghitung skor siswa.',
    'model = make_model(scenario, cfg).to(device)': 'Buat model baru untuk fold ini, lalu tempatkan di GPU atau CPU yang tersedia.',
    'return SingleScoreModel(scenario, encoder(': 'S1/S3 memuat RoBERTa; S2 hanya membuat lapisan prediksi di atas vektor BGE frozen.',
    'enc = [p for name, p in model.named_parameters()': 'Pilih bobot encoder RoBERTa yang menerima gradien.',
    'head = [p for name, p in model.named_parameters()': 'Pilih bobot head/proyeksi; kecepatannya bisa berbeda dari encoder.',
    'optimizer = torch.optim.AdamW(': 'AdamW memperbarui kelompok parameter sesuai learning rate masing-masing.',
    'targets = torch.as_tensor(': 'Ubah nilai guru dari skala 0–100 menjadi 0–1 hanya untuk menghitung loss.',
    'loss = torch.mean((prediction-targets)**2)': 'MSE mengukur selisih kuadrat antara prediksi dan target; optimizer berusaha memperkecilnya.',
    'scaler.scale(loss*': 'Hitung gradien dan bobot microbatch sesuai ukuran batch efektif aktual.',
    'scaler.step(optimizer)': 'Perbarui bobot; GradScaler bisa melewati langkah yang overflow pada AMP.',
    "if row['validation']['RMSE'] < best_rmse:": 'Simpan epoch dengan RMSE validasi terkecil; data uji akhir tidak ikut memilih epoch.',
    "return scaled*(self.cfg['score_max']": 'Kembalikan prediksi ke skala nilai guru sebelum ditampilkan atau dihitung RMSE/Pearson.',
    'clipped = np.clip(': 'Buat SALINAN skor sah untuk agreement; prediksi mentah tetap utuh.',
    'qwk = cohen_kappa_score(': 'QWK memakai jarak kategori kuadratik, bukan vektor embedding.',
    "result['AC2'] = gwet_ac2(": 'AC2 memakai kategori yang sama dengan QWK dan koreksi peluang berbeda.',
}

FUNCTION_NOTES = {
    'main':'Memeriksa pustaka yang tersedia, menyiapkan versi yang diperlukan, lalu memvalidasi import sebelum melanjutkan.',
    'active_requirements':'Pilih versi pustaka sesuai Python aktif agar aturan Python 3.12 dan 3.13 tidak tertukar.',
    'config':'Gabungkan parameter pilihan dengan nilai bawaan, lalu periksa batas skor, batch, dan panjang input.',
    'clean':'Rapikan Unicode dan karakter kontrol; ejaan, huruf besar, dan tanda baca jawaban tetap dipertahankan.',
    'read_dataset':'Baca jawaban dan nilai guru, periksa kelengkapan, lalu kelompokkan email/duplikat untuk mencegah kebocoran data.',
    'ExcelJournal':'Catat hasil setiap tahap secara privat; paket unduhan utamanya disusun menjadi lima Excel di akhir.',
    'SavedOutputs':'Baca angka hasil yang sudah tersimpan untuk laporan; tidak memanggil model atau melatih ulang.',
    'write_report':'Buat workbook dengan Panduan dan nama kolom yang jelas tanpa memotong komponen vektor.',
    'build_reports':'Susun lima laporan utama dari hasil sesi yang sudah selesai; tidak menghitung prediksi atau metrik ulang.',
    'encoder':'Memuat encoder dari nama checkpoint dan revision pada cfg; hasilnya objek model pembaca teks.',
    'check_gpu':'Memeriksa nama GPU, status T4, dan VRAM; Colab berhenti sebelum training jika GPU belum aktif.',
    'tokenizer':'Memuat pemecah teks milik checkpoint yang sama. Hasil tokenisasi menjadi input encoder.',
    'content_windows':'Membagi token isi yang panjang menjadi beberapa potongan; bagian yang bertumpang tindih menjaga konteks.',
    'batch_windows':'Menambahkan token khusus dan padding, lalu mengembalikan tensor input beserta pemetaan ke jawaban asal.',
    'SingleScoreModel':'Menyusun encoder dan lapisan prediksi untuk S1, S2, atau S3; ketiganya menghasilkan satu skor per jawaban.',
    'Engine':'Menghubungkan teks, tokenizer, dan model. Saat prediksi, hasil dikembalikan ke skala nilai guru.',
    'make_model':'Membuat model baru agar bobot hasil fold sebelumnya tidak terbawa ke fold berikutnya.',
    'optimizer_groups':'Memisahkan parameter RoBERTa pretrained, proyeksi BGE, dan head agar learning rate setiap bagian jelas.',
    'bge_features':'Menghitung vektor BGE-M3 sekali karena encodernya frozen; vektor dipakai bersama oleh S2 dan S3.',
    'roberta_pretrained_vectors':'Menyimpan vektor RoBERTa sebelum latihan sebagai pembanding untuk hasil sesudah fine-tuning.',
    'token_vectors':'Mengambil vektor lookup awal dan vektor setelah membaca konteks, berikut posisi token di teks.',
    'word_vectors':'Merata-ratakan potongan subword yang membentuk satu kata agar perbandingan bisa dibaca per kata.',
    'compare_repetitions':'Membandingkan dua kemunculan kata yang sama di model yang sama, sebelum dan sesudah konteks.',
    'metrics':'Membandingkan nilai guru dan prediksi: RMSE/Pearson memakai skor mentah; QWK/AC2 memakai salinan kategori.',
    'fit':'Melatih satu model pada satu pembagian data; RMSE validasi menentukan epoch yang disimpan.',
    'save_training_progress':'Menyimpan bobot, keadaan optimizer, urutan acak dan riwayat setelah epoch lengkap; satu cadangan tetap disimpan.',
    'load_training_progress':'Memeriksa kecocokan dan keutuhan checkpoint sebelum melanjutkan epoch berikutnya.',
    'train_all':'Menjalankan lima fold tiap skenario, latihan final, lalu evaluasi data uji setelah pilihan model dikunci.',
    'score_answer':'Membaca satu jawaban dengan ketiga model final untuk menghasilkan prediksi; tidak melatih model lagi.',
    'export_final_features':'Menyimpan vektor model yang sudah dilatih agar dapat dibandingkan dengan vektor sebelum latihan.',
    'test_score_comparison':'Membaca prediksi test tersimpan, menyamakan ID jawaban, lalu menampilkan nilai guru, prediksi mentah dan pembulatannya.',
}


def source(path, names):
    raw = (ROOT/path).read_text(); lines = raw.splitlines(); found = []
    for node in ast.parse(raw).body:
        if isinstance(node, (ast.FunctionDef, ast.ClassDef)) and node.name in names:
            first = min([node.lineno]+[x.lineno for x in getattr(node,'decorator_list',[])])
            block = '\n'.join(lines[first-1:node.end_lineno])
            # Imports relatif diganti nama yang sudah didefinisikan di cell sebelumnya.
            annotated = []
            for line in block.splitlines():
                if line.lstrip().startswith(('from .','from aes.')):
                    continue
                for prefix, explanation in COMMENTS.items():
                    if line.lstrip().startswith(prefix):
                        annotated.append(line[:len(line)-len(line.lstrip())]+'# '+explanation)
                annotated.append(line)
            note=FUNCTION_NOTES.get(node.name)
            found.append(('# '+note+'\n' if note else '')+'\n'.join(annotated))
    if len(found) != len(names):
        raise ValueError('Definisi source hilang: '+str(names))
    return '\n\n\n'.join(found)


def assignments(path,names):
    """Salin konstanta publik sebagai kode terbaca, bukan modul tersembunyi."""
    raw=(ROOT/path).read_text();blocks=[]
    for node in ast.parse(raw).body:
        if isinstance(node,ast.Assign) and any(isinstance(t,ast.Name) and t.id in names for t in node.targets):
            blocks.append(ast.get_source_segment(raw,node))
    if len(blocks)!=len(names):raise ValueError('Konstanta belum lengkap: '+str(names))
    return '\n\n'.join(blocks)


def references_markdown():
    papers=json.loads((ROOT/'docs/aes_dataset300_references.json').read_text())
    lines=['<details><summary>Rujukan metode — catatan bacaan, bukan input program</summary>','',
           'Rujukan ini mendukung penjelasan metode. Notebook tidak mengunduh atau membaca daftar pustaka untuk menghitung skor.','']
    for paper in papers:
        lines.append(f"- **{paper['id']} · {paper['title']}** ({paper['year']}). {paper['authors']}. {paper['used_for']} [Sumber]({paper['url']}).")
        lines.append('  Batas penerapan: '+paper['limitation'])
    return '\n'.join(lines+['','</details>'])


def standalone_run_folder():
    tree=ast.parse((ROOT/'src/aes/teacher_score/data.py').read_text())
    function=next(n for n in tree.body if isinstance(n,ast.FunctionDef) and n.name=='run_folder')
    body=[]
    for node in function.body:
        name=node.targets[0].id if isinstance(node,ast.Assign) and isinstance(node.targets[0],ast.Name) else None
        if name in ('package','code_files'):continue
        if name=='code_hash':node.value=ast.Name(id='NOTEBOOK_CODE_SHA256',ctx=ast.Load())
        body.append(node)
    function.body=body
    return '# Identitas hasil mengikuti data, konfigurasi, dan kode notebook ini.\n'+ast.unparse(ast.fix_missing_locations(function))


def build():
    cells = []
    def md(text):
        if '## 12 ·' in text:
            text='''## 12 · Pelatihan dan evaluasi tiga skenario

            Model dilatih dengan lima fold pada data development. Epoch terbaik dipilih dari RMSE validasi, kemudian ketiga model dilatih ulang pada development sebelum final test dibuka. Cara ini tetap sama seperti versi sebelumnya.

            Riwayat seluruh fold digabung per model pada **05_Riwayat_Pelatihan.xlsx**. Prediksi dan empat metrik tersedia pada **04_Prediksi_dan_Evaluasi.xlsx**. File rincian per fold tetap disimpan oleh program, tetapi tidak dimasukkan ke paket utama.'''
        elif '## 14 ·' in text:
            text='''## 14 · Lihat vektor setelah model belajar

            Tahap penghitungan tetap mengambil fitur model final setelah evaluasi. Untuk laporan utama, cukup gunakan vektor RoBERTa sebelum/sesudah latihan, BGE-M3, dan gabungan fusion pada **02_Vektor_Jawaban.xlsx**.

            Vektor kata dan kalimat contoh tersedia pada **03_Vektor_Kata_dan_Kalimat.xlsx**. Angka lapisan internal seperti dropout atau aktivasi tidak perlu dilampirkan ke dosen; rincian itu tetap disimpan secara privat oleh program.'''
        elif '## 15 ·' in text:
            text='''## 15 · Unduh hasil Excel yang sudah rapi

            Jalankan **satu cell terakhir di bawah**. Colab utama ini langsung menyusun hasil menjadi lima Excel dan mengunduh **Hasil_AES_Ringkas.zip**. Tidak perlu membuka notebook lain atau mengunggah hasil untuk dirapikan lagi.

            Paket akhir **Hasil_AES_Ringkas.zip** berisi:
            1. **Data dan Tokenisasi** — jawaban, nilai guru, pembagian data, serta token.
            2. **Vektor Jawaban** — seluruh komponen vektor per jawaban, sebelum/sesudah latihan dan fusion.
            3. **Vektor Kata dan Kalimat** — satu contoh jawaban dan perbandingan kata berulang.
            4. **Prediksi dan Evaluasi** — buka sheet **Perbandingan Model** terlebih dahulu.
            5. **Riwayat Pelatihan** — fold dan epoch digabung per model.

            Tiap file diawali sheet **Panduan**. Nama kolom sudah dirapikan; semua komponen vektor utama tetap lengkap. Rincian lapisan internal tidak masuk ZIP unduhan.

            Jika proses baru selesai sebagian, paket hanya memuat tahap yang tersedia, tanpa hasil pengganti. Untuk mengunduh ulang di sesi yang sama, cukup jalankan cell terakhir ini; cell ini tidak melatih atau mengevaluasi model ulang.'''
        replacements={'01_konfigurasi.xlsx':'05_Riwayat_Pelatihan.xlsx','01_rujukan_paper.xlsx':'05_Riwayat_Pelatihan.xlsx',
            '02_dataset_dan_cleaning.xlsx':'01_Data_dan_Tokenisasi.xlsx','03_split_dan_cv.xlsx':'01_Data_dan_Tokenisasi.xlsx',
            '04_tokenisasi.xlsx':'01_Data_dan_Tokenisasi.xlsx','05_roberta_pretrained.xlsx':'02_Vektor_Jawaban.xlsx',
            '06_bge_frozen.xlsx':'02_Vektor_Jawaban.xlsx','07_vektor_kata_kalimat_pretrained.xlsx':'03_Vektor_Kata_dan_Kalimat.xlsx',
            '08_arsitektur.xlsx':'05_Riwayat_Pelatihan.xlsx','13_prediksi_satu_jawaban.xlsx':'04_Prediksi_dan_Evaluasi.xlsx'}
        for old,new in replacements.items():text=text.replace(old,new)
        text=text.replace('**Tiap tahap penghitungan menyimpan Excel** dengan `Info`/provenance dan seluruh koordinat per kolom.',
                          '**Hasil akhir dirangkum dalam lima Excel** dengan nama sheet dan kolom yang mudah dibaca. Rincian setiap tahap tetap tersimpan secara privat.')
        text = re.sub(r'([A-Za-z]{3,})(\d)',r'\1 \2',text)
        text = re.sub(r'(\d)([A-Za-z]{3,})',r'\1 \2',text)
        for before,after in {'seluruhjawaban':'seluruh jawaban','tiapfold':'tiap fold','setiapfold':'setiap fold',
                             'setiapwindow':'setiap window','setiapepoch':'setiap epoch','semua skenario':'semua skenario',
                             'fullfine-tuned':'fully fine-tuned','modeleval':'model eval','checkpointfinal':'checkpoint final',
                             'nilaiguru':'nilai guru','teksbaru':'teks baru','melatihulang':'melatih ulang',
                             'out-of-range':'out-of-range','dataframe':'DataFrame','tidaktercakup':'tidak tercakup',
                             'skor/model':'skor/model','kelasXI':'kelas XI','finaltest':'final test',
                             'sebelumtest':'sebelum test','ditest':'di test','earlystopping':'early stopping',
                             'batchlain':'batch lain','metriklain':'metrik lain','fiturmanual':'fitur manual',
                             'fileunggahan':'file unggahan','bukanskor':'bukan skor','barubukan':'baru bukan',
                             'selesai.':'selesai.','testbaru':'test baru','Dropoutnonaktif':'Dropout nonaktif',
                             'bobot/cache':'bobot/cache','republik':'repo publik','repopublik':'repo publik',
                             'outputfusion':'output fusion','untukinspeksi':'untuk inspeksi','untukpelaporan':'untuk pelaporan'}.items():
            text=text.replace(before,after)
        cells.append(dict(cell_type='markdown',metadata={},source=textwrap.dedent(text).strip()+'\n'))
    def code(text, worker=True):
        cells.append(dict(cell_type='code',metadata={},execution_count=None,outputs=[],
                          source=('%%aes\n' if worker else '')+textwrap.dedent(text).strip()+'\n'))
    def defs(path,names,comment):
        code('# '+comment+'\n'+source(path,names))
    md('''
    # AES guru300 · RoBERTa, BGE-M3, fusion → prediksi dan evaluasi

    Urutan mengikuti **BGE-M3 [AES].ipynb** unggahan: instalasi/impor → Excel/dataset → split/tokenisasi → embedding → arsitektur → loss/optimizer → training → prediksi/evaluasi. **Kode proses ada di cell dengan komentar `#`; output tidak lagi mencetak source kode.**

    **S1** RoBERTa Indonesia fine-tuned. **S2** BGE-M3 frozen+MLP. **S3** fusion RoBERTa fine-tuned dan BGE frozen. Ini mempertahankan rancangan sebelumnya; notebook BGE unggahan sendiri memperbarui seluruh BGE. Dataset hanya mempunyai satu `rater`, jadi setiap model menghasilkan **satu skor**, bukan empat label rubrik yang tidak ada.

    Gunakan GPU Colab, Python 3.12/3.13. Jalankan berurutan. **Batas maksimal sama: 25 epoch per fold untuk S1/S2/S3**, dengan early stopping setelah lima epoch berturut-turut tanpa penurunan RMSE validasi. Bobot dengan RMSE validasi terkecil selalu dipakai. Ada 5 fold, 15 pelatihan CV dan 3 pelatihan final; proses bisa melewati sesi gratis. **Tiap tahap penghitungan menyimpan Excel** dengan `Info`/provenance dan seluruh koordinat per kolom. Preview notebook hanya cuplikan. Seluruh jawaban: embedding per jawaban; satu jawaban contoh: vektor token/kata/kalimat lengkap. Definisi fungsi menyiapkan proses, bukan menghasilkan skor dengan sendirinya.

    **Notebook mandiri:** semua kode pendukung ada di cell. Tidak ada clone/pull GitHub atau import modul proyek. Koneksi internet tetap dipakai saat memasang pustaka dan mengambil bobot pretrained dari Hugging Face.

    **Simpan dan lanjut otomatis:** di Colab, hubungkan Google Drive saat diminta. Dataset, hasil, dan checkpoint disimpan di `MyDrive/AES_300_Privat`; pustaka/cache bobot tetap lokal. Setiap epoch yang selesai disimpan lengkap, dengan satu cadangan. Bila runtime terputus, pilih GPU yang sama dan jalankan notebook ini lagi dari atas; epoch tersimpan dilanjutkan, fold selesai dilewati. Epoch yang terputus sebelum penyimpanan selesai diulang. Colab memeriksa GPU/T4 sebelum training; GPU dipilih melalui menu Runtime, bukan otomatis dialokasikan oleh kode.

    File dataset, Excel, checkpoint dan notebook terisi tetap privat. GitHub hanya tempat menyimpan salinan notebook. Penyimpanan Drive membutuhkan ruang kosong; penulisan checkpoint menambah waktu tiap epoch. Pesan “progres tersimpan” baru muncul sesudah penulisan selesai. Progres lama yang sudah hilang sebelum fitur ini tidak bisa dipulihkan.
    ''')
    md(references_markdown())
    md('''
    ## 1 · Siapkan ruang kerja

    Cell ini menyiapkan folder kerja untuk dataset dan hasil. Daftar versi pustaka ditulis dari teks yang ada di cell; tidak mengambil kode atau daftar pustaka dari repo. Semua kode model dan pendukungnya tersedia di notebook ini.
    ''')
    code(textwrap.dedent('''
    import os, sys, json, subprocess, html, secrets, atexit, argparse, hashlib, venv
    from pathlib import Path
    from IPython.display import display, HTML
    try:
        from packaging.requirements import Requirement
    except ImportError:
        from pip._vendor.packaging.requirements import Requirement
    IN_COLAB = 'google.colab' in sys.modules or Path('/content').is_dir()
    default_workspace='/content/AES_300' if IN_COLAB else str(Path.cwd()/'outputs/AES_300')
    ROOT=Path(os.environ.get('AES300_WORKSPACE',default_workspace)).expanduser().absolute()
    ROOT.mkdir(parents=True,exist_ok=True)
    def prepare_persistent_storage(workspace,in_colab):
        # Drive hanya untuk data/hasil/checkpoint; pustaka dan cache bobot tetap lokal agar lebih cepat.
        if in_colab:
            from google.colab import drive
            drive.mount('/content/drive')  # Colab meminta izin akun; setelah tersambung, penyimpanan otomatis.
            if not Path('/content/drive/MyDrive').is_dir():
                raise RuntimeError('Google Drive belum tersambung. Training tidak dimulai tanpa penyimpanan permanen.')
            default_storage='/content/drive/MyDrive/AES_300_Privat'
        else:
            default_storage=str(workspace)
        storage=Path(os.environ.get('AES300_PERSISTENT_DIR',default_storage)).expanduser().absolute()
        storage.mkdir(parents=True,exist_ok=True)
        # Gagal menulis harus terlihat sebelum training, bukan diam-diam menyimpan ke /content.
        probe=storage/'.aes_write_check'
        probe.write_text('write check');probe.unlink()
        return storage
    # Identitas kode notebook untuk membedakan hasil dari implementasi yang berbeda.
    NOTEBOOK_CODE_SHA256='__NOTEBOOK_CODE_SHA256__'
    print('Ruang kerja AES siap. Seluruh kode tersedia di notebook ini.')
    print('Versi Python:',sys.version.split()[0])
    ''')+'\n# Hanya pustaka runtime yang diperlukan; alat pengujian tidak dipasang untuk menjalankan AES.\n'+
         "REQUIREMENTS_TEXT=r'''\n"+'\n'.join(line for line in (ROOT/'requirements.dataset300.txt').read_text().splitlines()
                                                if not line.startswith(('pytest','nbformat','nbclient')))+"\n'''"+
         "\n(ROOT/'requirements.dataset300.txt').write_text(REQUIREMENTS_TEXT)\n",False)
    md('''
    ## 2 · Instalasi dan pemilih Excel

    Setup memeriksa versi pustaka yang sudah tersedia. Versi yang cocok dipakai ulang; versi yang berbeda dipasang dalam lingkungan khusus AES tanpa mencoba menghapus pustaka bawaan. PyTorch runtime dipertahankan agar akses GPU yang tersedia tidak diganti wheel CPU. Pemeriksaan mencakup versi, dependensi pustaka yang dipakai AES, dan import; bukan klaim semua aplikasi bawaan Colab memakai versi yang sama.

    Output utama berisi status singkat. **Rincian instalasi** bisa dibuka jika diperlukan. Jika pemeriksaan gagal, cell berhenti dengan alasan dan log, bukan menampilkan status siap. Python worker numerik terpisah menghindari NumPy lama yang sudah dimuat kernel.

    Awalan **`%%aes` menjalankan kode cell pada worker yang sama**, sehingga variabel bertahan antar-cell. Upload Excel sekali per runtime; bila file sudah ada tidak diminta lagi.
    Drive dihubungkan pada tahap ini. Jika izin Drive atau penulisan gagal, cell berhenti dan tidak diam-diam beralih ke penyimpanan sementara. Dataset yang sudah ada di folder privat Drive tidak perlu diunggah ulang setelah runtime berganti. Jangan menjalankan dua sesi training pada folder percobaan yang sama secara bersamaan.
    ''')
    code('# Pemeriksaan dan pemasangan pustaka tersedia langsung di sini.\n'+
         assignments('scripts/setup_dataset300.py',['INSPECT','IMPORTS'])+'\n\n'+
         source('scripts/setup_dataset300.py',['environment_folder','active_requirements','SetupError','main'])
            .replace('def main(argv=None):','def prepare_environment(argv=None):'),False)
    code('# Penghubung proses Python: menjaga variabel antar-cell %%aes tanpa file pendukung proyek.\n'+
         "WORKER_SOURCE=r'''\n"+WORKER_SOURCE+"\n'''\n",False)
    code(r'''
    PYTHON = os.environ.get('AES300_PYTHON','')
    def show_installation_log():
        path=ROOT/'outputs/dataset300_setup.log'
        contents=path.read_text() if path.is_file() else 'Log belum tersedia.'
        display(HTML('<details style="background:#f8fafc;color:#172035;padding:12px;border:1px solid #94a3b8">'
                     '<summary>Rincian instalasi — buka jika diperlukan</summary>'
                     '<pre style="max-height:320px;overflow:auto;white-space:pre-wrap;color:#172035">'
                     +html.escape(contents)+'</pre></details>'))
    try:
        ready=prepare_environment(['--verify-python',PYTHON] if PYTHON else [])
    except Exception:
        show_installation_log()
        raise
    if ready['runtime_version'] != list(sys.version_info[:2]) or not ready['imports_verified']:
        raise RuntimeError('Metadata setup tidak cocok runtime aktif.')
    PYTHON=ready['python']
    setup_rows=[
        ('Status','Siap — versi, dependensi, dan import AES lulus pemeriksaan'),
        ('Python',ready['python_version']),
        ('PyTorch',ready['torch']),
        ('Perangkat','GPU: '+ready['device'] if ready['cuda_available'] else 'CPU — GPU belum terdeteksi'),
        ('Pustaka AES','Versi yang cocok dipakai ulang' if ready['packages_action']=='reused'
                        else 'Versi khusus AES dipasang; pustaka bawaan tetap disimpan'),
        ('Pustaka utama','NumPy '+ready['versions']['numpy']+'; Transformers '+ready['versions']['transformers']),
        ('Model','Belum dimuat; tahap ini hanya menyiapkan lingkungan')]
    status_html=''.join('<tr><th style="padding:9px;text-align:left;border:1px solid #94a3b8">'+html.escape(label)
                        +'</th><td style="padding:9px;border:1px solid #94a3b8">'+html.escape(value)+'</td></tr>'
                        for label,value in setup_rows)
    display(HTML('<table style="background:#f8fafc!important;color:#172035!important;border-collapse:collapse">'
                 +status_html+'</table>'))
    show_installation_log()
    PERSISTENT_ROOT=prepare_persistent_storage(ROOT,IN_COLAB)
    print('Penyimpanan data, checkpoint dan Excel:',PERSISTENT_ROOT)
    print('Progres disimpan otomatis setiap epoch. Saat sesi terputus, jalankan kembali notebook yang sama dari atas.')
    DATASET_NAME='fix_dataset_300.xlsx'
    PRIVATE_INPUT=PERSISTENT_ROOT/'data/dataset300';PRIVATE_INPUT.mkdir(parents=True,exist_ok=True)
    DATA_PATH=Path(os.environ.get('AES300_DATASET',str(PRIVATE_INPUT/DATASET_NAME))).resolve()
    if not DATA_PATH.exists():
        if not IN_COLAB:
            raise FileNotFoundError(str(DATA_PATH))
        from google.colab import files
        previous=Path.cwd()
        try:
            os.chdir(PRIVATE_INPUT);uploaded=files.upload()
        finally:
            os.chdir(previous)
        if DATASET_NAME not in uploaded:
            raise ValueError('Upload '+DATASET_NAME)
    def display_aes_table(event):
        display(HTML(event['html']))
    if '_aes_process' in globals() and _aes_process.poll() is None:
        _aes_process.terminate()
    _aes_prefix='__AES_'+secrets.token_hex(12)+'__'
    child_env={**os.environ,'HF_HOME':str(ROOT/'outputs/cache/huggingface'),
               'AES300_DATASET':str(DATA_PATH),'AES300_RESULTS':str(PERSISTENT_ROOT/'outputs/dataset300_inline'),
               'AES300_REQUIRE_GPU':'1' if IN_COLAB else os.environ.get('AES300_REQUIRE_GPU','0'),
               'TOKENIZERS_PARALLELISM':'false',
               'OMP_NUM_THREADS':'2','MKL_NUM_THREADS':'2','PYTHONUNBUFFERED':'1'}
    child_env.pop('PYTHONPATH',None)
    _aes_process=subprocess.Popen([PYTHON,'-u','-c',WORKER_SOURCE,_aes_prefix,str(ROOT)],
                                  cwd=ROOT,env=child_env,stdin=subprocess.PIPE,stdout=subprocess.PIPE,
                                  stderr=subprocess.STDOUT,text=True,bufsize=1)
    while True:
        line=_aes_process.stdout.readline()
        if not line:
            raise RuntimeError('Worker gagal mulai; lihat log setup.')
        if line.startswith(_aes_prefix):
            print('Proses perhitungan siap. Python',json.loads(line[len(_aes_prefix):])['python']);break
        print(line,end='')
    def aes_cell(line,cell):
        if _aes_process.poll() is not None:
            raise RuntimeError('Worker berhenti; jalankan setup dan cell berikutnya dari awal.')
        _aes_process.stdin.write(json.dumps({'code':cell,'filename':'<notebook-cell>'})+'\n')
        _aes_process.stdin.flush()
        while True:
            output=_aes_process.stdout.readline()
            if not output:
                raise RuntimeError('Worker terhenti; jika RAM penuh pilih runtime lebih besar.')
            if not output.startswith(_aes_prefix):
                print(output,end='',flush=True);continue
            event=json.loads(output[len(_aes_prefix):])
            if event['kind']=='done':break
            if event['kind']=='error':raise RuntimeError(event['traceback'])
            if event['kind']=='table':display_aes_table(event)
    get_ipython().register_magic_function(aes_cell,'cell','aes')
    atexit.register(lambda:_aes_process.terminate() if _aes_process.poll() is None else None)
    print('Dataset siap. Jalankan cell %%aes berurutan.')
    ''',False)
    md('''
    ## 3 · Impor dan konfigurasi sebelum eksperimen

    Impor pustaka yang diperlukan, lalu tentukan checkpoint, epoch, dan batch. Fungsi pendukung untuk membaca data, menampilkan tabel, dan membuat Excel juga didefinisikan di cell notebook ini. Rujukan tetap menjadi catatan Markdown, tidak dibaca oleh program.

    Batas 25 epoch dan kesabaran lima epoch berlaku sama untuk ketiga skenario. Learning rate kandidat A memakai 0,00001 pada encoder RoBERTa dan 0,0001 pada proyeksi/head. Konfigurasi aktual dicatat setelah dataset selesai diperiksa.
    ''')
    code('''
    import os, sys, json, math, hashlib, random, warnings, re, itertools, unicodedata, zipfile, subprocess, gc, logging, html
    from collections import Counter
    from dataclasses import dataclass
    from pathlib import Path
    from importlib.metadata import version
    import numpy as np
    import pandas as pd
    import torch
    from torch import nn
    from torch.nn.functional import normalize
    from torch.utils.data import DataLoader
    from transformers import AutoModel, AutoTokenizer, get_linear_schedule_with_warmup
    from scipy.stats import pearsonr
    from sklearn.metrics import cohen_kappa_score
    from sklearn.model_selection import GroupKFold, GroupShuffleSplit
    from openpyxl import Workbook, load_workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    SCENARIOS=('S1','S2','S3')
    NOTEBOOK_CODE_SHA256='__NOTEBOOK_CODE_SHA256__'
    ''')
    code('# Nilai bawaan dan fungsi konfigurasi, tersedia langsung di notebook.\n'+
         assignments('src/aes/teacher_score/__init__.py',['SCOPE'])+'\n\n'+
         assignments('src/aes/teacher_score/data.py',['DEFAULTS','REQUIRED'])+'\n\n'+
         source('src/aes/teacher_score/data.py',['config']))
    code('# Bantuan membaca tabel: penamaan kolom, angka, dan penjelasan kata berulang.\n'+
         assignments('src/aes/teacher_score/presentation.py',['LABELS','REPEATED_COLUMNS','MODELS','ROLES'])+'\n\n'+
         assignments('src/aes/teacher_score/display.py',['REPEATED_NOTES','METRIC_NOTES'])+'\n\n'+
         'REPEATED_HEADERS=REPEATED_COLUMNS\n\n'+
         source('src/aes/teacher_score/presentation.py',['friendly_header','friendly_value','repeated_word_rows'])+'\n\n'+
         source('src/aes/teacher_score/display.py',['preview_table','cell_text','render_table'])+'\n\n'+
         source('src/aes/teacher_score/models.py',['check_gpu']))
    code('''

    # Nama checkpoint adalah alamat model pretrained di Hugging Face, bukan objek model yang sudah dimuat.
    ROBERTA_CHECKPOINT='flax-community/indonesian-roberta-base'  # RoBERTa bahasa Indonesia untuk S1 dan S3.
    BGE_CHECKPOINT='BAAI/bge-m3'  # BGE-M3 untuk menghasilkan embedding di S2 dan S3.
    # Revision mengunci versi bobot/tokenizer agar eksperimen berikutnya memakai sumber yang sama.
    ROBERTA_REVISION=DEFAULTS['roberta']['revision']
    BGE_REVISION=DEFAULTS['bge']['revision']

    # Ubah parameter DI SINI sebelum latihan; jangan memilihnya dari hasil data uji.
    # Batas maksimal sama untuk tiga skenario; setiap fold memulai model baru.
    EPOCHS,FOLDS=25,5
    # Berhenti jika RMSE validasi tidak turun selama lima epoch berturut-turut.
    EARLY_STOPPING_PATIENCE=5
    # Kandidat A: encoder pretrained berubah perlahan; proyeksi dan head baru belajar lebih cepat.
    ENCODER_LR=1e-5
    PROJECTION_LR=HEAD_LR=1e-4
    # Proses satu jawaban per microbatch, lalu akumulasikan gradien hingga batch efektif delapan jawaban.
    MICRO_BATCH,EFFECTIVE_BATCH=1,8
    NEW_ANSWER=''  # Kosong: satu jawaban held-out untuk contoh scoring/vektor.
    AUDIT_ONLY=os.environ.get('AES300_AUDIT_ONLY')=='1'
    RUN_MODELS=RUN_TRAINING=RUN_WORD_VECTORS=RUN_FINAL_FEATURES=not AUDIT_ONLY
    cfg=config(dict(dataset=os.environ['AES300_DATASET'],output=os.environ['AES300_RESULTS'],
                    sheet='Sheet1',expected_rows=300,score_min=0,score_max=100,
                    epochs=EPOCHS,folds=FOLDS,micro_batch=MICRO_BATCH,effective_batch=EFFECTIVE_BATCH,
                    early_stopping_patience=EARLY_STOPPING_PATIENCE,
                    encoder_lr=ENCODER_LR,projection_lr=PROJECTION_LR,head_lr=HEAD_LR,
                    roberta={'checkpoint':ROBERTA_CHECKPOINT,'revision':ROBERTA_REVISION},
                    bge={'checkpoint':BGE_CHECKPOINT,'revision':BGE_REVISION},
                    device='auto',new_answer=NEW_ANSWER,vector_source='pretrained',lesson_version='candidate_a_layernorm_25_v1',
                    runtime_versions={p:version(p) for p in ('torch','transformers','numpy')}))
    # Colab harus mendeteksi GPU sebelum latihan panjang dimulai. Pengujian lokal tetap boleh CPU.
    gpu=check_gpu(require_gpu=RUN_MODELS and os.environ.get('AES300_REQUIRE_GPU')=='1')
    show(pd.DataFrame([
        {'Pemeriksaan':'Perangkat aktif','Hasil':gpu['name']},
        {'Pemeriksaan':'CUDA tersedia','Hasil':'Ya' if gpu['cuda_available'] else 'Tidak'},
        {'Pemeriksaan':'GPU T4','Hasil':'T4 terdeteksi' if gpu['is_t4'] else 'GPU lain terdeteksi' if gpu['cuda_available'] else 'Belum aktif'},
        {'Pemeriksaan':'Memori GPU (GiB)','Hasil':str(gpu['memory_gib']) if gpu['memory_gib'] is not None else 'Tidak tersedia'}
    ]),'Pemeriksaan GPU sebelum training',columns=2)
    ''')
    md('''
    ## 4 · Baca Excel dan cleaning

    Hanya `response_text` masuk encoder. Soal/kunci/email/skor tidak menjadi fitur. Cleaning mempertahankan ejaan, kapital, tanda baca/paragraf; sumber tidak ditimpa. Dataset satu nilai tidak memberi empat label rubrik; file bertopik metodologi penelitian ini juga tidak membuktikan populasi SMAkelasXI.

    **Output:** dataset asli, audit, teks sebelum/sesudah dan target loss `(rater−0)/100`→`02_dataset_dan_cleaning.xlsx`. Rujukan m01 menempatkan pemahaman/persiapan data sebelum modeling.
    ''')
    code('# Baca dan audit data tanpa mengubah file Excel asli.\n'+
         source('src/aes/cleaning.py',['clean'])+'\n\n'+
         source('src/aes/teacher_score/data.py',['digest','dump','Bundle','read_dataset'])+'\n\n'+standalone_run_folder())
    code('# Catat hasil tahap demi tahap; laporan unduhan diringkas menjadi lima Excel di bagian terakhir.\n'+
         source('src/aes/teacher_score/excel.py',['excel_value','write_excel','vector_frame','ExcelJournal']))
    code('''
    bundle=read_dataset(cfg)  # Periksa enam header, label, duplikat dan kelompok email.
    output=run_folder(cfg,bundle)
    journal=ExcelJournal(output/'excel',{'scope':SCOPE,'dataset_sha256':bundle.source_sha256,
                         'run_folder':str(output),'notebook_code_sha256':NOTEBOOK_CODE_SHA256,
                         'privacy':'PRIVAT; bukan untuk repo publik'})
    journal.save('01_konfigurasi',{'Parameter':pd.DataFrame([{'name':k,'value':v} for k,v in cfg.items()]),
                                 'Versi':pd.DataFrame([{'library':p,'version':version(p)} for p in ('torch','transformers','numpy','pandas','scikit-learn','openpyxl')])},
                 'Konfigurasi aktual; belum ada hasil training')
    df=pd.read_excel(cfg['dataset'],sheet_name=cfg['sheet'])
    df=df.loc[:,~df.columns.astype(str).str.startswith('Unnamed:')]
    cleaning=pd.DataFrame([{'row_id':r['row_id'],'teks_asli':r['raw_text'],'teks_model':r['text'],
                           'berubah':r['raw_text']!=r['text'],'skor_guru':r['score'],
                           'target_loss':(r['score']-cfg['score_min'])/(cfg['score_max']-cfg['score_min'])} for r in bundle.rows])
    journal.save('02_dataset_dan_cleaning',{'Dataset_asli':df,'Audit':pd.DataFrame([bundle.summary]),
                                          'Sebelum_sesudah':cleaning},'Pembacaan/cleaning aktual; sumber utuh')
    show(cleaning[['row_id','teks_model','skor_guru']],'Contoh jawaban dan nilai guru')
    print(len(df),'jawaban tersedia. Setiap jawaban memiliki satu nilai guru; setiap model menghasilkan satu prediksi skor.')
    ''')
    md('''
    ## 5 · Split80/20 dan5-fold validasi

    Email/duplikat tidak menyeberang split. Kandidat dipilih mendekati80/20 berdasarkan jumlah baris saja, tanpa melihat skor/model. Lima fold berada di development; tidak diperlukan split validasi ketiga. **Output:**`03_split_dan_cv.xlsx`, setiapbaris/kelompok/fold. Ini menguji jawaban kelompok baru pada pertanyaan yang sama, bukan unseen-prompt. Rujukan m09 membahas dependensi fold, bukan kewajiban universal5fold.
    ''')
    code(source('src/aes/teacher_score/data.py',['make_split'])+'''

manifest=make_split(bundle,cfg)
dump(output/'split_manifest.json',manifest)
roles={i:'development' for i in manifest['development']}
roles.update({i:'test' for i in manifest['test']})
split_rows=pd.DataFrame([{'row_id':r['row_id'],'role':roles[i],'connected_group':int(bundle.groups[i])} for i,r in enumerate(bundle.rows)])
fold_rows=pd.DataFrame([{'fold':f['fold'],'row_id':bundle.rows[i]['row_id'],'role':role} for f in manifest['folds'] for role,key in [('train','train'),('validation','validation')] for i in f[key]])
journal.save('03_split_dan_cv',{'Split_setiap_baris':split_rows,'Fold_setiap_baris':fold_rows,
                              'Ringkasan':pd.DataFrame([{k:v for k,v in manifest.items() if k not in ('folds','development','test')}])},'Split dan5fold sama untukS1/S2/S3')
show(split_rows.groupby('role').size().rename('rows').reset_index(),'Pembagian aktual')
trace_index=manifest['test'][cfg['trace_position']]
trace_text=cfg.get('new_answer','').strip() or bundle.rows[trace_index]['text']
print('Jawaban contoh:',friendly_value('row_id',bundle.rows[trace_index]['row_id']),'. Contoh ini dipakai untuk melihat token, vektor, dan prediksi setelah model dipilih.')
''')
    md('''
    ## 6 · Tokenizer dan window yang benar-benar dipakai

    RoBERTa512=510isi+2special, overlap128, seluruh token akhir tercakup. BGE memakai tokenizer berbeda, batas eksperimen1024 termasukspecial; truncation dilaporkan. Dasar tokenizer/encoder: m04/p14; batas/overlap ini keputusan proyek. **Output:**token,ID,attentionmask,padding,window/coverage→`04_tokenisasi.xlsx`.
    ''')
    defs('src/aes/windows.py',['content_windows','batch_windows','mean_by_essay'],'Buat window, tambahkan special sekali, mean CLS secara differentiable.')
    defs('src/aes/teacher_score/models.py',['device_name','release','ExpectedPoolerNotice','encoder','tokenizer'],'Checkpoint terkunci; RoBERTa requires_grad=True, BGE frozen.')
    code(source('src/aes/teacher_score/lesson.py',['export_tokens'])+'''

if RUN_MODELS:
    token_roberta,token_bge=export_tokens(bundle,cfg,journal)
    show(token_roberta,'Token RoBERTa');show(token_bge,'Token BGE-M3')
else:
    print('Pemeriksaan data saja: tokenisasi dan model belum dijalankan.')
''')
    md('''
    ## 7 · RoBERTa pretrained → embedding Excel

    `last_hidden_state[:,0,:]` adalah CLS tiapwindow; mean antarwindow menjadi fitur jawaban. **Output:**seluruhjawaban×768 dan setiapCLSwindow→`05_roberta_pretrained.xlsx`. Ini SEBELUM fine-tuning; bukan nilai siswa dan bukan similarity yang otomatis terkalibrasi. Semua koordinat tersimpan, preview hanya sedikitkolom.
    ''')
    code(source('src/aes/teacher_score/lesson.py',['roberta_pretrained_vectors'])+'''

if RUN_MODELS:
    roberta_before=roberta_pretrained_vectors(bundle,cfg,journal)
    print('RoBERTa: setiap jawaban diubah menjadi vektor berisi 768 angka, sebelum pelatihan dengan nilai guru.')
    show(vector_frame(roberta_before,[r['row_id'] for r in bundle.rows]),'Vektor jawaban RoBERTa sebelum pelatihan')
''')
    md('''
    ## 8 · BGE-M3 frozen → embedding Excel

    **Output:**seluruhjawaban×1024→`06_bge_frozen.xlsx`. Dense CLS+L2 mengikuti p14§3.2, `e=norm(H[0])` (PDF gabungan3135). Vektor belum menjadi skor; head belajar memetakannya. Tidak memakai sparse/ColBERT atau melatih contrastive retrieval ulang.
    ''')
    code(source('src/aes/teacher_score/models.py',['bge_features'])+'''

if RUN_MODELS:
    bge=bge_features(bundle,cfg,output)
    journal.save('06_bge_frozen',{'Vektor_jawaban_1024':vector_frame(bge.numpy(),[r['row_id'] for r in bundle.rows]),
                                'Encoding':pd.DataFrame([json.loads((output/'bge_encoding.json').read_text())])},
                 'BGE dense frozen; semua koordinat',source=cfg['bge'],pooling='CLS then L2')
    print('BGE-M3: setiap jawaban diubah menjadi vektor berisi 1024 angka. Bobot encoder tetap selama pelatihan.')
    show(vector_frame(bge.numpy(),[r['row_id'] for r in bundle.rows]),'Vektor jawaban BGE-M3')
''')
    md('''
    ## 9 · Vektor token/kata/kalimat satu jawaban → Excel

    Di tahap ini kita melihat satu jawaban contoh dengan kedua model. **Vektor sebelum konteks** berasal dari tabel embedding model. **Vektor setelah konteks** sudah dipengaruhi posisi dan kata di sekitarnya. Untuk satu kata yang terpecah menjadi beberapa token, vektornya dirata-ratakan dari potongan tersebut. Token yang terbaca di beberapa window dirata-ratakan dahulu.

    Setiap baris tabel pertama adalah satu kata; hanya lima komponen vektor ditampilkan di layar. Seluruh angka tersedia di **03_Vektor_Kata_dan_Kalimat.xlsx**. Tabel berikutnya membandingkan dua kemunculan kata yang sama di model yang sama, dengan cara membaca di atas tabel.

    Kalimat dibaca secara terpisah dan diwakili token CLS. Kata di luar batas input ditandai tidak tercakup; sel kosong tidak berarti vektor nol. Pemisahan kalimat memakai tanda baca sederhana sehingga bisa keliru pada singkatan. Jangan membandingkan cosine langsung antara ruang RoBERTa 768 angka dan BGE-M3 1024 angka.
    ''')
    code("WORD = re.compile(r\"\\b\\w+(?:[-’']\\w+)*\\b\", re.UNICODE)\n"+source('src/aes/teacher_score/vectors.py',['sentence_spans','token_vectors','word_vectors','cosine','compare_repetitions','vector_lab'])+'\n'+source('src/aes/teacher_score/lesson.py',['export_word_sentence_vectors'])+'''

if RUN_WORD_VECTORS:
    before_words=vector_lab(trace_text,{**cfg,'vector_source':'pretrained'},output)
    export_word_sentence_vectors(before_words,journal,'07_vektor_kata_kalimat_pretrained')
    for role in ('roberta','bge'):
        word_rows=before_words[role]['words']
        preview=pd.DataFrame([{'word_index':w['word_index'],'word':w['word'],'covered':w['covered'],
                              **{f'v_{i:04d}':w['contextual_first5'][i] if w['contextual_first5'] is not None else None for i in range(5)}} for w in word_rows])
        show(preview,('RoBERTa' if role=='roberta' else 'BGE-M3')+' — contoh vektor kata')
        pairs=list(repeated_word_rows(before_words[role]['repeated_words'],role))
        show(pd.DataFrame(pairs,columns=REPEATED_COLUMNS),
             'Perbandingan kata yang berulang',columns=9)
''')
    md('''
    ## 10 · Model S1/S2/S3 dan cara menghasilkan skor

    **S1** memakai 768 angka RoBERTa untuk menghasilkan satu skor. **S2** memakai 1024 angka BGE-M3, lalu MLP, yaitu jaringan kecil untuk mempelajari pemetaan vektor menjadi skor. **S3** memberi LayerNorm tambahan pada vektor RoBERTa, memadatkan BGE menjadi 256 angka dengan proyeksi dan LayerNorm, menyambung kedua cabang, lalu menghasilkan satu skor. LayerNorm adalah operasi umum, bukan khusus BGE; tambahan ini hanya dipasang sebelum fusion S3. Vektor RoBERTa asli dan setelah LayerNorm sama-sama disimpan.

    **Encoder** adalah bagian yang membaca teks dan menghasilkan vektor. **Head**, atau lapisan prediksi, mengubah vektor menjadi skor. Fine-tuned berarti bobot encoder ikut diperbarui saat latihan; frozen berarti bobot encoder tetap. Fungsi `forward` di cell benar-benar dipakai saat pelatihan. Lapisan prediksi yang baru dibuat belum menghasilkan skor yang layak digunakan.

    p10 mendukung konsep menggabungkan fitur, tetapi memakai RoBERTa+fiturmanual, bukan replikasi RoBERTa+BGE ini. BGE dalam unggahan fullyfine-tuned; di sini S2/S3frozen sesuai rancangan sebelumnya, head/proyeksi tetapbelajar. **Output spesifikasi:**`08_arsitektur.xlsx`, jelas diberi statusDESIGN, bukan angka hasilmodel.
    ''')
    defs('src/aes/teacher_score/models.py',['SingleScoreModel','Engine'],'Satu head satu nilai; Engine menyiapkan tensor dan mengembalikan satuan0–100.')
    defs('src/aes/teacher_score/experiment.py',['seed_everything','make_model'],'Model baru pada setiapfold/refit, bukan meneruskan bobot foldlain.')
    code('''
    specifications=pd.DataFrame([
        {'scenario':'S1','encoder':'RoBERTa: bobot ikut dilatih','vector':768,'head':'Linear: 768 angka → 1 skor','output':1},
        {'scenario':'S2','encoder':'BGE-M3: bobot tetap','vector':1024,'head':'MLP: 1024 → 256 → 1 skor','output':1},
        {'scenario':'S3','encoder':'RoBERTa ikut dilatih; BGE-M3 tetap','vector':1024,'head':'RoBERTa 768 → LayerNorm; BGE → proyeksi 256 + LayerNorm; concat 1024 → 1 skor','output':1}])
    journal.save('08_arsitektur',{'Spesifikasi':specifications},'Spesifikasi model, bukan prediksi',status='DESIGN / CONFIGURATION')
    show(specifications,'Arsitektur dan bagian yang belajar')
    ''')
    md('''
    ## 11 · Empat metrik, loss, optimizer dan batch

    RMSE/Pearson memakai prediksi mentah0–100. QWK/AC2 kuadratik memakai salinan clip0–100 lalu integerterdekat(tie.5kebawah),101kategori dipertahankan. Undefined=kosong+alasan. AC2 sensitif kategoridan bisa tinggi saat QWK buruk; baca bersama metriklain(m06). Rumus AC2 juga diverifikasi dengan sourceirrCAC; bab rumus tidak tersedia dalam cuplikan bukum08.

    `fit` di bawah adalah fungsi latihan sebenarnya. **Learning rate kandidat A:** RoBERTa 0,00001 (`1e-5`), head dan proyeksi 0,0001 (`1e-4`). LayerNorm tambahan yang baru dibuat ikut kelompok head. Encoder BGE frozen sehingga tidak diperbarui. AdamW memakai MSE, warmup 10%, clip gradien 1; scheduler menurunkan learning rate setelah warmup. Microbatch 1 diakumulasi menjadi efektif 8. Target 0–1 hanya untuk loss, RMSE tetap poin asli. **Output sesudah training:** riwayat epoch, loss, metrik, epoch pilihan dan alasan berhenti → Excel tiap fold/refit.

    **Early stopping:** maksimal 25 epoch untuk semua skenario. Setiap RMSE validasi turun, bobot disimpan dan hitungan tanpa perbaikan kembali nol. Setelah lima epoch berturut-turut tidak membaik, fold dihentikan dan bobot terbaik dipulihkan. Ini tidak menggunakan data testing dan tidak menjamin fusion akan selalu menang.

    Saat berjalan, **epoch** berarti satu putaran data latih. **Loss latih (MSE, target 0–1)** adalah kesalahan yang diperkecil saat model belajar, dengan nilai guru dibagi 100. **RMSE validasi** sudah dalam poin skor asli 0–100. Keduanya memakai skala berbeda, sehingga angkanya tidak dibandingkan langsung.
    ''')
    defs('src/aes/metrics.py',['gwet_ac2'],'Agreement kuadratik Gwet, bukan sekadar korelasi.')
    defs('src/aes/teacher_score/models.py',['metrics'],'Empat metrik tanpa menimpa prediksi mentah.')
    defs('src/aes/teacher_score/experiment.py',['optimizer_groups','save_training_progress','load_training_progress','clear_training_progress','fit'],'backward menghitung gradien; optimizer.step memperbarui parameter; checkpoint otomatis menjaga progres setiap epoch.')
    md('''
    ## 12 · TrainingS1→S2→S3, refit, finaltest

    CV development memilih epoch minimumRMSE. Refit memakai ceil median epoch5fold. Semua konfigurasi/refit selesai sebelumtest, tanpa earlystopping/tuning ditest. p08 juga membedakan fungsi validation dantesting; bentuk5fold ini keputusan proyek, bukan menyalin split70/15/15paper. Fold yang berbagi data tidak independen(m09).

    Sesudah CV S3 selesai, alur kembali ke S1 untuk **pelatihan model final**, lalu S2 dan S3. Model dibuat ulang dari pretrained dan dilatih pada seluruh 80% data latih. Jumlah epoch mengikuti median pilihan CV (dibulatkan ke atas), bukan wajib 25. Tidak ada early stopping pada tahap final karena tidak ada fold validasi terpisah. Setelah semua model final terkunci, barulah 20% testing dievaluasi.

    **Checkpoint otomatis:** tiap epoch lengkap menyimpan bobot saat ini dan bobot terbaik, optimizer, scheduler, GradScaler, riwayat, urutan acak dan penghitung early stopping. Saat dijalankan lagi, kode memeriksa data/fold/pengaturan/perangkat sebelum melanjutkan epoch berikutnya. Dua generasi terakhir disimpan selama tahap aktif; bila salinan terbaru tidak utuh, cadangan sebelumnya diperiksa. Setelah fold atau refit selesai disimpan, checkpoint sementara tahap itu dibersihkan. Model final tetap disimpan. Fold dan refit yang sudah selesai tidak dilatih ulang.

    Versi ini mendapat folder hasil baru berdasarkan kode dan pengaturan; hasil lama tidak ditimpa atau dipakai sebagai hasil training revisi. Untuk melanjutkan, gunakan notebook, dataset, konfigurasi, versi pustaka dan perangkat GPU yang sama. Karena hasil testing versi lama sudah dilihat, pemilihan pengaturan selanjutnya tetap memakai CV; hasil testing lanjutan bukan bukti konfirmasi yang sepenuhnya belum pernah disentuh dalam pengembangan.

    **Output:**`09_S1_fold_1.xlsx` sampai`09_S3_fold_5.xlsx`(riwayat setiapepoch,prediksi/metrikvalidation), `10_S1_refit.xlsx`dst., `11_S1_test.xlsx`dst., `12_evaluasi_tiga_skenario.xlsx`. Tiapfold selesai langsung disimpan. Resume hasil cocok sumber/config/code memakai test tersimpan.
    ''')
    defs('src/aes/teacher_score/experiment.py',['save_model','load_model','train_all'],'CV semua skenario, refitsemua, lalu finaltest dengan checkpoint terkunci.')
    code(source('src/aes/teacher_score/lesson.py',['ordinal_copy','prediction_frame','export_training_event','test_score_comparison','export_completed_training'])+'''

if RUN_TRAINING:
    if not RUN_MODELS:raise RuntimeError('Aktifkan RUN_MODELS sebelumtraining.')
    def on_training_stage(kind,scenario,result):
        export_training_event(kind,scenario,result,bundle,cfg,journal)
    report=train_all(bundle,manifest,bge,cfg,output,observer=on_training_stage)
    export_completed_training(report,bundle,cfg,output,journal)
    show(pd.DataFrame(report['cv']),'Hasil validasi silang (lima fold)',rows=15)
    final_table=pd.DataFrame(report['rows'])[['Skenario','RMSE','Pearson','QWK','AC2']]
    show(final_table,'Perbandingan model pada data uji',rows=3)
    print('Perbandingan skor: nilai guru asli, prediksi mentah skala 0–100, dan salinan yang dibulatkan. '
          'RMSE/Pearson memakai prediksi mentah; QWK/AC2 memakai salinan kategori. Tepat .5 dibulatkan ke bawah.')
    # Semua jawaban uji memakai nilai guru asli; tidak menjalankan prediksi ulang.
    comparison=test_score_comparison(report,bundle,cfg,output)
    show(comparison,'Nilai guru, prediksi mentah, dan pembulatan pada data uji',rows=len(comparison),columns=8)
else:
    print('Pelatihan belum dijalankan. Prediksi dan evaluasi belum tersedia.')
''')
    md('''
    ## 13 · Prediksi satu jawaban dengan3model final

    Sama seperti akhir fileunggahan: tampilkan prediksi. `NEW_ANSWER`kosong memakai contohheld-out; teksbaru tidak memerlukan nilaiguru dan tidak melatihulang. **Output:**3skormentah,kategorisalinan,out-of-range→`13_prediksi_satu_jawaban.xlsx`.
    ''')
    code(source('src/aes/teacher_score/experiment.py',['score_answer'])+'''

if RUN_TRAINING:
    prediction_rows=score_answer(trace_text,cfg,output)
    # Contoh dari data uji mempunyai nilai guru. Jawaban baru belum mempunyai nilai guru.
    for row in prediction_rows:
        row['skor_guru']=None if cfg.get('new_answer','').strip() else float(bundle.y[trace_index])
    if cfg.get('new_answer','').strip():
        print('Jawaban baru: nilai guru belum tersedia, sehingga kolom Nilai Guru kosong.')
    journal.save('13_prediksi_satu_jawaban',{'Input_privat':pd.DataFrame([{'text':trace_text}]),
                                           'Tiga_skenario':pd.DataFrame(prediction_rows)},'Prediksi checkpointfinal; bukan trainingulang')
    show(pd.DataFrame(prediction_rows)[['Skenario','skor_guru','prediksi_mentah','kategori_salinan','di_luar_rentang']],
         'Nilai guru dan prediksi satu jawaban (skala 0–100)',columns=5)
''')
    md('''
    ## 14 · Fitur SESUDAHtraining dan outputfusion → Excel

    Setelah finaltest, ekstrak fitur modeleval:RoBERTa768,BGE1024,proyeksi256sebelum/sesudahLayerNorm,concat1024,hiddenMLP/GELU/dropouteval,outputhead. Hanya komponen yang ada pada skenario diekspor. Dropout nonaktif saatinspeksi, berbeda daritraining.

    **Output:**`14_S1_fitur_setelah_training.xlsx`,S2,S3. Cocokkan`row_id`dengan05untuk membandingkan sebelum/sesudah. BGE tetapfrozen. Semua jawaban untukinspeksi, bukan evaluasitestbaru. Vektor kata/kalimat contohsesudahfine-tuning→`15_S1_...xlsx`dan`15_S3_...xlsx`.
    ''')
    code(source('src/aes/teacher_score/lesson.py',['export_final_features'])+'''

if RUN_TRAINING and RUN_FINAL_FEATURES:
    print('Mengambil vektor jawaban dari ketiga model yang sudah dilatih...')
    export_final_features(bundle,bge,cfg,output,journal)
    print('Vektor',len(bundle.rows),'jawaban siap disajikan dalam 02_Vektor_Jawaban.xlsx.')
    if RUN_WORD_VECTORS:
        for scenario in ('S1','S3'):
            after_words=vector_lab(trace_text,{**cfg,'vector_source':scenario},output)
            export_word_sentence_vectors(after_words,journal,f'15_{scenario}_vektor_kata_setelah_training')
''')
    md('''
    ## 15 · Unduh hasil Excel yang sudah rapi
    ''')
    report_code=source('src/aes/teacher_score/presentation.py',[
        'Table','SavedOutputs','guide','write_report','data_report','answer_vectors_report',
        'word_vectors_report','evaluation_report','training_report','build_reports'])
    # Daftar pustaka sudah menjadi Markdown, bukan sumber runtime atau sheet hasil.
    start=report_code.index("    if source.has('01_rujukan_paper.xlsx','Rujukan'):")
    end=report_code.index('\n    return sheets',start)
    report_code=report_code[:start]+report_code[end:]
    code('# Susun lima laporan dengan nama kolom yang bisa dibaca dan seluruh komponen vektor utama.\n'+report_code)
    code('''
    # Susun laporan dari hasil sesi ini. Tidak memanggil training atau evaluasi lagi.
    aes_cell('', """
    summary=build_reports(journal.folder,output/'hasil_ringkas',verbose=False)
    dump(ROOT/'outputs/dataset300_download.json',{
        'archive':summary['archive'],'run_folder':str(output),
        'excel_count':len(summary['files']),'files':[Path(p).name for p in summary['files']]})
    """)

    # Unduh paket yang baru dibuat, bukan ZIP rincian teknis atau hasil sesi lain.
    import zipfile
    download=json.loads((ROOT/'outputs/dataset300_download.json').read_text())
    with zipfile.ZipFile(download['archive']) as package:
        names=package.namelist()
    if sorted(names)!=sorted(download['files']):
        raise RuntimeError('Isi ZIP tidak cocok dengan daftar laporan. Unduhan dihentikan.')
    print('Hasil siap diunduh:',Path(download['archive']).name)
    print('Isi paket:',len(names),'file Excel yang sudah dirapikan:')
    descriptions={
        '01_Data_dan_Tokenisasi.xlsx':'Jawaban, nilai guru, pembagian data, dan token kedua model.',
        '02_Vektor_Jawaban.xlsx':'Embedding jawaban RoBERTa, BGE-M3, dan fusion.',
        '03_Vektor_Kata_dan_Kalimat.xlsx':'Vektor kata/kalimat dan perbandingan kata yang berulang.',
        '04_Prediksi_dan_Evaluasi.xlsx':'Nilai guru, prediksi, RMSE, Pearson, QWK, dan Gwet AC2.',
        '05_Riwayat_Pelatihan.xlsx':'Seluruh fold dan epoch digabung per model.'}
    for name in names:
        print('•',name)
        print('  ',descriptions.get(name,''))
    if len(names)<5:
        print('Paket belum lengkap: hanya tahap yang sudah selesai yang disertakan.')
    if '04_Prediksi_dan_Evaluasi.xlsx' in names:
        print('Mulai membaca hasil akhir dari 04_Prediksi_dan_Evaluasi.xlsx → Perbandingan Model.')
    if IN_COLAB:
        from google.colab import files
        files.download(download['archive'])
    else:print('Lokasi file:',download['archive'])
    ''',False)
    # Hash dari kode yang dapat dieksekusi; penjelasan Markdown tidak mengubah identitas hasil.
    normalized=[ast.dump(ast.parse(c['source'].removeprefix('%%aes\n')),include_attributes=False)
                for c in cells if c['cell_type']=='code']
    fingerprint=hashlib.sha256('\n'.join(normalized).encode()).hexdigest()
    for cell in cells:
        cell['source']=cell['source'].replace('__NOTEBOOK_CODE_SHA256__',fingerprint)
        if cell['cell_type']=='code':
            tree=ast.parse(cell['source'].removeprefix('%%aes\n'))
            for node in ast.walk(tree):
                if isinstance(node,ast.ImportFrom) and node.module and (node.module=='aes' or node.module.startswith('aes.')):
                    raise ValueError('Notebook masih mengimpor kode proyek: '+node.module)
            if '__file__' in cell['source'] or "['git'," in cell['source']:
                raise ValueError('Notebook masih bergantung pada file kode atau Git.')
    nb=dict(cells=cells,metadata={'kernelspec':{'display_name':'Python3','language':'python','name':'python3'},
             'language_info':{'name':'python'},'colab':{'name':'03_AES_dataset_guru_300.ipynb'},
             'aes_standalone':{'code_sha256':fingerprint,'repository_required':False}},nbformat=4,nbformat_minor=4)
    content=json.dumps(nb,ensure_ascii=False,indent=1)+'\n'
    for directory in ('notebooks','aes-learning/notebooks'):
        (ROOT/directory/'03_AES_dataset_guru_300.ipynb').write_text(content,encoding='utf-8')
    return nb


if __name__=='__main__':
    print('Notebookinline:',len(build()['cells']),'cell;2mirror identik tanpaoutput.')
