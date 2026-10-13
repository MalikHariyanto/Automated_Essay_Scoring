"""Workbook review privat: usulan AI tidak menjadi score_* guru."""
import csv,hashlib,json
from pathlib import Path
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table,TableStyleInfo
from openpyxl.comments import Comment

import argparse
parser=argparse.ArgumentParser(description="Ekspor review privat yang SUDAH tersedia; tidak menghasilkan label guru atau menilai esai otomatis.")
parser.add_argument('--csv',type=Path,required=True)
parser.add_argument('--rubric',type=Path,required=True)
parser.add_argument('--review',type=Path,required=True)
parser.add_argument('--output',type=Path,required=True)
args=parser.parse_args()
csv_path=args.csv;rub_path=args.rubric
transfer=json.loads(args.review.read_text(encoding='utf-8'))
R={}
for essay_id,trait,score,reason,higher,quotes,review in transfer['reviews']:
 if (essay_id,trait) in R:raise ValueError('Review trait duplikat.')
 R[(essay_id,trait)]={'score':score,'reason':reason,'higher':higher,'quotes':quotes,'review':review}
source_bytes=csv_path.read_bytes();rub_bytes=rub_path.read_bytes()
rows=list(csv.DictReader(source_bytes.decode('utf-8-sig').splitlines(keepends=True)))
rubric=json.loads(rub_bytes);traits=[t['name'] for t in rubric['traits']]
assert traits==['isi','organisasi','kebahasaan','mekanik']
assert len(R)==len(rows)*4
hashes={'csv':hashlib.sha256(source_bytes).hexdigest(),'rubric':hashlib.sha256(rub_bytes).hexdigest()}
if hashes != transfer['source_hashes']:
 raise ValueError('CSV/JSON berbeda dari sumber review. Jangan memasangkan usulan ini dengan versi data lain.')
if len({row['essay_id'] for row in rows}) != len(rows):raise ValueError('essay_id duplikat.')
if set(R) != {(row['essay_id'],trait) for row in rows for trait in traits}:raise ValueError('Mapping esai/trait tidak cocok.')
now=transfer['timestamp_bangkok']
deferred_ids=[row['essay_id'] for row in rows if all(R[(row['essay_id'],trait)]['score'] is None for trait in traits)]
numeric_count=sum(value['score'] is not None for value in R.values())
continuations=transfer.get('ai_continuations',{})
scopes=transfer.get('assessment_scopes',{})
if set(continuations)-{row['essay_id'] for row in rows}:raise ValueError('ID rekonstruksi tidak ada dalam sumber.')
review_texts={row['essay_id']:row['text_raw']+continuations.get(row['essay_id'],{}).get('added_text','') for row in rows}
for row in rows:
 essay_id=row['essay_id']
 if row.get('text_completeness')=='continuation_unassigned' and any(R[(essay_id,t)]['score'] is not None for t in traits):
  if essay_id not in continuations:raise ValueError('Teks tidak lengkap: skor hanya boleh pada versi rekonstruksi berlabel, bukan teks asli.')
 if R[(essay_id,'mekanik')]['score'] is not None and row.get('transcription_status')!='verified':
  if not scopes.get(essay_id,{}).get('mechanic_transcription_simulation'):
   raise ValueError('Skor mekanik pada transkripsi draf memerlukan label simulasi eksplisit.')
reviews=[]
for row in rows:
 for trait in rubric['traits']:
  a=R[(row['essay_id'],trait['name'])]
  assert a['score'] is None or a['score'] in trait['categories']
  quotes=[]
  for quote in a['quotes']:
   start=review_texts[row['essay_id']].find(quote)
   assert start>=0,(row['essay_id'],trait['name'],quote)
   paragraph=review_texts[row['essay_id']][:start].count('\n\n')+1
   quotes.append({'quote':quote,'paragraph':paragraph,'start_char_zero_based':start,'end_char_exclusive':start+len(quote)})
  reviews.append({'essay_id':row['essay_id'],'trait':trait['name'],'ai_provisional_score':a['score'],
     'assessment_source':'AI_ASSISTANT_NOT_TEACHER','status':'provisional_for_teacher_review' if a['score'] is not None else 'deferred_requires_source_review',
     'rubric_level_original':trait['levels'][str(a['score'])] if a['score'] is not None else rubric['missing_rule'],
     'reason':a['reason'],'higher_level_or_defer_reason':a['higher'],'evidence':quotes,'teacher_next_step':a['review'],
     'assessment_scope':scopes.get(row['essay_id'],{}),
     'evidence_text_source':'AI_RECONSTRUCTION' if row['essay_id'] in continuations else 'ORIGINAL_DRAFT_TRANSCRIPTION'})
report={'assessment_source':'AI_ASSISTANT_NOT_TEACHER','teacher_validated':False,'timestamp_bangkok':now,
        'rubric_version':rubric['version'],'rubric_confirmed_original':rubric['confirmed'],'source_hashes':hashes,
        'teacher_score_columns_unchanged':True,'original_rows_preserved':len(rows),'reviews':reviews}


wb=Workbook();guide=wb.active;guide.title='Panduan'
green='166758';blue='E8F3FA';amber='FFF0DB';grey='F1F4F2';white='FFFFFF';ink='183B34'
header_fill=PatternFill('solid',fgColor=green)
line=Side(style='thin',color='DCE5DF')
def put(ws,row):
 ws.append(row)
 for c in ws[ws.max_row]:
  if isinstance(c.value,str):c.data_type='s' # Teks/provenance tidak dieksekusi sebagai formula.
  c.alignment=Alignment(vertical='top',wrap_text=True)
  c.font=Font(name='Calibri',size=11,color=ink)
  c.border=Border(bottom=line)
def finish(ws,widths,table=None,height=56):
 ws.freeze_panes='C2';ws.sheet_view.showGridLines=False
 for cell in ws[1]:cell.fill=header_fill;cell.font=Font(name='Calibri',bold=True,color=white,size=11);cell.alignment=Alignment(wrap_text=True,vertical='center')
 ws.row_dimensions[1].height=38
 for index,width in enumerate(widths,1):ws.column_dimensions[__import__('openpyxl').utils.get_column_letter(index)].width=width
 for n in range(2,ws.max_row+1):
  ws.row_dimensions[n].height=height
  if n%2==0:
   for c in ws[n]:c.fill=PatternFill('solid',fgColor='F6F9F7')
 ws.auto_filter.ref=ws.dimensions
 ws.sheet_properties.pageSetUpPr.fitToPage=True
 ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A4;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
 ws.print_title_rows='1:1'
 if table:
  t=Table(displayName=table,ref=ws.dimensions);t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium4',showRowStripes=True);ws.add_table(t)
def add_score_validation(ws,columns):
 for column in columns:
  dv=DataValidation(type='whole',operator='between',formula1=1,formula2=5,allow_blank=True)
  dv.errorTitle='Skor tidak sah';dv.error='Gunakan satu integer 1–5. Jika belum dinilai, biarkan kosong.';dv.showErrorMessage=True;dv.errorStyle='stop'
  dv.promptTitle='Keputusan guru';dv.prompt='Isi setelah guru menilai; jangan salin usulan otomatis.';dv.showInputMessage=True
  ws.add_data_validation(dv);dv.add(f'{column}2:{column}{len(rows)+1}')

put(guide,['USULAN PENILAIAN AI — BUKAN SKOR GURU','Keterangan'])
items=[
 ('Tujuan','Membantu guru membaca esai dan menilai dengan rubrik. Usulan bukan ground truth guru dan tidak membuat penelitian siap training.'),
 ('Yang benar-benar dilakukan',f'Seluruh {len(rows)} text_raw dibaca satu per satu. Best fit per trait, bukan mean subindikator. Setiap usulan memiliki kutipan literal dan lokasi paragraf. Nilai huruf lama diabaikan sebagai target.'),
 ('Batas penilaian',f'{numeric_count} usulan angka. Bila transkripsi belum verified, mekanik berangka hanya simulasi pada teks draf. Skor versi rekonstruksi tidak menjadi skor teks asli. ID tanpa usulan angka: {", ".join(deferred_ids)}.'),
 ('Sumber penilai','AI_ASSISTANT_NOT_TEACHER. Tidak ada guru yang diwakili atau pengesahan yang diasumsikan.'),
 ('Waktu usulan (Asia/Bangkok)',now),('Versi rubrik',rubric['version']),
 ('Rubrik asli confirmed',str(rubric['confirmed']).lower()),
 ('scoring_rule ASLI',rubric['scoring_rule']),('missing_rule ASLI',rubric['missing_rule']),
 ('Urutan target','isi, organisasi, kebahasaan, mekanik; ordinal 1–5; tidak ada skor total/holistik.'),
 ('Cara membaca','Ringkasan_AI menampilkan kandidat. Review_trait memuat best fit, batas level, bukti dan tindak lanjut. Nilai kosong berarti ditunda, bukan nol/1.'),
 ('Langkah guru','Periksa naskah/foto dan rubrik; nilai mandiri, lalu bandingkan usulan. Isi Keputusan_guru dengan rater/tanggal/bukti. Guru boleh berbeda dari usulan AI.'),
 ('Cara menghasilkan CSV sesudah validasi','Salin KEPUTUSAN GURU berdasarkan essay_id ke Untuk_CSV_Guru. Perbaiki transkripsi/status hanya dengan bukti. Simpan sheet itu sebagai CSV UTF-8, bukan hanya sheet ringkasan. Notebook sekarang membaca CSV, bukan .xlsx.'),
 ('Syarat Colab penelitian','Empat skor guru valid pada esai yang sama, rubrik disahkan dengan identitas/tanggal guru, teks lengkap/verified, identitas/duplikasi/group diperiksa, review hash terkini dan split sah. Menyalin usulan ke score_* saja tidak memenuhi syarat.'),
 ('Larangan representasi','Jangan menyebut angka AI sebagai skor guru, mengubah confirmed/ready_for_training agar lolos, atau mengubah A/A+ menjadi skor trait. Tidak ada loss/metrik penelitian dihitung dari usulan ini.'),
 ('Data_asli',f'Seluruh kolom CSV dan seluruh {len(rows)} baris disalin tanpa mengisi score_* atau mengubah status. Untuk_CSV_Guru masih merupakan salinan awal yang sama, menunggu guru.'),
 ('Rekonstruksi AI','Rekonstruksi_AI memisahkan teks asli dan tambahan yang diusulkan. Tidak membuktikan halaman hilang atau penulis yang sama. Skor guru tetap kosong.'),
 ('Bukti yang belum ada','Tidak memeriksa foto/naskah asli atau tugas lengkap. Tidak memverifikasi sumber eksternal klaim gizi/harga. Kalimat bertanda masalah dapat berasal dari transkripsi; reviewer harus membedakannya.'),
 ('Kemiripan','Usulan yang sama pada teks mirip tidak membuktikan siswa yang sama atau validitas label. Jangan menggabungkan/menghapus atau menyebar group usulan ke split tanpa review.'),
 ('Hash CSV asli (SHA256)',hashes['csv']),('Hash JSON asli (SHA256)',hashes['rubric']),
 ('Privasi','Workbook mengandung teks/provenance siswa. Berikan ke guru/peneliti yang berwenang dan simpan lokal/Drive privat; jangan commit ke GitHub publik.')]
for item in items:put(guide,list(item))
finish(guide,[37,115],height=65);guide.freeze_panes='B2'

summary=wb.create_sheet('Ringkasan_AI')
put(summary,['essay_id','topic','text_completeness','ai_usul_isi','ai_usul_organisasi','ai_usul_kebahasaan','ai_usul_mekanik','status_usulan','catatan_review'])
for row in rows:
 candidates=[R[(row['essay_id'],t)]['score'] for t in traits]
 note='Usulan AI pada transkripsi draf, bukan skor guru. Mekanik berangka adalah simulasi; bandingkan dengan foto.'
 if row['essay_id'] in continuations:note='SIMULASI VERSI REKONSTRUKSI AI; teks asli masih terpotong dan tidak memperoleh label guru.'
 elif row['text_completeness']=='continuation_unassigned':note='Seluruh skor ditunda: teks asli terpotong dan lanjutan belum diketahui.'
 put(summary,[row['essay_id'],row.get('topic',''),row['text_completeness'],*candidates,'AI_RECONSTRUCTION_SIMULATION' if row['essay_id'] in continuations else 'AI_DRAFT_NOT_TEACHER_PENDING_REVIEW',note])
finish(summary,[12,16,32,16,21,22,20,36,78],table='UsulanAI',height=64)
for r in range(2,len(rows)+2):
 for col in range(4,8):
  summary.cell(r,col).fill=PatternFill('solid',fgColor=amber if summary.cell(r,col).value is None else blue)
  summary.cell(r,col).comment=Comment('Usulan AI, bukan skor guru. Kosong berarti ditunda; jangan isi 0.','Catatan review')

trait_sheet=wb.create_sheet('Review_trait')
put(trait_sheet,['essay_id','trait','ai_usulan_1_5','status','deskriptor_JSON_ASLI','alasan_best_fit_AI','batas_level_atau_penundaan','kutipan_bukti_LITERAL','lokasi_bukti','tindak_lanjut_guru','sumber_teks_bukti','cakupan_AI'])
for item in reviews:
 evidence='\n\n'.join(f"P{x['paragraph']}: {x['quote']}" for x in item['evidence'])
 location='; '.join(f"P{x['paragraph']}; karakter {x['start_char_zero_based']}:{x['end_char_exclusive']} (0-based)" for x in item['evidence'])
 put(trait_sheet,[item['essay_id'],item['trait'],item['ai_provisional_score'],item['status'],item['rubric_level_original'],item['reason'],item['higher_level_or_defer_reason'],evidence,location,item['teacher_next_step'],item['evidence_text_source'],json.dumps(item['assessment_scope'],ensure_ascii=False)])
finish(trait_sheet,[12,19,17,35,65,75,75,65,43,70,35,55],table='BuktiUsulan',height=132)

if continuations:
 reconstruction=wb.create_sheet('Rekonstruksi_AI')
 put(reconstruction,['essay_id','teks_asli_TIDAK_DIUBAH','tambahan_AI_BUKAN_TULISAN_ASLI','teks_versi_rekonstruksi','rujukan_gaya_dari_dataset','status_sumber','penulis_sama_terverifikasi'])
 for row in rows:
  essay_id=row['essay_id']
  if essay_id in continuations:
   item=continuations[essay_id]
   put(reconstruction,[essay_id,row['text_raw'],item['added_text'],review_texts[essay_id],item.get('reference_essay_id',''),
                       'USULAN_REKONSTRUKSI_AI_BUKAN_HALAMAN_ASLI','false'])
 finish(reconstruction,[14,100,90,110,30,65,30],height=300)
 for i in range(2,reconstruction.max_row+1):reconstruction.cell(i,3).fill=PatternFill('solid',fgColor=amber)

teacher=wb.create_sheet('Keputusan_guru')
put(teacher,['essay_id','score_isi','score_organisasi','score_kebahasaan','score_mekanik','rater_id','scoring_date','evidence_isi','evidence_organisasi','evidence_kebahasaan','evidence_mekanik','keputusan_review','catatan_guru'])
for row in rows:put(teacher,[row['essay_id'],None,None,None,None,None,None,None,None,None,None,'belum_direview',None])
finish(teacher,[12,16,21,22,20,26,22,55,55,55,55,34,65],table='KeputusanGuru',height=70)
add_score_validation(teacher,['B','C','D','E'])
dv=DataValidation(type='list',formula1='"belum_direview,dinilai_guru,perlu_perbaikan_transkripsi,tunda_halaman_tidak_lengkap"',allow_blank=False);dv.showErrorMessage=True;teacher.add_data_validation(dv);dv.add(f'L2:L{len(rows)+1}')

headers=list(rows[0])
for title in ['Untuk_CSV_Guru','Data_asli']:
 ws=wb.create_sheet(title);put(ws,headers)
 for row in rows:put(ws,[row[key] for key in headers])
 widths=[12 if h=='essay_id' else 98 if h in ['text_raw','text_clean'] else 18 if h.startswith('score_') else 45 if h.startswith('evidence_') else 30 for h in headers]
 finish(ws,widths,height=225)
 if title=='Untuk_CSV_Guru':
  for col,h in enumerate(headers,1):
   if h.startswith('score_'):
    letter=__import__('openpyxl').utils.get_column_letter(col);add_score_validation(ws,[letter])
    for row in range(2,len(rows)+2):ws.cell(row,col).fill=PatternFill('solid',fgColor='E7F2EC');ws.cell(row,col).comment=Comment('Menunggu KEPUTUSAN GURU. Tidak diisi dari usulan AI.','Review guru')
 ws.freeze_panes='C2'

rub_sheet=wb.create_sheet('Rubrik_asli')
put(rub_sheet,['trait','min','max','scale_type','categories','definition_ASLI','indicators_ASLI','exclude_ASLI','level','descriptor_ASLI'])
for t in rubric['traits']:
 for level in t['categories']:
  put(rub_sheet,[t['name'],t['min'],t['max'],t['scale_type'],json.dumps(t['categories']),t['definition'],'\n'.join(t['indicators']),t['exclude'],level,t['levels'][str(level)]])
finish(rub_sheet,[18,8,8,15,22,70,60,60,10,100],table='RubrikAsli',height=120)
raw=wb.create_sheet('JSON_asli');put(raw,['Keterangan','Isi ASLI'])
for item in [('Nama file',rub_path.name),('SHA256',hashes['rubric']),('JSON byte UTF-8 dibaca tanpa mengubah isi',rub_bytes.decode('utf-8'))]:put(raw,list(item))
finish(raw,[50,150],height=70);raw.row_dimensions[4].height=300

# Audit disalin dari laporan privat, tidak menghitung metrik penelitian.
audit_sheet=wb.create_sheet('Audit')
put(audit_sheet,['Aspek','Hasil aktual / kebutuhan'])
for row in transfer['audit_rows']:put(audit_sheet,row)
finish(audit_sheet,[40,130],height=65)
for row in range(2,audit_sheet.max_row+1):
 if audit_sheet.cell(row,1).value in ['statuses','near_duplicates']:audit_sheet.row_dimensions[row].height=225

if transfer.get('reference_access'):
 refs=wb.create_sheet('Akses_rujukan')
 put(refs,['rujukan','URL_yang_dicoba','hasil_akses_aktual','status_verifikasi_fakta'])
 for item in transfer['reference_access']:
  put(refs,[item['key'],item['url'],item.get('error',str(item.get('status','unknown'))),
            'BELUM_DIVERIFIKASI' if 'error' in item else 'Akses tercatat; baca dan tinjau isi sumber sebelum verifikasi fakta'])
 finish(refs,[30,90,80,70],height=85)

path=args.output
path.parent.mkdir(parents=True,exist_ok=True)
wb.active=wb.sheetnames.index('Ringkasan_AI')
wb.save(path)
# Periksa kembali nilai, urutan, keutuhan sumber, dan semua kutipan.
check=load_workbook(path,data_only=False)
assert check['Data_asli'].max_row==len(rows)+1 and check['Review_trait'].max_row==len(reviews)+1
for i,row in enumerate(rows,2):
 for j,key in enumerate(headers,1):assert (check['Data_asli'].cell(i,j).value or '')==row[key],(i,key)
 for title in ['Data_asli','Untuk_CSV_Guru']:
  for t in traits:assert check[title].cell(i,headers.index('score_'+t)+1).value is None
 for j in range(2,6):assert check['Keputusan_guru'].cell(i,j).value is None
assert json.loads(check['JSON_asli']['B4'].value)==rubric
assert sum(item['ai_provisional_score'] is not None for item in reviews)==numeric_count
assert all(item['ai_provisional_score'] is None for item in reviews if item['essay_id'] in deferred_ids)
assert csv_path.read_bytes()==source_bytes and rub_path.read_bytes()==rub_bytes
# String formulas tidak muncul dari teks/provenance unggahan.
assert not any(c.data_type=='f' for ws in check for row in ws for c in row)
proof={'workbook':path.name,'workbook_sha256':hashlib.sha256(path.read_bytes()).hexdigest(),
 'source_hashes':hashes,'original_sources_unchanged':True,'essays':len(rows),'trait_reviews':len(reviews),
 'provisional_ai_numeric_candidates':numeric_count,'teacher_labels_written':0,'mechanic_scores_deferred':sum(R[(row['essay_id'],'mekanik')]['score'] is None for row in rows),
 'all_traits_deferred_ids':deferred_ids,'reconstruction_ids':list(continuations),
 'mechanic_transcription_simulation_ids':[key for key,value in scopes.items() if value.get('mechanic_transcription_simulation')],'literal_evidence_all_verified':True,'rubric_original_roundtrip':True,
 'research_training_or_evaluation_run':False,'public_repository_upload':False,'sheets':check.sheetnames}
path.with_suffix('.verification.json').write_text(json.dumps(proof,ensure_ascii=False,indent=2))
print(json.dumps(proof,ensure_ascii=False,indent=2))
