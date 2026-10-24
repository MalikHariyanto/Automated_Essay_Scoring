"""Panduan belajar dari potongan source E6, tanpa data atau hasil siswa."""
import hashlib
import json
from pathlib import Path
import subprocess
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Alignment,Font,PatternFill,Border,Side
from openpyxl.worksheet.table import Table,TableStyleInfo
from openpyxl.utils import get_column_letter

ROOT=Path(__file__).resolve().parents[1]
SOURCE=ROOT/'notebooks/experiments/AES_300_E6_S3_LR5e6_CPU_Reuse.ipynb'
TARGET=ROOT/'downloads/Panduan_Belajar_Kode_E6.xlsx'
PIN='2257cb0f601148764d7844fc6ffca925a9b2ab95'
URL=f'https://github.com/MalikHariyanto/Automated_Essay_Scoring/blob/{PIN}/notebooks/experiments/AES_300_E6_S3_LR5e6_CPU_Reuse.ipynb'
# index cell, teks pencarian, nama bagian, fungsi singkat, panjang snippet, kemunculan.
ROWS = [
    (12, 'NEW_ANSWER=', 'Jawaban baru', 'Menentukan teks baru; kosong memakai contoh held-out.', 1, 1),
    (3, 'def prepare_persistent_storage', 'Penyimpanan Google Drive', 'Menyiapkan folder data, checkpoint, dan hasil privat.', 1, 1),
    (5, 'def prepare_environment', 'Lingkungan Python', 'Memeriksa dan menyiapkan pustaka untuk menjalankan AES.', 1, 1),
    (6, 'WORKER_SOURCE', 'Worker Python', 'Menyiapkan proses Python terpisah untuk perhitungan model.', 1, 1),
    (7, 'ready=prepare_environment', 'Menjalankan setup', 'Menjalankan pemeriksaan lingkungan sebelum model mulai dipakai.', 1, 1),
    (7, 'PERSISTENT_ROOT=', 'Menentukan folder hasil', 'Menghubungkan penyimpanan hasil dengan folder Drive privat.', 1, 1),
    (7, 'DATASET_NAME=', 'Nama dataset', 'Menentukan nama file Excel yang akan dibaca.', 1, 1),
    (7, 'files.upload()', 'Upload dataset', 'Meminta dataset ketika file belum tersedia di Drive.', 1, 1),
    (7, "'AES300_REQUIRE_GPU':'0'", 'CPU tanpa kuota GPU', 'Menonaktifkan kewajiban GPU pada proses perhitungan AES.', 1, 1),
    (7, 'def aes_cell(', 'Penghubung %%aes', 'Meneruskan isi cell ke worker Python yang sama.', 1, 1),
    (9, 'from transformers import', 'Pustaka model', 'Mengimpor pemuat model, tokenizer, serta scheduler training.', 1, 1),
    (10, 'DEFAULTS =', 'Pengaturan bawaan', 'Menyediakan nilai awal konfigurasi eksperimen sebelum diubah.', 1, 1),
    (10, 'def config(', 'Validasi pengaturan', 'Menggabungkan konfigurasi dan memeriksa batas parameter eksperimen.', 1, 1),
    (11, 'def check_gpu(', 'Fungsi pemeriksaan perangkat', 'Mendefinisikan pemeriksaan hardware; bukan langkah melatih model.', 1, 1),
    (12, 'ROBERTA_CHECKPOINT=', 'Nama model RoBERTa', 'Menentukan checkpoint pretrained RoBERTa bahasa Indonesia.', 1, 1),
    (12, 'BGE_CHECKPOINT=', 'Nama model BGE-M3', 'Menentukan checkpoint pretrained untuk fitur semantik BGE.', 1, 1),
    (12, 'ROBERTA_REVISION=', 'Versi pretrained', 'Mengunci versi bobot dan tokenizer yang digunakan.', 2, 1),
    (12, 'EPOCHS,FOLDS=', 'Epoch dan jumlah fold', 'Membatasi latihan per fold serta jumlah putaran validasi.', 1, 1),
    (12, 'EARLY_STOPPING_PATIENCE=', 'Early stopping', 'Menetapkan batas epoch tanpa rekor RMSE validasi baru.', 1, 1),
    (12, 'ENCODER_LR=', 'Learning rate S1', 'Menetapkan langkah pembaruan encoder RoBERTa pada S1.', 2, 1),
    (12, 'S3_ENCODER_LR=', 'Learning rate S3', 'Memperkecil langkah encoder RoBERTa S3 menjadi 5e-6.', 1, 1),
    (12, 'PROJECTION_LR=HEAD_LR=', 'Learning rate head/proyeksi', 'Menetapkan langkah lapisan pemetaan fitur menjadi skor.', 1, 1),
    (12, 'MICRO_BATCH,EFFECTIVE_BATCH=', 'Batch dan akumulasi', 'Mengolah satu jawaban, mengakumulasi gradien hingga delapan.', 1, 1),
    (12, "device='cpu'", 'Memilih CPU', 'Memastikan model dan training memakai perangkat CPU.', 1, 1),
    (14, 'def clean(', 'Definisi pembersihan teks', 'Menyiapkan pembersihan Unicode tanpa memperbaiki ejaan siswa.', 1, 1),
    (14, 'class Bundle', 'Wadah data: bundle', 'Menyatukan baris, target, kelompok, dan bukti sumber.', 7, 1),
    (14, 'def read_dataset(', 'Definisi pembaca dataset', 'Menyiapkan pembacaan serta pemeriksaan kolom dan nilai.', 1, 1),
    (14, 'y = np.asarray', 'Target atau y', 'Mengambil skor guru sebagai target numerik supervised learning.', 1, 1),
    (14, 'return Bundle(', 'Keluaran pembaca dataset', 'Mengembalikan teks, target, kelompok, serta metadata sumber.', 1, 1),
    (15, 'def write_excel(', 'Penulis Excel', 'Mendefinisikan penyimpanan tabel hasil ke workbook Excel.', 1, 1),
    (15, 'class ExcelJournal', 'Pencatat hasil tahap', 'Mencatat keluaran setiap tahap beserta informasi sumbernya.', 1, 1),
    (16, 'bundle=read_dataset(', 'Membaca dataset sungguhan', 'Memanggil fungsi pembaca dan membentuk bundle aktual.', 1, 1),
    (18, 'GroupShuffleSplit(', 'Pembagian development/test', 'Memisahkan kelompok jawaban sebelum model dipilih melalui validasi.', 2, 1),
    (18, 'GroupKFold(', 'Validasi silang lima fold', 'Membagi development bergantian menjadi data latihan dan validasi.', 2, 1),
    (18, 'manifest=make_split(', 'Menjalankan pembagian data', 'Menghasilkan daftar anggota development, test, dan lima fold.', 1, 1),
    (20, 'def content_windows(', 'Window jawaban saja', 'Menyiapkan potongan token jawaban untuk RoBERTa S1/S3.', 1, 1),
    (20, 'def mean_by_essay(', 'Pooling antar-window', 'Merata-ratakan CLS semua window milik satu jawaban.', 4, 1),
    (21, 'AutoTokenizer.from_pretrained(', 'Memuat tokenizer', 'Memuat aturan tokenisasi mengikuti checkpoint dan revision model.', 1, 1),
    (21, 'AutoModel.from_pretrained(', 'Memuat encoder', 'Memuat objek encoder beserta bobot model pretrained.', 1, 1),
    (21, 'requires_grad_', 'Kebijakan fine-tuning/frozen', 'Menentukan parameter encoder yang boleh menerima pembaruan.', 1, 1),
    (22, 'def export_tokens(', 'Ekspor tokenisasi', 'Mendefinisikan tabel token, ID, window, serta segmen input.', 1, 1),
    (24, 'def roberta_pretrained_vectors(', 'Fitur RoBERTa sebelum latihan', 'Menyiapkan ekstraksi CLS dan fitur 768 komponen.', 1, 1),
    (24, 'last_hidden_state[:, 0, :]', 'Mengambil CLS awal', 'Mengambil token pertama setiap window sebagai representasi.', 1, 1),
    (26, 'def bge_features(', 'Fitur BGE-M3', 'Menyiapkan ekstraksi dan cache vektor BGE yang frozen.', 1, 1),
    (26, 'functional.normalize(', 'Normalisasi L2 BGE', 'Membagi vektor BGE berdasarkan panjangnya sebelum disimpan.', 1, 1),
    (28, 'def word_vectors(', 'Vektor per kata', 'Menggabungkan subtoken menjadi vektor setiap kata contoh.', 1, 1),
    (28, 'def compare_repetitions(', 'Perbandingan kata berulang', 'Membandingkan kemiripan serta perbedaan vektor kata sama.', 1, 1),
    (28, 'def vector_lab(', 'Laboratorium vektor', 'Menyiapkan pemeriksaan token, kata, dan kalimat contoh.', 1, 1),
    (30, 'class SingleScoreModel', 'Arsitektur tiga skenario', 'Mendefinisikan encoder dan head pemberi skor tiap skenario.', 1, 1),
    (30, 'roberta.requires_grad_(True)', 'RoBERTa full fine-tuning', 'Mengizinkan pembaruan bobot encoder RoBERTa S1 dan S3.', 1, 1),
    (30, 'self.dropout =', 'Lapisan dropout', 'Mengacak penonaktifan fitur saat training untuk regularisasi.', 1, 1),
    (30, 'self.head = nn.Linear(768, 1)', 'Head regresi linear S1', 'Memetakan 768 fitur RoBERTa menjadi satu keluaran.', 1, 1),
    (30, 'self.head = nn.Sequential(', 'Head regresi MLP S2', 'Memetakan fitur BGE melalui lapisan nonlinier menjadi skor.', 1, 1),
    (30, 'self.roberta_norm =', 'LayerNorm RoBERTa S3', 'Menormalkan komponen RoBERTa sebelum digabung dengan cabang BGE.', 1, 1),
    (30, 'self.projection =', 'Proyeksi BGE S3', 'Memadatkan 1024 fitur BGE menjadi 256 fitur.', 1, 1),
    (30, 'self.head = nn.Sequential(', 'Head regresi MLP S3', 'Memetakan fitur fusion melalui MLP menjadi satu skor.', 1, 2),
    (30, 'cls = self.roberta', 'Fitur contextual RoBERTa', 'Menjalankan encoder dan mengambil CLS setiap window.', 1, 1),
    (30, 'essay = mean_by_essay', 'Fitur per jawaban', 'Menggabungkan seluruh CLS window menjadi satu vektor jawaban.', 1, 1),
    (30, "features['concat'] =", 'Fusion yang sebenarnya', 'Menyambung 768 fitur RoBERTa dengan 256 fitur proyeksi.', 1, 1),
    (30, "prediction = self.head(features['concat'])", 'Prediksi fusion S3', 'Mengirim fitur gabungan ke head MLP pemberi skor.', 1, 1),
    (30, 'class Engine', 'Penghubung data dan model', 'Mengubah indeks jawaban menjadi tensor lalu memanggil model.', 1, 1),
    (30, 'self.model.eval()', 'Mode prediksi', 'Menonaktifkan dropout saat membaca model yang sudah dilatih.', 1, 1),
    (30, 'return scaled*', 'Mengembalikan skala skor', 'Mengubah keluaran latihan menjadi poin skor guru 0–100.', 1, 1),
    (31, 'def seed_everything(', 'Seed atau angka acak', 'Mengatur pengacakan Python, NumPy, serta Torch untuk pengulangan.', 1, 1),
    (31, 'def make_model(', 'Model baru setiap fold', 'Mendefinisikan inisialisasi model baru untuk setiap latihan.', 1, 1),
    (34, 'def gwet_ac2(', 'Metrik Gwet AC2', 'Menghitung kesepakatan kategori dengan koreksi peluang Gwet.', 1, 1),
    (35, 'def metrics(', 'Empat metrik evaluasi', 'Menghitung RMSE, Pearson, QWK, dan Gwet AC2.', 1, 1),
    (36, 'def optimizer_groups(', 'Kelompok learning rate', 'Memisahkan encoder, proyeksi, serta head sebelum optimisasi.', 1, 1),
    (36, "rates = {'roberta':", 'Pemilihan LR S3', 'Memakai s3_encoder_lr untuk S3; encoder_lr untuk S1.', 1, 1),
    (36, 'def save_training_progress(', 'Autosave checkpoint', 'Menyimpan keadaan belajar sesudah setiap epoch selesai.', 1, 1),
    (36, 'def load_training_progress(', 'Memulihkan checkpoint', 'Memverifikasi checkpoint terbaru atau mencoba cadangan yang utuh.', 1, 1),
    (36, 'def fit(', 'Definisi training satu model', 'Menyiapkan seluruh proses latihan satu fold atau refit.', 1, 1),
    (36, 'optimizer = torch.optim.AdamW(', 'Optimizer AdamW', 'Menyiapkan mesin pembaruan parameter beserta learning rate-nya.', 1, 1),
    (36, 'targets = torch.as_tensor', 'Normalisasi target guru', 'Mengubah target skor 0–100 menjadi skala latihan 0–1.', 2, 1),
    (36, 'loss = torch.mean(', 'Loss training MSE', 'Menghitung rata-rata kuadrat selisih prediksi dan target.', 1, 1),
    (36, 'scaler.scale(loss*', 'Backward dan akumulasi', 'Menghitung gradien microbatch dengan bobot batch efektif.', 1, 1),
    (36, 'clip_grad_norm_', 'Gradient clipping', 'Membatasi besar gradien sebelum parameter model diperbarui.', 1, 1),
    (36, 'scaler.step(optimizer)', 'Pembaruan bobot', 'Menjalankan langkah optimizer; CPU tidak memakai AMP aktif.', 1, 1),
    (36, 'scheduler.step()', 'Scheduler learning rate', 'Memperbarui learning rate sesuai jadwal langkah training.', 1, 1),
    (36, "if row['validation']['RMSE'] < best_rmse:", 'Pemilihan epoch terbaik', 'Membandingkan RMSE validasi dengan rekor terendah sebelumnya.', 3, 1),
    (38, 'def train_all(', 'Alur CV, refit, test', 'Mengatur latihan dan evaluasi seluruh skenario secara berurutan.', 1, 1),
    (38, 'selected_epochs[scenario] =', 'Menetapkan epoch refit', 'Mengambil pembulatan ke atas median epoch terbaik CV.', 1, 1),
    (38, "checkpoint=output/scenario/'progress_refit.json'", 'Latihan final atau refit', 'Melatih model final berdasarkan epoch yang dipilih CV.', 2, 1),
    (39, 'def ordinal_copy(', 'Pembulatan skor kategori', 'Membuat salinan kategori untuk menghitung QWK dan AC2.', 3, 1),
    (39, 'def test_score_comparison(', 'Nilai guru versus prediksi', 'Menyiapkan perbandingan guru, prediksi mentah, serta pembulatan.', 1, 1),
    (39, 'report=train_all(', 'Menjalankan eksperimen', 'Memanggil CV, refit, dan test atau memulihkan hasil.', 1, 1),
    (41, 'def score_answer(', 'Scoring satu jawaban', 'Menyiapkan prediksi tiga model final tanpa latihan ulang.', 1, 1),
    (41, 'prediction_rows=score_answer(', 'Menjalankan scoring', 'Memanggil prediksi dengan jawaban dan konteks yang sesuai.', 1, 1),
    (43, 'def export_final_features(', 'Fitur setelah training', 'Membaca serta mengekspor vektor dari checkpoint final.', 1, 1),
    (45, 'def build_reports(', 'Merapikan laporan akhir', 'Menggabungkan jurnal teknis menjadi lima workbook utama.', 1, 1),
    (46, 'summary=build_reports(', 'Menyiapkan paket unduhan', 'Menyusun lima Excel serta arsip ZIP hasil akhir.', 1, 1),
    (46, "files.download(download['archive'])", 'Mengunduh hasil', 'Mengirim arsip hasil eksperimen ke perangkat pengguna.', 1, 1),
    (12, 'REUSE_RUN=', 'Folder sumber S1/S2', 'Menentukan sumber kontrol lama yang harus diverifikasi.', 1, 1),
    (30, 'windows, mapping, _ = batch_windows', 'Input RoBERTa jawaban saja', 'Menyiapkan token jawaban; E6 tidak memakai konteks pair.', 2, 1),
    (38, 'def inspect_reuse_runtime(', 'Kompatibilitas versi sumber', 'Memeriksa versi pustaka sumber sebelum memakai model final.', 1, 1),
    (38, 'def inspect_control_source(', 'Verifikasi hasil S1/S2', 'Memeriksa dataset, split, konfigurasi, serta checkpoint kontrol.', 1, 1),
    (38, 'def prepare_reused_controls(', 'Mengambil kontrol S1/S2', 'Memverifikasi dan menyalin hasil kontrol yang sudah selesai.', 1, 1),
    (38, 'reused=prepare_reused_controls(', 'Menjalankan reuse kontrol', 'Memakai S1/S2 terverifikasi tanpa melatihnya kembali.', 1, 1),
    (39, 'reuse=json.loads(', 'Asal model dalam laporan', 'Mencatat asal S1/S2 dan pelatihan baru S3.', 1, 1),
]


def build():
    raw=SOURCE.read_bytes();n=json.loads(raw)
    assert raw==subprocess.check_output(['git','show',PIN+':'+str(SOURCE.relative_to(ROOT))],cwd=ROOT)
    codes={i:''.join(c['source']) for i,c in enumerate(n['cells']) if c['cell_type']=='code'}
    ordinal={i:k+1 for k,i in enumerate(codes)}
    def chapter(i):
        for c in reversed(n['cells'][:i]):
            s=''.join(c['source'])
            if c['cell_type']=='markdown' and s.startswith('## '):return s.splitlines()[0].removeprefix('## ')
        return 'Pengantar'
    entries=[]
    for i,anchor,name,function,length,occurrence in ROWS:
        lines=codes[i].splitlines();hits=[j for j,l in enumerate(lines) if anchor in l]
        if len(hits)<occurrence:raise ValueError(f'Anchor tidak ditemukan: cell{i}: {anchor}')
        start=hits[occurrence-1];snippet='\n'.join(lines[start:start+length])
        # WORKER_SOURCE memuat protokol panjang; tampilkan hanya nama assignment, cocok source.
        if anchor=='WORKER_SOURCE':snippet=lines[start].split('=',1)[0]+'='
        entries.append((i,start,name,function,snippet,f'Bab {chapter(i)}\nBlok kode {ordinal[i]}/30 · baris {start+1}\nCtrl+F: {anchor}'))
    entries.sort(key=lambda e:(e[0],e[1]))
    w=Workbook();ws=w.active;ws.title='Urutan Kode'
    ws.append(['PETA BELAJAR KODE · E6']);ws.merge_cells('A1:F1')
    ws.append(['Jawaban saja + MLP · LR RoBERTa S3 5e-6 · reuse S1/S2']);ws.merge_cells('A2:F2')
    ws.append(['Baca dari atas. Kolom kode disalin dari notebook; Ctrl+F memudahkan pencarian.']);ws.merge_cells('A3:F3')
    ws.append(['Definisi “def/class” menyiapkan fungsi/kelas; pemanggilan menjalankan prosesnya.']);ws.merge_cells('A4:F4')
    headers=['No.','Nama bagian','Fungsi singkat','Kode yang tertulis di notebook','Letak dan cara mencari','Jenis bagian']
    ws.append(headers)
    for no,(i,start,name,function,snippet,location) in enumerate(entries,1):
        kind='Definisi fungsi/kelas' if snippet.lstrip().startswith(('def ','class ')) else 'Pengaturan atau operasi'
        ws.append([no,name,function,snippet,location,kind])
    tab=Table(displayName='UrutanBelajarE6',ref=f'A5:F{ws.max_row}');tab.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);ws.add_table(tab)
    ws.freeze_panes='C6'
    widths=[7,30,49,89,51,25]
    for col,width in enumerate(widths,1):ws.column_dimensions[get_column_letter(col)].width=width
    for row in ws.iter_rows(min_row=6):
        for c in row:c.alignment=Alignment(vertical='top',wrap_text=True);c.font=Font(name='Calibri',size=11,color='15213E')
        row[3].font=Font(name='Consolas',size=10,color='084C72');row[1].font=Font(name='Calibri',size=11,bold=True,color='15213E')
        ws.row_dimensions[row[0].row].height=max(64,16*(row[3].value.count('\n')+2))
    model=w.create_sheet('Bedah Model')
    model.append(['Skenario','Nama bagian','Fungsi','Kode asli','Bentuk fitur / catatan'])
    modelrows=[
      ('S1','Encoder RoBERTa','Membaca jawaban saja.',30,'cls = self.roberta',1,'CLS tiap window mempunyai 768 komponen; mean antar-window menjadi vektor jawaban.'),
      ('S1','Head linear','Mengubah fitur menjadi satu skor.',30,'self.head = nn.Linear(768, 1)',1,'768 → 1. S1 tidak memakai head MLP tambahan.'),
      ('S2','Encoder BGE frozen','Menghasilkan fitur jawaban tetap.',26,'functional.normalize(',1,'1024 komponen. Bobot encoder tidak ikut checkpoint training head.'),
      ('S2','Head MLP','Mempelajari hubungan fitur dengan skor.',30,'self.head = nn.Sequential(',1,'1024 → 256 → GELU → dropout → 1. Head tetap dilatih.'),
      ('S3','Input jawaban saja','Menyiapkan token jawaban untuk RoBERTa.',30,'windows, mapping, _ = batch_windows',1,'E6 tidak memakai konteks pair. S1/S3 membaca jawaban saja; S1 hasil reuse.'),
      ('S3','LayerNorm RoBERTa','Menormalkan komponen sebelum fusion.',30,'self.roberta_norm =',1,'768 → 768; berada di luar encoder.'),
      ('S3','Proyeksi BGE','Memadatkan fitur sebelum fusion.',30,'self.projection =',1,'1024 → 256, kemudian LayerNorm.'),
      ('S3','Concat','Menyambung dua kelompok fitur.',30,"features['concat'] =",1,'768 + 256 = 1024; tidak menjumlahkan skor S1/S2.'),
      ('S3','Head MLP','Mengubah fitur fusion menjadi skor.',30,'self.head = nn.Sequential(',2,'1024 → 256 → GELU → dropout → 1. MLP adalah head, bukan encoder.'),
      ('S1/S3','Learning rate encoder','Mengatur langkah fine-tuning.',36,"rates = {'roberta':",1,'S1=1e-5; S3=5e-6. Head/proyeksi=1e-4.'),
    ]
    for scenario,name,function,i,anchor,occurrence,note in modelrows:
        line=[l for l in codes[i].splitlines() if anchor in l][occurrence-1];model.append([scenario,name,function,line,note])
    for j,width in enumerate([12,30,43,95,65],1):model.column_dimensions[get_column_letter(j)].width=width
    model.freeze_panes='C2';model.auto_filter.ref=model.dimensions
    for row in model.iter_rows(min_row=2):
        model.row_dimensions[row[0].row].height=66
        for c in row:c.alignment=Alignment(vertical='top',wrap_text=True);c.font=Font(name='Consolas' if c.column==4 else 'Calibri',size=10)
    guide=w.create_sheet('Cara Baca')
    instructions=[
      ('Mulai','Buka Urutan Kode. Baca Nama bagian → Fungsi → Kode → Letak. Setiap baris membahas satu bagian, bukan satu program baru. Potongan kode adalah penunjuk belajar, tidak untuk dijalankan terpisah.'),
      ('Mencari di Colab','Gunakan Ctrl+F lalu salin teks setelah “Ctrl+F” pada kolom Letak. Anchor yang sama dapat muncul beberapa kali; pilih bab, blok dan nomor baris yang sesuai.'),
      ('Nomor blok','Blok kode 1–30 hanya menghitung code cell. Bukan nomor bab atau nomor semua cell Markdown+code. Baris dihitung dari baris pertama blok, termasuk %%aes.'),
      ('Definisi vs eksekusi','def menyiapkan fungsi dan class menyiapkan kelas. Cari pemanggilan, misalnya bundle=read_dataset(cfg) atau report=train_all(...), untuk melihat kapan proses berjalan.'),
      ('Fitur dan target','bundle.texts memuat teks jawaban; bundle.y memuat skor guru. E6 tidak memasukkan soal/referensi ke encoder. Bundle adalah wadah data, bukan model.'),
      ('Pooling','Kode menggunakan CLS per window lalu rata-rata antar-window. Tidak menggunakan ContextPooler mDeBERTa dari screenshot pembanding.'),
      ('MLP dan linear','MLP tetap merupakan head regresi. Linear terakhir MLP menghasilkan satu skor. Encoder utama tetap RoBERTa atau BGE.'),
      ('Skala angka','Loss MSE target 0–1. Prediksi dikembalikan ke skor 0–100. RMSE/Pearson memakai prediksi mentah; pembulatan hanya salinan kategori QWK/AC2.'),
      ('Penyimpanan','E6 ini menyimpan file per epoch; JSON merujuk dua checkpoint terbaru sebagai cadangan. File lama masih ada. final.pt dan metadata tetap diperlukan. Ini berbeda dari dua slot A/B E7/E8.'),
      ('Reuse S1/S2','E6 memakai hasil S1/S2 yang sudah selesai. Kode pembentukan model dan training tetap ada, tetapi S1/S2 tidak dilatih ulang pada run E6. Hanya S3 dilatih baru.'),
      ('Tabel ini','Panduan source, bukan hasil training baru. Tidak memuat jawaban, identitas, atau nilai individual siswa. Penjelasan mengikuti E6, bukan MLP awal atau E7 tiga seed.'),
      ('Source GitHub',URL),
      ('SHA256 notebook',hashlib.sha256(raw).hexdigest()),
    ]
    guide.append(['Petunjuk','Penjelasan'])
    for row in instructions:guide.append(row)
    guide.column_dimensions['A'].width=27;guide.column_dimensions['B'].width=120;guide.freeze_panes='A2'
    for row in guide.iter_rows(min_row=2):
        guide.row_dimensions[row[0].row].height=51
        for c in row:c.alignment=Alignment(vertical='top',wrap_text=True);c.font=Font(name='Calibri',size=11)
    
    for row in guide.iter_rows(min_row=2):
        if row[0].value=='Source GitHub':row[1].hyperlink=URL;row[1].style='Hyperlink'
    for sheet,header in [(ws,5),(model,1),(guide,1)]:
        sheet.sheet_view.showGridLines=False;sheet.sheet_view.zoomScale=85
        for c in sheet[header]:c.fill=PatternFill('solid',fgColor='14244D');c.font=Font(name='Calibri',size=11,bold=True,color='FFFFFF');c.alignment=Alignment(wrap_text=True,vertical='center')
        sheet.row_dimensions[header].height=31
        sheet.sheet_properties.pageSetUpPr.fitToPage=True;sheet.page_setup.orientation='landscape';sheet.page_setup.paperSize=sheet.PAPERSIZE_A3;sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0
        sheet.print_title_rows=f'1:{header}'
    for r in (1,2,3,4):
        ws.cell(r,1).font=Font(name='Calibri',size=17 if r==1 else 11,bold=r==1,color='14244D');ws.row_dimensions[r].height=29
    w.properties.title='Panduan Belajar Kode E6';w.properties.subject='Nama bagian, fungsi, source dan lokasi notebook AES';w.properties.creator='Panduan belajar AES'
    TARGET.parent.mkdir(parents=True,exist_ok=True);w.save(TARGET)
    check=load_workbook(TARGET,read_only=True);assert check.sheetnames==['Urutan Kode','Bedah Model','Cara Baca'];assert check.worksheets[0].max_row==5+len(entries);check.close()
    print(json.dumps({'file':str(TARGET.relative_to(ROOT)),'rows':len(entries),'model_rows':len(modelrows),'sheets':3,'source_sha256':hashlib.sha256(raw).hexdigest(),'bytes':TARGET.stat().st_size}))


if __name__=='__main__':build()
